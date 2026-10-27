"""Сервис: выгрузка списка в XLSX (срез 006; PLAN-006, «Выгрузка в Excel»; CONTRACT-006, 5.7, 10.2, 12, 15.6).

Набор записей и столбцы — те же, что у таблицы (`grid.read` → `grid.records`, столбцы сотрудника без скрытых), но без
среза страницы. Книга пишется потоком (`Workbook(write_only=True)`), записи читаются `iterator(chunk_size=500)`;
подписи людей, узлов, записей и имена файлов собираются пачкой на каждые 500 строк (заметка R-6213). Число и деньги —
числовые ячейки (деньги с форматом валюты), дата — ячейка даты, «дата и время» — текст по Москве (так файл загружается
обратно). Весь текст, включая заголовки, защищён от формул апострофом впереди (`FORMULA_START`). В журнал выгрузка не
пишется (FR-542).
"""
import io
from datetime import date
from decimal import Decimal, InvalidOperation
from itertools import islice
from urllib.parse import quote

from django.apps import apps
from django.contrib.auth import get_user_model
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from apps.core.models import UserState
from apps.core.services.audit_search import FORMULA_START
from apps.lists.models import Choice, List, Record, RecordFile
from apps.lists.services import access, grid, schema
from apps.lists.services import values as values_service
from apps.lists.services.errors import ListsError

LIMIT = 10000
CHUNK = 500
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TITLE_HEADER = "Название"
UNAVAILABLE = "запись недоступного списка"
DATE_FORMAT = "DD.MM.YYYY"
MAX_CELL_TEXT = 32767  # предел длины текста ячейки Excel
REFERENCE_TYPES = ("choice", "person", "unit", "record")


def _safe(value) -> str:
    """Текст ячейки: без запрещённых в XML знаков, не длиннее предела Excel, с апострофом перед `=`, `+`, `-`, `@`,
    табуляцией и возвратом каретки."""
    text = ILLEGAL_CHARACTERS_RE.sub("", "" if value is None else str(value))[:MAX_CELL_TEXT]
    return "'" + text if text.startswith(FORMULA_START) else text


def _money_format(field) -> str:
    sign = values_service.CURRENCY_SIGNS.get((field.settings or {}).get("currency", "RUB"), "₽")
    return f'#,##0.00 "{sign}"'


class _Labels:
    """Подписи ссылочных значений: кэш на всю выгрузку, в базу — только за ещё не виденными кодами (пачкой)."""

    def __init__(self, viewer):
        self.viewer = viewer
        self.cache: dict[tuple[str, int], str | None] = {}
        self.readable: dict[int, bool] = {}

    def fill(self, field, ids: set[int]) -> None:
        kind = field.type
        fresh = [pk for pk in ids if self._key(field, pk) not in self.cache]
        if not fresh:
            return
        found: dict[int, str] = {}
        if kind == "choice":
            found = dict(Choice.objects.filter(field=field, pk__in=fresh).values_list("pk", "title"))
        elif kind == "person":
            # уволенный — с пометкой, как на экране: иначе при обратной загрузке ФИО молча нашло бы однофамильца
            # среди действующих (правка Opus по R-6220-1)
            found = {pk: name + (" (уволен)" if state == UserState.DISMISSED else "")
                     for pk, name, state in get_user_model().objects.filter(pk__in=fresh).values_list(
                         "pk", "full_name", "state")}
        elif kind == "unit":
            refs = {}
            if apps.is_installed("apps.org"):
                from apps.org.services import structure

                refs = structure.unit_refs(fresh)
            found = {pk: refs[pk].name if pk in refs else f"узел №{pk}" for pk in fresh}
        else:
            for pk, title, list_id in Record.objects.filter(pk__in=fresh).values_list("pk", "title", "list_id"):
                if list_id not in self.readable:
                    self.readable[list_id] = access.can_read(self.viewer, List(pk=list_id))
                found[pk] = title if self.readable[list_id] else UNAVAILABLE
        for pk in fresh:
            self.cache[self._key(field, pk)] = found.get(pk)

    @staticmethod
    def _key(field, pk: int) -> tuple[str, int]:
        """Подпись списка значений зависит от поля, остальные — только от типа."""
        return (f"choice{field.pk}" if field.type == "choice" else field.type, pk)

    def text(self, field, ids) -> str:
        names = (self.cache.get(self._key(field, pk)) for pk in ids if isinstance(pk, int))
        return "; ".join(name for name in names if name)


def _cell(sheet, field, stored, labels: _Labels, files: dict):
    """Ячейка по типу поля: `None` — пустая, число/деньги/дата — не текстом, остальное — защищённым текстом."""
    kind = field.type
    if kind == "bool":
        return _text_cell(sheet, "Да" if stored else "Нет")
    if kind == "file":
        return _text_cell(sheet, "; ".join(files.get(field.pk, [])))
    if stored is None or stored == "" or stored == []:
        return None
    if kind in ("number", "money"):
        try:
            number = Decimal(str(stored))
        except InvalidOperation:
            return _text_cell(sheet, stored)
        cell = WriteOnlyCell(sheet, value=number)
        if kind == "money":
            cell.number_format = _money_format(field)
        return cell
    if kind == "date":
        try:
            moment = date.fromisoformat(str(stored))
        except ValueError:
            return _text_cell(sheet, stored)
        cell = WriteOnlyCell(sheet, value=moment)
        cell.number_format = DATE_FORMAT
        return cell
    if kind == "datetime":
        return _text_cell(sheet, values_service.display(field, stored, None))
    if kind in REFERENCE_TYPES:
        return _text_cell(sheet, labels.text(field, stored if isinstance(stored, (list, tuple)) else []))
    return _text_cell(sheet, stored)


def _text_cell(sheet, value):
    text = _safe(value)
    return WriteOnlyCell(sheet, value=text) if text else None


def _width(field) -> int:
    if field is None:
        return 40
    return {"text": 40, "string": 24, "record": 30, "file": 30}.get(field.type, 18)


def _chunks(queryset):
    rows = queryset.iterator(chunk_size=CHUNK)
    while True:
        part = list(islice(rows, CHUNK))
        if not part:
            return
        yield part


def build(lst, viewer, params) -> HttpResponse:
    """Книга XLSX по параметрам таблицы; записей больше `LIMIT` — `ListsError("export_too_many")`.

    Нет права читать список — `Http404` (из `grid.read`)."""
    gp = grid.read(params, lst, viewer)
    queryset = grid.records(lst, viewer, gp)
    if queryset.count() > LIMIT:
        raise ListsError("export_too_many")
    by_code = {field.code: field for field in schema.visible_fields(lst)}
    columns = [None if code == "title" else by_code[code] for code in grid.columns(lst, viewer)
               if code == "title" or code in by_code]
    fields = [field for field in columns if field is not None]
    reference_fields = [field for field in fields if field.type in REFERENCE_TYPES]
    file_fields = [field.pk for field in fields if field.type == "file"]
    labels = _Labels(viewer)

    book = Workbook(write_only=True)
    sheet = book.create_sheet("Список")
    for index, field in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = _width(field)
    sheet.append([_text_cell(sheet, TITLE_HEADER if field is None else field.name) for field in columns])
    for part in _chunks(queryset):
        for field in reference_fields:
            ids = {pk for row in part for pk in (row.values or {}).get(field.code) or [] if isinstance(pk, int)}
            labels.fill(field, ids)
        files: dict[int, dict[int, list[str]]] = {}
        if file_fields:
            found = (RecordFile.objects.filter(record_id__in=[row.pk for row in part], field_id__in=file_fields)
                     .order_by("uploaded_at", "pk").values_list("record_id", "field_id", "name"))
            for record_id, field_id, name in found:
                files.setdefault(record_id, {}).setdefault(field_id, []).append(name)
        for row in part:
            stored = row.values or {}
            sheet.append([_text_cell(sheet, row.title) if field is None else
                          _cell(sheet, field, stored.get(field.code), labels, files.get(row.pk, {}))
                          for field in columns])
    buffer = io.BytesIO()
    book.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_TYPE)
    response["Content-Disposition"] = "attachment; filename*=UTF-8''" + quote(f"{lst.name}.xlsx", safe="")
    return response
