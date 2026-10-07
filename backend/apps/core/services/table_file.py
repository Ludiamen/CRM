"""Чтение файла таблицы CSV и XLSX (PLAN-003, «Чтение файла таблицы», FR-262, FR-265).

Общее для загрузки сотрудников и загрузки заявителей: формат определяется по содержимому, оглавление архива
проверяется до чтения, `openpyxl` читает книгу только для чтения и без формул, CSV читается строгим UTF-8.
Пределы (размер, строки, части и распакованный объём архива) и тексты отказов — параметры вызывающего;
своих пределов и текстов здесь нет. Файл разбирается в памяти.
"""
import csv
import datetime
import io
import zipfile
from dataclasses import dataclass
from decimal import Decimal


class TableFileError(Exception):
    """Отказ чтения таблицы: `code` — код (`too_big`, `bad_format`, `not_utf8`, `too_many_rows`, `zip_bomb`,
    `xml_bomb`), `message` — текст из `texts` вызывающего; `str(error) == message`."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class TableFile:
    """Прочитанная таблица: формат по содержимому, строка заголовков, строки данных (пустые не считаются)."""

    file_format: str
    header: list[str]
    rows: list[list[str]]


def _cell_text(value) -> str:
    """Значение ячейки XLSX текстом: дата — ДД.ММ.ГГГГ, целое — без «.0» и экспоненты, дробное — без экспоненты."""
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(repr(value)), "f")
    return str(value)


def _xlsx_error(error: Exception, bomb_type, texts: dict) -> TableFileError:
    """Ошибка чтения книги -> `xml_bomb`, если в цепочке причин (openpyxl оборачивает) есть отказ defusedxml."""
    link, hops = error, 0
    while link is not None and hops < 10:
        if bomb_type and isinstance(link, bomb_type):
            return TableFileError("xml_bomb", texts["xml_bomb"])
        link, hops = link.__cause__ or link.__context__, hops + 1
    return TableFileError("bad_format", texts["bad_format"])


def _read_xlsx(content: bytes, max_rows: int, texts: dict) -> list[list[str]]:
    """Первый лист книги: заголовок и строки данных, не больше `max_rows + 1` строк данных (читаем лениво)."""
    try:
        from defusedxml.common import DefusedXmlException
    except ImportError:  # defusedxml в зависимостях; без него защиты нет, но и класса нет
        DefusedXmlException = ()
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise _xlsx_error(error, DefusedXmlException, texts) from None
    try:
        sheet = book.worksheets[0]
        out = []
        for raw in sheet.iter_rows(values_only=True):
            cells = [_cell_text(value) for value in raw]
            if out and not any(cell.strip() for cell in cells):
                continue  # пустые строки данных не считаются
            out.append(cells)
            if len(out) > max_rows + 1:  # заголовок и строка сверх предела: дальше не читаем
                break
        return out
    except IndexError:
        return []
    except Exception as error:
        raise _xlsx_error(error, DefusedXmlException, texts) from None
    finally:
        book.close()


def _read_csv(content: bytes, max_rows: int, texts: dict) -> list[list[str]]:
    """CSV: строгое UTF-8 (с BOM или без) без нулевых байтов; разделитель «;» или «,» по строке заголовков."""
    if b"\x00" in content:
        raise TableFileError("not_utf8", texts["not_utf8"])
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise TableFileError("not_utf8", texts["not_utf8"]) from None
    first = text.split("\n", 1)[0]
    delimiter = ";" if first.count(";") >= first.count(",") and ";" in first else ","
    out = []
    try:
        for raw in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter):
            if out and not any(cell.strip() for cell in raw):
                continue
            out.append(raw)
            if len(out) > max_rows + 1:
                break
    except csv.Error:
        raise TableFileError("bad_format", texts["bad_format_csv"]) from None
    return out


def _open_archive(content: bytes):
    """ZIP-архив или `None`, если содержимое не архив."""
    try:
        return zipfile.ZipFile(io.BytesIO(content))
    except (zipfile.BadZipFile, ValueError, OSError):
        return None


def read(content: bytes, *, max_bytes: int, max_rows: int, max_parts: int, max_unpacked: int,
         texts: dict) -> TableFile:
    """Читает файл таблицы. Порядок отказов: размер, формат (архив не XLSX), оглавление архива, чтение.

    Строк данных отдаётся не больше `max_rows + 1`; о числе строк `read` не судит — это делает `check_rows`
    после проверки колонок вызывающим.
    """
    if len(content) > max_bytes:
        raise TableFileError("too_big", texts["too_big"])
    archive = _open_archive(content)
    if archive is not None:
        with archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise TableFileError("bad_format", texts["bad_format_archive"])
            parts = archive.infolist()
            if len(parts) > max_parts or sum(part.file_size for part in parts) > max_unpacked:
                raise TableFileError("zip_bomb", texts["zip_bomb"])
        file_format, table = "xlsx", _read_xlsx(content, max_rows, texts)
    else:
        file_format, table = "csv", _read_csv(content, max_rows, texts)
    return TableFile(file_format=file_format, header=table[0] if table else [], rows=table[1:])


def check_rows(rows, max_rows: int, texts: dict) -> None:
    """Отказ `too_many_rows`, если строк данных больше `max_rows`."""
    if len(rows) > max_rows:
        raise TableFileError("too_many_rows", texts["too_many_rows"])
