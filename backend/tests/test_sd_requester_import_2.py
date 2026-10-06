"""T-3128: загрузка заявителей — подтверждение, флажок удаления, журнал, страница (FR-261…266). Сценарий 10 спеки среза 003.

Службы — `apps.servicedesk.services.requester_import` и `retention` (CONTRACT-003 § 6.23, 6.24), маршруты — § 12.6,
зацепки — § 13.8, журнал — § 9.2, тексты — § 14.2. Предпросмотр, разбор файла и отказы файла — в `test_sd_requester_import`
(T-3127). Сбой строки — подмена `requester_import.apply_row` (§ 3.5).
"""
import csv
import io
import json
import re
import secrets
from html.parser import HTMLParser
from pathlib import Path

import pytest
from django.conf import settings as django_settings
from django.core.files.uploadedfile import SimpleUploadedFile

from tests import support
from tests import support_org as org
from tests import support_servicedesk as sd
from tests.support import model

pytestmark = pytest.mark.django_db

DOMAIN = "corp.example"
HEADERS = ["Почта", "Фамилия", "Имя", "Отчество", "Подразделение"]
IMPORT = "/servicedesk/requesters/import/"
HELP_TEXT = ("ФИО, почта и подразделение работников общества хранятся, чтобы определять подразделение "
             "заявителя по адресу письма")
CSV_TYPE = "text/csv; charset=utf-8"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# --- разбор разметки (помощник только этого файла) ---------------------------------------------

class _El:
    def __init__(self, tag, attrs, parent):
        self.tag, self.attrs, self.parent, self.kids = tag, attrs, parent, []

    def walk(self):
        yield self
        for kid in self.kids:
            if isinstance(kid, _El):
                yield from kid.walk()

    def text(self) -> str:
        parts = [kid if isinstance(kid, str) else kid.text() for kid in self.kids]
        return " ".join(" ".join(parts).split())

    def find_all(self, tag=None, **want):
        """Элементы с тегом `tag` и атрибутами `want` (`data_x` -> `data-x`; True — любое значение)."""
        wanted = {key.replace("_", "-"): value for key, value in want.items()}
        found = []
        for el in self.walk():
            if tag is not None and el.tag != tag:
                continue
            if all(name in el.attrs and (value is True or el.attrs[name] == str(value)) for name, value in wanted.items()):
                found.append(el)
        return found


class _Builder(HTMLParser):
    _VOID = {"input", "br", "img", "hr", "meta", "link", "area", "base", "col", "source", "wbr"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = _El("#root", {}, None)
        self.cur = self.root

    def handle_starttag(self, tag, attrs):
        el = _El(tag, {k: (v if v is not None else "") for k, v in attrs}, self.cur)
        self.cur.kids.append(el)
        if tag not in self._VOID:
            self.cur = el

    def handle_startendtag(self, tag, attrs):
        self.cur.kids.append(_El(tag, {k: (v if v is not None else "") for k, v in attrs}, self.cur))

    def handle_endtag(self, tag):
        node = self.cur
        while node is not self.root and node.tag != tag:
            node = node.parent
        if node is not self.root:
            self.cur = node.parent

    def handle_data(self, data):
        self.cur.kids.append(data)


def dom(response) -> _El:
    builder = _Builder()
    builder.feed(response.content.decode("utf-8", "replace"))
    return builder.root


def body(response) -> bytes:
    if getattr(response, "streaming", False):
        return b"".join(response.streaming_content)
    return response.content


# --- данные и вызовы ----------------------------------------------------------------------------

def service():
    from apps.servicedesk.services import requester_import

    return requester_import


def retention():
    from apps.servicedesk.services import retention as module

    return module


def admin():
    """Администратор загрузки: право «Настройка сервис-деска» (`servicedesk.manage`)."""
    return support.make_with_permissions("servicedesk.manage", full_name="Загружающий Админ")


def setup(*extra_domains: str) -> None:
    """Список корпоративных доменов: приём заявок нужен только для проверки домена строк."""
    sd.configure_servicedesk(domains=[DOMAIN, *extra_domains])


def addr(tag: str = "u", domain: str = DOMAIN) -> str:
    return f"{tag}-{secrets.token_hex(3)}@{domain}"


def line(email: str, last: str = "", first: str = "", middle: str = "", unit: str = "") -> list:
    return [email, last, first, middle, unit]


def full(last: str = "", first: str = "", middle: str = "") -> str:
    return " ".join(part for part in (last, first, middle) if part)


def csv_of(*rows) -> bytes:
    return org.staff_csv([list(r) for r in rows], headers=HEADERS)


def unit(name: str, *, archived: bool = False):
    ClientUnit = model("servicedesk.ClientUnit")
    return ClientUnit.objects.create(name=name, name_key=org.name_key(name), archived=archived)


def preview(actor, *rows, name: str = "people.csv"):
    return service().make_preview(csv_of(*rows), file_name=name, actor=actor)


def confirm(item, actor, *, delete_missing: bool = False):
    return service().confirm(item, delete_missing=delete_missing, actor=actor)


def numbered(item) -> dict:
    return {r["number"]: r for r in item.rows}


def reload(item):
    return model("servicedesk.RequesterImport").objects.get(pk=item.pk)


def requester(email: str):
    return model("servicedesk.Requester").objects.filter(email=email.lower()).first()


def events(action=None, *, prefix=None, **filters) -> list:
    qs = model("core.AuditEvent").objects.filter(**filters).order_by("id")
    if action:
        qs = qs.filter(action=action)
    if prefix:
        qs = qs.filter(action__startswith=prefix)
    return list(qs)


def import_events() -> int:
    return len(events(prefix="requester")) + len(events(prefix="client_unit"))


def with_ticket(req):
    return sd.make_ticket(requester=req)


def page(actor, item):
    response = support.logged_in(actor).get(f"{IMPORT}{item.pk}/")
    assert response.status_code == 200, "страница загрузки не открылась"
    return response, dom(response)


def files_tree() -> list:
    found = []
    root = django_settings.PROTECTED_ROOT
    if Path(root).exists():
        found += [str(p) for p in Path(root).rglob("*") if p.is_file()]
    return sorted(found)


def digits_of(el, attr: str) -> set:
    """Числа, названные элементом: значение атрибута и целые числа его текста."""
    return {el.attrs.get(attr, "")} | set(re.findall(r"\d+", el.text()))


# --- доступ и страница --------------------------------------------------------------------------

def test_import_pages_403_without_manage_right():
    """FR-262: все пути загрузки заявителей без права «Настройка сервис-деска» — 403 (в том числе у несуществующего
    номера), без входа — 302; POST без права ничего не создаёт и не подтверждает."""
    setup()
    boss = admin()
    item = preview(boss, line(addr("a"), "Иванов", "Иван"))
    Requester = model("servicedesk.Requester")
    Import = model("servicedesk.RequesterImport")
    before_imports, before_requesters = Import.objects.count(), Requester.objects.count()
    for user in (support.make_with_permissions("servicedesk.all"), support.make_employee()):
        client = support.logged_in(user)
        for path in (IMPORT, f"{IMPORT}{item.pk}/", f"{IMPORT}{item.pk + 1000}/", f"{IMPORT}sample.csv",
                     f"{IMPORT}sample.xlsx"):
            assert client.get(path).status_code == 403, f"GET {path}"
        upload = SimpleUploadedFile("people.csv", csv_of(line(addr("b"))), content_type="text/csv")
        assert client.post(IMPORT, {"file": upload}).status_code == 403
        assert client.post(f"{IMPORT}{item.pk}/", {}).status_code == 403
    assert Import.objects.count() == before_imports
    assert reload(item).state == "preview"
    assert Requester.objects.count() == before_requesters
    anonymous = support.make_client()
    for path in (IMPORT, f"{IMPORT}{item.pk}/", f"{IMPORT}sample.csv"):
        response = anonymous.get(path)
        assert response.status_code == 302, f"без входа {path}"


def test_import_page_text_names_data_and_purpose():
    """FR-266: страница загрузки называет хранимые данные и цель дословной фразой; есть форма выбора файла и ссылки на
    образцы CSV и XLSX."""
    boss = admin()
    response = support.logged_in(boss).get(IMPORT)
    assert response.status_code == 200
    root = dom(response)
    helps = root.find_all(data_format_help=True)
    assert helps, "нет [data-format-help]"
    assert HELP_TEXT in " ".join(helps[0].text().split())
    forms = root.find_all("form", data_form="import-upload")
    assert len(forms) == 1
    files = [i for i in forms[0].find_all("input", name="file") if i.attrs.get("type") == "file"]
    assert len(files) == 1
    assert root.find_all("a", data_sample="csv")[0].attrs["href"] == f"{IMPORT}sample.csv"
    assert root.find_all("a", data_sample="xlsx")[0].attrs["href"] == f"{IMPORT}sample.xlsx"


def test_sample_csv_and_xlsx_headers_and_two_fictional_rows():
    """FR-262: образцы отдаются вложением с нужным типом; в них пять заголовков «Почта», «Фамилия», «Имя», «Отчество»,
    «Подразделение» и две строки; CSV — UTF-8 с BOM и разделителем «;»; образцы читаются самой загрузкой."""
    from openpyxl import load_workbook

    service_ = service()
    client = support.logged_in(admin())
    got = client.get(f"{IMPORT}sample.csv")
    assert got.status_code == 200
    assert got["Content-Type"] == CSV_TYPE
    assert got["Content-Disposition"].lower().startswith("attachment")
    raw = body(got)
    assert raw == service_.sample_csv()
    assert raw.startswith(b"\xef\xbb\xbf"), "CSV образца — UTF-8 с BOM"
    table = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"), newline=""), delimiter=";"))
    table = [r for r in table if any(r)]
    assert table[0] == HEADERS
    assert len(table) == 3, "заголовок и две строки"
    assert all(r[0] and "@" in r[0] for r in table[1:])
    parsed = service_.parse_file(raw)
    assert parsed.file_format == "csv" and len(parsed.rows) == 2

    got = client.get(f"{IMPORT}sample.xlsx")
    assert got.status_code == 200
    assert got["Content-Type"] == XLSX_TYPE
    assert got["Content-Disposition"].lower().startswith("attachment")
    raw = body(got)
    sheet = load_workbook(io.BytesIO(raw), read_only=True).active
    lines = [["" if c is None else str(c) for c in r] for r in sheet.iter_rows(values_only=True)]
    lines = [r for r in lines if any(r)]
    assert lines[0] == HEADERS
    assert len(lines) == 3
    parsed = service_.parse_file(raw)
    assert parsed.file_format == "xlsx" and len(parsed.rows) == 2


# --- подтверждение ------------------------------------------------------------------------------

def test_confirm_skips_error_rows_applies_others_with_total():
    """FR-263: подтверждение пропускает строки с ошибкой (нет почты, чужой домен, повтор почты в другом регистре),
    применяет остальные; итог «добавлено 2, обновлено 0, без изменений 1, с ошибками 3» — в записи и на странице."""
    setup()
    boss = admin()
    known = sd.make_requester(addr("old"), full_name=full("Сидоров", "Пётр"))
    a, b, foreign = addr("a"), addr("b"), addr("x", "other.example")
    item = preview(boss, line(a, "Иванов", "Иван"), line(b, "Петров", "Пётр"),
                   line(known.email, "Сидоров", "Пётр"), line("", "Безпочты"), line(foreign, "Чужой"),
                   line(a.upper(), "Повторов"))
    Requester = model("servicedesk.Requester")
    assert Requester.objects.count() == 1, "предпросмотр ничего не записывает"
    done = confirm(item, boss)
    assert done.state == "done"
    rows = numbered(done)
    assert [rows[n]["result"] for n in range(1, 7)] == ["added", "added", "unchanged", "error", "error", "error"]
    assert rows[6]["reason"] == "ошибка: повтор почты"
    assert all(rows[n]["reason"].startswith("ошибка") for n in (4, 5, 6))
    for key, value in {"added": 2, "updated": 0, "unchanged": 1, "errors": 3, "deleted": 0}.items():
        assert done.summary[key] == value, key
    assert requester(a).full_name == "Иванов Иван" and requester(b).full_name == "Петров Пётр"
    assert requester(foreign) is None
    assert Requester.objects.count() == 3
    _, root = page(boss, done)
    summary = root.find_all(data_result_summary=True)[0]
    assert "добавлено 2, обновлено 0, без изменений 1, с ошибками 3" in summary.text()
    assert (summary.attrs["data-added"], summary.attrs["data-errors"]) == ("2", "3")


def test_confirm_rows_revalidated():
    """FR-263: строки проверяются заново при подтверждении: подразделение, ушедшее в архив, чужой домен после правки
    списка и запись, появившаяся после предпросмотра (вместо «будет добавлен» — «обновлён»)."""
    setup("second.example")
    boss = admin()
    stock = unit("Склад")
    first, second, third = addr("r1"), addr("r2"), addr("r3", "second.example")
    item = preview(boss, line(first, "Первов", "Иван", unit="Склад"), line(second, "Новый", "Имя"),
                   line(third, "Третьев", "Тим"))
    assert numbered(item)[2]["result"] == "will_add"
    assert all(r["result"] == "will_add" for r in item.rows)
    model("servicedesk.ClientUnit").objects.filter(pk=stock.pk).update(archived=True)
    sd.make_requester(second, full_name="Старый Имя")
    setup()
    done = confirm(item, boss)
    rows = numbered(done)
    assert rows[1]["result"] == "error" and rows[1]["reason"] == "ошибка: подразделение в архиве"
    assert rows[2]["result"] == "updated"
    assert rows[3]["result"] == "error"
    assert requester(first) is None and requester(third) is None
    assert requester(second).full_name == "Новый Имя"
    assert (done.summary["added"], done.summary["updated"], done.summary["errors"]) == (0, 1, 2)


def test_confirm_each_valid_row_in_own_savepoint(monkeypatch):
    """FR-263: каждая годная строка пишется в своей точке сохранения: сбой `apply_row` у одной строки (после частичной
    записи) откатывает только её, соседние строки применены, итог — одна ошибка."""
    setup()
    boss = admin()
    module = service()
    real = module.apply_row
    first, bad, last = addr("one"), addr("bad"), addr("three")
    called = []
    Requester = model("servicedesk.Requester")

    def flaky(row, *, actor):
        called.append(row["email"])
        assert actor == boss
        if row["email"] == bad:
            sd.make_requester(bad, full_name="Частично записан", origin="table")
            raise RuntimeError("сбой строки")
        return real(row, actor=actor)

    item = preview(boss, line(first, "Первый"), line(bad, "Сбойный"), line(last, "Третий"))
    monkeypatch.setattr(module, "apply_row", flaky)
    done = confirm(item, boss)
    assert sorted(called) == sorted([first, bad, last])
    rows = numbered(done)
    assert [rows[n]["result"] for n in (1, 2, 3)] == ["added", "error", "added"]
    assert rows[2]["reason"].startswith("ошибка")
    assert requester(first) is not None and requester(last) is not None
    assert requester(bad) is None, "частичная запись сбойной строки откатана"
    assert (done.summary["added"], done.summary["errors"]) == (2, 1)
    assert sorted(e.object_id for e in events("requester.created")) == sorted(
        str(requester(email).pk) for email in (first, last))
    assert done.state == "done"


def test_confirm_creates_missing_units():
    """FR-263: недостающие подразделения создаются при подтверждении (по первому написанию, один раз), заявители
    получают их; существующее подразделение не дублируется."""
    setup()
    boss = admin()
    existing = unit("Склад")
    ClientUnit = model("servicedesk.ClientUnit")
    a, b, c = addr("a"), addr("b"), addr("c")
    item = preview(boss, line(a, "Один", unit="Новый отдел"), line(b, "Два", unit="НОВЫЙ отдел"),
                   line(c, "Три", unit="Склад"))
    assert item.new_units == ["Новый отдел"]
    assert ClientUnit.objects.count() == 1
    confirm(item, boss)
    assert ClientUnit.objects.count() == 2
    created = ClientUnit.objects.exclude(pk=existing.pk).get()
    assert created.name == "Новый отдел"
    assert requester(a).client_unit_id == created.pk and requester(b).client_unit_id == created.pk
    assert requester(c).client_unit_id == existing.pk


def test_empty_name_and_unit_cells_do_not_erase_values():
    """FR-263: пустые ФИО и подразделение в файле прежние значения не стирают (строка «без изменений»); одна пустая
    часть ФИО — не пустое ФИО: заполненная часть заменяет имя."""
    setup()
    boss = admin()
    stock = unit("Склад")
    keep = sd.make_requester(addr("keep"), full_name="Иванов Иван", client_unit=stock, origin="table")
    part = sd.make_requester(addr("part"), full_name="Старов Иван", client_unit=stock, origin="table")
    item = preview(boss, line(keep.email), line(part.email, "Петров"))
    done = confirm(item, boss)
    rows = numbered(done)
    assert rows[1]["result"] == "unchanged" and rows[2]["result"] == "updated"
    keep.refresh_from_db()
    part.refresh_from_db()
    assert keep.full_name == "Иванов Иван" and keep.client_unit_id == stock.pk
    assert part.full_name == "Петров" and part.client_unit_id == stock.pk
    assert (done.summary["unchanged"], done.summary["updated"]) == (1, 1)
    assert [e.object_id for e in events("requester.updated")] == [str(part.pk)]


def test_same_file_again_all_rows_unchanged():
    """FR-263: тот же файл повторно — все строки «без изменений», ничего нового не создано и в журнал не записано."""
    setup()
    boss = admin()
    rows = [line(addr("a"), "Первов", "Иван", "Петрович", "Отдел А"), line(addr("b"), "Второв", unit="Отдел Б"),
            line(addr("c"), "Третьев", "Тим")]
    confirm(preview(boss, *rows), boss)
    Requester, ClientUnit = model("servicedesk.Requester"), model("servicedesk.ClientUnit")
    counts = (Requester.objects.count(), ClientUnit.objects.count())
    base = max(e.id for e in events())
    again = preview(boss, *rows)
    assert all(r["result"] == "unchanged" for r in again.rows)
    assert again.new_units == []
    done = confirm(again, boss)
    assert all(r["result"] == "unchanged" for r in done.rows)
    assert (done.summary["added"], done.summary["updated"], done.summary["unchanged"], done.summary["errors"]) == (0, 0, 3, 0)
    assert (Requester.objects.count(), ClientUnit.objects.count()) == counts
    assert [e.action for e in events() if e.id > base] == ["requester_import.completed"]


def test_repeat_confirmation_shows_previous_total_writes_nothing():
    """FR-263: повторное подтверждение выполненной загрузки (в том числе с флажком) — тот же ответ 302 на тот же адрес,
    итог прежний, ничего не записано и не удалено."""
    setup()
    boss = admin()
    item = preview(boss, line(addr("a"), "Первов"), line(addr("b"), "Второв"))
    client = support.logged_in(boss)
    path = f"{IMPORT}{item.pk}/"
    first = client.post(path, {})
    assert first.status_code == 302 and first["Location"] == path
    done = reload(item)
    assert done.state == "done" and done.summary["added"] == 2
    survivor = sd.make_requester(addr("later"), full_name="Позже Пришёл", origin="table")
    Requester = model("servicedesk.Requester")
    counts = (Requester.objects.count(), len(events(prefix="requester")), len(events(prefix="client_unit")))
    second = client.post(path, {"delete_missing": "on"})
    assert second.status_code == 302 and second["Location"] == path
    again = reload(item)
    assert again.summary == done.summary and again.delete_missing == done.delete_missing
    assert Requester.objects.filter(pk=survivor.pk).exists(), "повтор с флажком ничего не удаляет"
    assert (Requester.objects.count(), len(events(prefix="requester")), len(events(prefix="client_unit"))) == counts
    assert len(events("requester_import.completed")) == 1
    assert service().confirm(reload(item), delete_missing=True, actor=boss).summary == done.summary
    _, root = page(boss, item)
    assert root.find_all(data_result_summary=True)[0].attrs["data-added"] == "2"


# --- флажок удаления ----------------------------------------------------------------------------

def test_flag_off_deletes_nobody():
    """FR-264: флажок выключен (по умолчанию) — никто не удалён, даже записи «таблица» без заявок, которых нет в файле;
    предпросмотр при этом сообщает, сколько удалил бы флажок."""
    setup()
    boss = admin()
    lonely = sd.make_requester(addr("lone"), full_name="Одинокий Иван", origin="table")
    item = preview(boss, line(addr("new"), "Новый"))
    assert item.delete_missing is False and item.summary["would_delete"] == 1
    done = confirm(item, boss)
    assert done.summary["deleted"] == 0
    assert model("servicedesk.Requester").objects.filter(pk=lonely.pk).exists()
    assert events("requester.deleted") == []


def test_flag_on_deletes_table_records_without_tickets_not_in_file():
    """FR-264: флажок включён — удалены записи «таблица» без заявок, чьих адресов нет в файле; запись, чей адрес в файле,
    остаётся; в журнале — `requester.deleted` с причиной «загрузка таблицы»."""
    setup()
    boss = admin()
    gone1 = sd.make_requester(addr("g1"), full_name="Уйдёт Один", origin="table")
    gone2 = sd.make_requester(addr("g2"), full_name="Уйдёт Два", origin="table")
    stay = sd.make_requester(addr("stay"), full_name="Остаётся Иван", origin="table")
    ids = (gone1.pk, gone2.pk)
    item = preview(boss, line(stay.email, "Остаётся", "Иван"))
    assert item.summary["would_delete"] == 2
    done = confirm(item, boss, delete_missing=True)
    Requester = model("servicedesk.Requester")
    assert done.delete_missing is True and done.summary["deleted"] == 2
    assert not Requester.objects.filter(pk__in=ids).exists()
    assert Requester.objects.filter(pk=stay.pk).exists()
    deleted = events("requester.deleted")
    assert sorted(e.object_id for e in deleted) == sorted(str(i) for i in ids)
    for event in deleted:
        assert event.object_type == "requester" and event.actor_id == boss.pk
        assert event.before["origin"] == "table"
        assert event.after["reason"] == "загрузка таблицы"


def test_flag_counts_all_nonempty_rows_including_error_rows_as_in_file():
    """FR-264: «есть в файле» — адрес любой непустой строки, в том числе строки с ошибкой: запись, чей адрес встречается
    только в строке с ошибкой, не удаляется и не считается в `would_delete`."""
    setup()
    boss = admin()
    unit("Архивный", archived=True)
    in_error = sd.make_requester(addr("err"), full_name="Ошибочный Иван", origin="table")
    absent = sd.make_requester(addr("abs"), full_name="Отсутствует Иван", origin="table")
    item = preview(boss, line(in_error.email, "Ошибочный", unit="Архивный"))
    assert numbered(item)[1]["result"] == "error"
    assert item.summary["would_delete"] == 1
    done = confirm(item, boss, delete_missing=True)
    Requester = model("servicedesk.Requester")
    assert Requester.objects.filter(pk=in_error.pk).exists()
    assert not Requester.objects.filter(pk=absent.pk).exists()
    assert done.summary["deleted"] == 1 and done.summary["errors"] == 1


def test_flag_never_deletes_ticket_origin_records():
    """FR-264: записи с отметкой «заявка» флажок не трогает: не удаляются и не меняют отметку, `would_delete` их не
    считает."""
    setup()
    boss = admin()
    mine = sd.make_requester(addr("tick"), full_name="Из Заявки", origin="ticket")
    item = preview(boss, line(addr("new"), "Новый"))
    assert item.summary["would_delete"] == 0
    done = confirm(item, boss, delete_missing=True)
    mine.refresh_from_db()
    assert mine.origin == "ticket"
    assert done.summary["deleted"] == 0


def test_flag_table_record_with_tickets_becomes_ticket_origin_not_deleted():
    """FR-261, FR-264: запись «таблица» с заявками, адреса которой нет в файле, при флажке не удаляется, а получает
    отметку «заявка»; заявка остаётся связанной с записью; запись без заявок удаляется."""
    setup()
    boss = admin()
    kept = sd.make_requester(addr("kept"), full_name="С Заявкой", origin="table")
    ticket = with_ticket(kept)
    free = sd.make_requester(addr("free"), full_name="Без Заявок", origin="table")
    item = preview(boss, line(addr("new"), "Новый"))
    done = confirm(item, boss, delete_missing=True)
    Requester = model("servicedesk.Requester")
    kept.refresh_from_db()
    assert kept.origin == "ticket"
    ticket.refresh_from_db()
    assert ticket.requester_id == kept.pk
    assert not Requester.objects.filter(pk=free.pk).exists()
    assert done.summary["deleted"] == 1
    assert done.summary["origin_changed"] == 1
    switched = [e for e in events("requester.updated") if e.object_id == str(kept.pk)]
    assert len(switched) == 1
    assert switched[0].before == {"origin": "table"} and switched[0].after == {"origin": "ticket"}


def test_import_ticket_origin_record_becomes_table_origin_not_counted_updated():
    """FR-261, FR-263: адрес записи «заявка» в годной строке меняет отметку на «таблица»; сама по себе смена отметки
    обновлением не считается (строка «без изменений»), правка ФИО — считается."""
    setup()
    boss = admin()
    same = sd.make_requester(addr("same"), full_name="Иванов Иван", origin="ticket")
    renamed = sd.make_requester(addr("ren"), full_name="Старов Иван", origin="ticket")
    item = preview(boss, line(same.email, "Иванов", "Иван"), line(renamed.email, "Новов", "Иван"))
    rows = numbered(item)
    assert rows[1]["result"] == "unchanged" and rows[1]["origin_change"] is True
    assert rows[2]["result"] == "will_update"
    assert item.summary["origin_changed"] == 0, "«заявка» → «таблица» в origin_changed не входит"
    done = confirm(item, boss)
    same.refresh_from_db()
    renamed.refresh_from_db()
    assert same.origin == "table" and renamed.origin == "table"
    assert renamed.full_name == "Новов Иван"
    assert (done.summary["updated"], done.summary["unchanged"], done.summary["added"]) == (1, 1, 0)
    assert done.summary["origin_changed"] == 0
    assert numbered(done)[1]["result"] == "unchanged"
    changed = [e for e in events("requester.updated") if e.object_id == str(same.pk)]
    assert len(changed) == 1
    assert changed[0].before == {"origin": "ticket"} and changed[0].after == {"origin": "table"}


def test_preview_and_total_show_deleted_and_origin_changed_counts():
    """FR-263, FR-264: предпросмотр показывает, сколько удалит флажок (`[data-would-delete]`); итог — «удалено D» и
    «сменили отметку на «заявка»: E», а также атрибуты `data-deleted`, `data-origin-changed`."""
    setup()
    boss = admin()
    for tag in ("d1", "d2"):
        sd.make_requester(addr(tag), full_name=f"Удаляемый {tag}", origin="table")
    with_ticket(sd.make_requester(addr("tk"), full_name="С Заявкой", origin="table"))
    ticket_origin = sd.make_requester(addr("to"), full_name="Иванов Иван", origin="ticket")
    ticket_origin_2 = sd.make_requester(addr("to2"), full_name="Петров Пётр", origin="ticket")
    item = preview(boss, line(addr("new"), "Новый"), line(ticket_origin.email, "Иванов", "Иван"),
                   line(ticket_origin_2.email, "Петров", "Пётр"))
    assert item.summary["would_delete"] == 2
    assert item.summary["origin_changed"] == 1, "в предпросмотре — «таблица» с заявкой, адреса нет в файле (флажок будет)"
    assert item.summary["deleted"] == 0
    _, root = page(boss, item)
    assert root.find_all(data_import_preview=True)
    shown = root.find_all(data_would_delete=True)
    assert shown and "2" in digits_of(shown[0], "data-would-delete")
    confirm_form = root.find_all("form", data_form="import-confirm")[0]
    flags = [i for i in confirm_form.find_all("input", name="delete_missing") if i.attrs.get("type") == "checkbox"]
    assert len(flags) == 1 and "checked" not in flags[0].attrs
    assert not [i for i in confirm_form.find_all("input", name="confirm")]
    assert confirm_form.find_all("button")
    done = confirm(item, boss, delete_missing=True)
    assert done.summary["deleted"] == 2
    _, root = page(boss, done)
    summary = root.find_all(data_result_summary=True)[0]
    assert summary.attrs["data-deleted"] == "2"
    assert summary.attrs["data-origin-changed"] == "1"
    assert done.summary["origin_changed"] == 1
    text = summary.text()
    assert "добавлено 1, обновлено 0, без изменений 2, с ошибками 0" in text
    assert "удалено 2" in text
    assert "сменили отметку на «заявка»: 1" in text


# --- журнал -------------------------------------------------------------------------------------

def test_audit_requester_events_by_record_number_without_address_and_name():
    """FR-263: события `requester.created`, `requester.updated`, `requester.deleted` — по номеру записи, от имени
    администратора; ни адреса, ни ФИО нет ни в одном событии загрузки."""
    setup()
    boss = admin()
    old = sd.make_requester(addr("upd"), full_name="Стариков Семён", origin="table")
    gone = sd.make_requester(addr("del"), full_name="Удалёнов Дмитрий", origin="table")
    old_pk, gone_pk, gone_email = old.pk, gone.pk, gone.email
    new_email = addr("new")
    base = max([e.id for e in events()] or [0])
    item = preview(boss, line(new_email, "Зайцев", "Антон", unit="Отдел Х"), line(old.email, "Волков", "Семён"))
    assert import_events() == 0, "предпросмотр журнал не пишет"
    confirm(item, boss, delete_missing=True)
    created = requester(new_email)
    made = events("requester.created")
    assert [e.object_id for e in made] == [str(created.pk)]
    assert made[0].object_type == "requester" and made[0].actor_id == boss.pk
    assert made[0].after["origin"] == "table" and made[0].after["client_unit"] == "Отдел Х"
    updated = [e for e in events("requester.updated")]
    assert [e.object_id for e in updated] == [str(old_pk)]
    assert updated[0].before == {} and updated[0].after == {"name_changed": True}
    removed = events("requester.deleted")
    assert [e.object_id for e in removed] == [str(gone_pk)]
    fresh = [e for e in events() if e.id > base]
    flat = " ".join(json.dumps([e.before, e.after, e.object_id], ensure_ascii=False) for e in fresh).lower()
    for secret in (new_email, old.email, gone_email, "зайцев", "антон", "волков", "стариков", "удалёнов"):
        assert secret.lower() not in flat, secret


def test_audit_import_completed_who_and_totals():
    """FR-263: в конце — одно событие `requester_import.completed`: кто, имя файла, флажок и итоги; предпросмотр и
    повторное подтверждение его не пишут."""
    setup()
    boss = admin()
    sd.make_requester(addr("gone"), full_name="Уйдёт Иван", origin="table")
    known = sd.make_requester(addr("old"), full_name="Третьев", origin="table")
    with_ticket(sd.make_requester(addr("tk"), full_name="С Заявкой", origin="table"))
    item = preview(boss, line(addr("a"), "Первов"), line(addr("b"), "Второв"), line(known.email, "Третьев"),
                   line(addr("x", "other.example"), "Чужой"), name="people.csv")
    assert events("requester_import.completed") == []
    done = confirm(item, boss, delete_missing=True)
    confirm(done, boss, delete_missing=True)
    found = events("requester_import.completed")
    assert len(found) == 1
    event = found[0]
    assert event.object_type == "requester_import" and event.object_id == str(item.pk)
    assert event.actor_id == boss.pk
    assert event.after["file_name"] == "people.csv" and event.after["delete_missing"] is True
    totals = event.after["totals"]
    assert set(totals) == {"added", "updated", "unchanged", "errors", "deleted", "origin_changed"}
    assert {k: totals[k] for k in ("added", "updated", "unchanged", "errors", "deleted")} == {
        "added": 2, "updated": 0, "unchanged": 1, "errors": 1, "deleted": 1}
    assert totals["origin_changed"] == 1
    assert done.summary["origin_changed"] == 1


def test_audit_client_unit_created_for_created_units():
    """FR-263: на каждое созданное загрузкой подразделение — событие `client_unit.created` (вид «справочник», номер
    `client-units:<id>`, название, администратор); на существующее — нет; предпросмотр событий не пишет."""
    setup()
    boss = admin()
    unit("Склад")
    ClientUnit = model("servicedesk.ClientUnit")
    item = preview(boss, line(addr("a"), "Один", unit="Отдел А"), line(addr("b"), "Два", unit="Отдел Б"),
                   line(addr("c"), "Три", unit="Склад"), line(addr("d"), "Четыре", unit="Отдел А"))
    assert events(prefix="client_unit") == []
    confirm(item, boss)
    made = events("client_unit.created")
    assert len(made) == 2
    for event in made:
        assert event.object_type == "sd_dictionary" and event.actor_id == boss.pk
        assert event.object_id.startswith("client-units:")
    by_name = {e.after["name"]: e for e in made}
    assert set(by_name) == {"Отдел А", "Отдел Б"}
    for name, event in by_name.items():
        assert event.object_id == f"client-units:{ClientUnit.objects.get(name=name).pk}"


# --- файл и сроки хранения ----------------------------------------------------------------------

def test_file_not_stored_after_parse(protected_root):
    """FR-265: файл после разбора не хранится: у строки загрузки нет полей файла, в томе `protected` (временном) ничего
    не появилось, строк `StoredFile` не прибавилось — ни после загрузки через страницу, ни после подтверждения."""
    from django.db import models

    setup()
    boss = admin()
    fields = model("servicedesk.RequesterImport")._meta.get_fields()
    assert not [f for f in fields if isinstance(f, (models.FileField, models.BinaryField))]
    StoredFile = model("core.StoredFile")
    tree, stored = files_tree(), StoredFile.objects.count()
    upload = SimpleUploadedFile("people.csv", csv_of(line(addr("a"), "Первов")), content_type="text/csv")
    response = support.logged_in(boss).post(IMPORT, {"file": upload})
    assert response.status_code == 302
    item = model("servicedesk.RequesterImport").objects.get()
    assert response["Location"] == f"{IMPORT}{item.pk}/"
    assert item.file_name == "people.csv" and item.file_format == "csv"
    assert (files_tree(), StoredFile.objects.count()) == (tree, stored)
    confirm(item, boss)
    assert (files_tree(), StoredFile.objects.count()) == (tree, stored)


def test_unconfirmed_preview_purged_after_a_day():
    """FR-265: неподтверждённый предпросмотр старше суток стирается `purge_requester_previews`, моложе суток и
    выполненная загрузка остаются; ссылка на стёртый предпросмотр — 404."""
    setup()
    boss = admin()
    Import = model("servicedesk.RequesterImport")
    old = preview(boss, line(addr("a"), "Старов"))
    young = preview(boss, line(addr("b"), "Молодов"))
    done = confirm(preview(boss, line(addr("c"), "Готов")), boss)
    sd.backdate(old, "created_at", hours=25)
    sd.backdate(young, "created_at", hours=23)
    sd.backdate(done, "created_at", days=5)
    assert retention().purge_requester_previews() == 1
    assert not Import.objects.filter(pk=old.pk).exists()
    assert Import.objects.filter(pk=young.pk).exists() and Import.objects.filter(pk=done.pk).exists()
    client = support.logged_in(boss)
    assert client.get(f"{IMPORT}{old.pk}/").status_code == 404
    assert client.post(f"{IMPORT}{old.pk}/", {}).status_code == 404
    assert client.get(f"{IMPORT}{done.pk}/").status_code == 200


def test_executed_row_keeps_number_total_and_reason_only():
    """FR-265: после выполнения у строки остаются номер, итог и причина (без адреса и ФИО); итог загрузки сохранён;
    на странице — `[data-result-summary]` и строки с `data-row-result` и причиной только у ошибок."""
    setup()
    boss = admin()
    mail = addr("a")
    foreign = addr("x", "other.example")
    item = preview(boss, line(mail, "Первов", "Иван"), line(foreign, "Чужов", "Пётр"))
    done = reload(confirm(item, boss))
    assert done.state == "done"
    for row in done.rows:
        assert set(row) == {"number", "result", "reason"}, row
    rows = numbered(done)
    assert rows[1]["result"] == "added" and rows[2]["result"] == "error"
    assert rows[2]["reason"].startswith("ошибка")
    dumped = json.dumps([done.rows, done.summary, done.new_units], ensure_ascii=False).lower()
    for secret in (mail, foreign, "первов", "чужов"):
        assert secret.lower() not in dumped, secret
    for key in ("added", "updated", "unchanged", "errors", "deleted"):
        assert key in done.summary
    response, root = page(boss, done)
    summary = root.find_all(data_result_summary=True)[0]
    assert (summary.attrs["data-added"], summary.attrs["data-errors"]) == ("1", "1")
    shown = {r.attrs["data-row-number"]: r for r in root.find_all("tr", data_import_row=True)}
    assert {n: r.attrs["data-row-result"] for n, r in shown.items()} == {"1": "added", "2": "error"}
    assert not shown["1"].find_all(data_row_reason=True)
    assert shown["2"].find_all(data_row_reason=True)[0].text().startswith("ошибка")
    page_text = response.content.decode("utf-8").lower()
    assert mail.lower() not in page_text and foreign.lower() not in page_text
