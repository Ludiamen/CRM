"""T-6119: загрузка из таблицы — подтверждение, итог, журнал, образец (FR-514, FR-522, FR-533, FR-534, FR-535, FR-542).
Сценарий 9 (п. 9, 11-13) и сценарий 11 (п. 8) спеки среза 006.

Задания загрузки собираются службами `imports.upload` и `imports.preview` (CONTRACT-006, 5.7), подтверждение —
`imports.confirm` и адресом `/lists/<id>/import/<job>/confirm/` (10.2). Порядок замков — по SQL (5.2, решение 6 задач),
время — часы `clock` (решение 5), тексты — литералы раздела 12. Код среза на уровне модуля не импортируется.
"""
import importlib
import io
import json
import re
from datetime import datetime
from html.parser import HTMLParser
from urllib.parse import unquote

import pytest
from django.contrib.postgres.search import SearchQuery
from django.http import Http404

from tests import support
from tests import support_lists as ls

pytestmark = pytest.mark.django_db

FILE = "перечень-оборудования.xlsx"
HEADER = ["Название", "Серийный номер", "Кабинет"]
MOMENT = ls.moscow(2026, 3, 11, 0, 30)  # 10.03.2026 21:30 UTC: дата по Москве уже следующая


# --- разбор разметки (помощник только этого файла) ------------------------------------------------------

_VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "source", "track", "wbr"}


class Node:
    def __init__(self, tag, attrs, parent):
        self.tag = tag
        self.attrs = {name: (value if value is not None else "") for name, value in attrs}
        self.parent = parent
        self.children = []

    def walk(self):
        for child in self.children:
            if isinstance(child, Node):
                yield child
                yield from child.walk()

    def text(self) -> str:
        parts = []

        def collect(node):
            for child in node.children:
                if isinstance(child, Node):
                    collect(child)
                else:
                    parts.append(child)

        collect(self)
        return " ".join("".join(parts).split())


class _Tree(HTMLParser):
    def __init__(self):
        super().__init__()
        self.root = Node("#root", [], None)
        self.cur = self.root

    def handle_starttag(self, tag, attrs):
        node = Node(tag, attrs, self.cur)
        self.cur.children.append(node)
        if tag not in _VOID:
            self.cur = node

    def handle_startendtag(self, tag, attrs):
        self.cur.children.append(Node(tag, attrs, self.cur))

    def handle_endtag(self, tag):
        node = self.cur
        while node is not None and node.tag != tag:
            node = node.parent
        if node is not None and node.parent is not None:
            self.cur = node.parent

    def handle_data(self, data):
        self.cur.children.append(data)


def tree(response) -> Node:
    parser = _Tree()
    parser.feed(response.content.decode("utf-8"))
    return parser.root


def q(scope, tag=None, **attrs) -> list:
    """Потомки `scope` по тегу и атрибутам; имя с дефисом пишется через `_`; `True` — атрибут есть."""
    wanted = {name.replace("_", "-"): value for name, value in attrs.items()}
    found = []
    for node in scope.walk():
        if tag is not None and node.tag != tag:
            continue
        if all(name in node.attrs and (value is True or node.attrs[name] == value) for name, value in wanted.items()):
            found.append(node)
    return found


# --- помощники файла ------------------------------------------------------------------------------------

def _imports():
    return importlib.import_module("apps.lists.services.imports")


def _world():
    """Список с уникальной строкой «Серийный номер» и строкой «Кабинет»; правящий сотрудник."""
    lst = ls.make_list()
    uniq = ls.add_field(lst, "string", "Серийный номер", unique=True)
    room = ls.add_field(lst, "string", "Кабинет")
    return lst, uniq, room, ls.as_editor(lst)


def _mapping(uniq, room):
    """Сопоставление «номер столбца с 0 — код поля или title» (5.7, 10.2)."""
    return {"0": "title", "1": uniq.code, "2": room.code}


def _job(lst, actor, header, rows, mapping, name=FILE):
    """Файл → задание → предпросмотр (состояние `preview`); возвращает свежее задание."""
    imp = _imports()
    job = imp.upload(lst, actor, name=name, content=ls.xlsx_of(header, rows))
    return imp.preview(job, mapping, actor)


def _confirm_url(lst, job) -> str:
    return f"/lists/{lst.pk}/import/{job.pk}/confirm/"


def _records(lst):
    return support.model("lists.Record").objects.filter(list=lst)


def _titles(lst) -> set:
    return set(_records(lst).values_list("title", flat=True))


def _changes(record, kind=None):
    rows = support.model("lists.Change").objects.filter(record=record)
    return rows.filter(kind=kind) if kind else rows


def _all_changes(lst):
    return support.model("lists.Change").objects.filter(record__list=lst)


def _unique_values(lst):
    return support.model("lists.UniqueValue").objects.filter(list=lst)


def _audit(action: str, job):
    return list(support.model("core.AuditEvent").objects.filter(action=action, object_id=str(job.pk)))


def _first(sqls, predicate):
    for number, sql in enumerate(sqls):
        if predicate(sql):
            return number
    return None


def _for_update_at(sqls, table: str):
    """Номер первого `SELECT … FROM "<table>" … FOR UPDATE` или `None`."""
    head = f'FROM "{table}"'.upper()
    return _first(sqls, lambda sql: sql.lstrip().upper().startswith("SELECT")
                  and "FOR UPDATE" in sql.upper() and head in sql.upper())


def _digits(node) -> str:
    """Единственное число в тексте элемента (форма вокруг числа контрактом не задана)."""
    found = re.findall(r"\d+", node.text())
    assert len(found) == 1, node.text()
    return found[0]


def _has_attr(node, name: str) -> bool:
    return name in node.attrs or any(name in found.attrs for found in node.walk())


def _sheet_rows(sheet, width: int) -> list[list]:
    """Строки листа openpyxl; пустая ячейка (в том числе хвостовая) — `""`, строка дополнена до `width`."""
    out = []
    for row in sheet.iter_rows(values_only=True):
        cells = ["" if cell is None else cell for cell in row][:width]
        out.append(cells + [""] * (width - len(cells)))
    return out


def _book(content: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(content))


# --- 1. подтверждает только автор, не позже суток -------------------------------------------------------

def test_confirm_by_other_editor_than_preview_author_returns_404():
    """FR-534: подтверждение чужого задания (другой правящий) — 404 без данных; ничего не загружено; автор потом
    подтверждает без помех."""
    lst, uniq, room, author = _world()
    other = ls.as_editor(lst)
    job = _job(lst, author.user, HEADER, [["Принтер", "SN-1", "101"], ["Сканер", "SN-2", "102"]],
               _mapping(uniq, room))

    response = other.client.post(_confirm_url(lst, job))
    assert response.status_code == 404
    body = response.content.decode()
    assert "Принтер" not in body and "SN-1" not in body and FILE not in body
    with pytest.raises(Http404):
        _imports().confirm(lst, job.pk, other.user)
    assert _records(lst).count() == 0
    job.refresh_from_db()
    assert job.state == "preview" and len(job.rows) == 2

    response = author.client.post(_confirm_url(lst, job))
    assert response.status_code == 200
    assert _titles(lst) == {"Принтер", "Сканер"}


def test_confirm_job_older_than_one_day_is_404_and_nothing_loaded(clock):
    """FR-534: задание старше суток (`created_at < now - 1 день`) — 404, ничего не загружено, `rows` целы; задание
    моложе суток подтверждается."""
    clock.set(ls.moscow(2026, 3, 11, 12, 0))
    lst, uniq, room, author = _world()
    job_old = _job(lst, author.user, HEADER, [["Старое", "SN-1", "101"]], _mapping(uniq, room))
    clock.advance(hours=2)
    job_new = _job(lst, author.user, HEADER, [["Свежее", "SN-2", "102"]], _mapping(uniq, room))
    clock.advance(hours=23)  # возраст: старому 25 часов, новому 23
    client = support.logged_in(author.user)

    response = client.post(_confirm_url(lst, job_old))
    assert response.status_code == 404
    assert "Старое" not in response.content.decode()
    with pytest.raises(Http404):
        _imports().confirm(lst, job_old.pk, author.user)
    assert _records(lst).count() == 0
    job_old.refresh_from_db()
    assert job_old.state == "preview" and len(job_old.rows) == 1

    response = client.post(_confirm_url(lst, job_new))
    assert response.status_code == 200
    assert _titles(lst) == {"Свежее"}


# --- 2. замки, повтор -----------------------------------------------------------------------------------

def test_confirm_lock_order_shared_list_lock_then_job_then_list_row_for_update():
    """FR-534: порядок замков `imports.confirm` — разделяемый замок списка, `FOR UPDATE` задания, `FOR UPDATE` строки
    списка; схема читается после замка (раздел 5.2)."""
    lst, uniq, room, author = _world()
    job = _job(lst, author.user, HEADER, [["Принтер", "SN-1", "101"], ["Сканер", "SN-2", "102"]],
               _mapping(uniq, room))

    sqls = ls.sql_of(_imports().confirm, lst, job.pk, author.user)

    locks = ls.list_locks(sqls)
    assert locks and locks[0] == ls.ListLock("shared", ls.list_no(lst.pk)), locks
    assert all(lock.kind == "shared" for lock in locks), "подтверждение берёт только разделяемый замок списка"
    lock_at = _first(sqls, lambda sql: "pg_advisory_xact_lock_shared" in sql)
    job_at = _for_update_at(sqls, "lists_importjob")
    list_at = _for_update_at(sqls, "lists_list")
    assert job_at is not None, "нет SELECT … FOR UPDATE задания"
    assert list_at is not None, "нет SELECT … FOR UPDATE строки списка"
    assert lock_at < job_at <= list_at, (lock_at, job_at, list_at)
    field_at = _first(sqls, lambda sql: sql.lstrip().upper().startswith("SELECT") and 'FROM "LISTS_FIELD"' in sql.upper())
    assert field_at is not None and lock_at < field_at, "схема читается после замка списка"
    assert _records(lst).count() == 2


def test_confirm_twice_gives_single_summary_and_no_duplicates():
    """FR-534: повторное подтверждение (службой и адресом) даёт прежний итог, записей, истории, уникальных значений и
    журнала не прибавляется, загруженное не перезаписывается."""
    lst, uniq, room, author = _world()
    job = _job(lst, author.user, HEADER, [["Принтер", "SN-1", "101"], ["Сканер", "SN-2", "102"]],
               _mapping(uniq, room))
    imp = _imports()

    first = imp.confirm(lst, job.pk, author.user)
    assert first == {"created": 2, "updated": 0, "errors": 0}
    _records(lst).filter(title="Принтер").update(title="Принтер правлен")
    snapshot = (_records(lst).count(), _all_changes(lst).count(), _unique_values(lst).count())

    second = imp.confirm(lst, job.pk, author.user)
    assert second == first
    assert (_records(lst).count(), _all_changes(lst).count(), _unique_values(lst).count()) == snapshot
    assert _titles(lst) == {"Принтер правлен", "Сканер"}, "повтор не перезаписывает записи"
    assert len(_audit("lists.import_done", job)) == 1

    job2 = _job(lst, author.user, HEADER, [["Монитор", "SN-3", "103"]], _mapping(uniq, room))
    pages = []
    for _ in range(2):
        response = author.client.post(_confirm_url(lst, job2))
        assert response.status_code == 200
        notice = q(tree(response), None, data_lists_notice="import-summary")[0]
        pages.append(tuple(_digits(q(notice, None, data_import_summary=name)[0])
                           for name in ("created", "updated", "errors")))
    assert pages == [("1", "0", "0")] * 2
    assert _records(lst).filter(title="Монитор").count() == 1
    assert len(_audit("lists.import_done", job2)) == 1


# --- 3. строки: перепроверка, ошибки --------------------------------------------------------------------

def test_confirm_rebuilds_references_employee_dismissed_after_preview_is_row_error():
    """FR-534: справочники при подтверждении собираются заново — сотрудник, уволенный после предпросмотра, даёт ошибку
    строки; остальные строки загружены."""
    lst = ls.make_list()
    owner = ls.add_field(lst, "person", "Ответственный")
    author = ls.as_editor(lst)
    first = support.make_employee(email=ls.email_of("petrov-imp"))
    second = support.make_employee(email=ls.email_of("sidorov-imp"))
    job = _job(lst, author.user, ["Название", "Ответственный"],
               [["Сервер А", first.email], ["Сервер Б", second.email], ["Сервер В", first.email]],
               {"0": "title", "1": owner.code})
    assert _imports().counts(job)["will_create"] == 3
    support.model("core.User").objects.filter(pk=second.pk).update(state="dismissed")

    result = _imports().confirm(lst, job.pk, author.user)

    assert result == {"created": 2, "updated": 0, "errors": 1}
    assert _titles(lst) == {"Сервер А", "Сервер В"}
    assert {record.values[owner.code][0] for record in _records(lst)} == {first.pk}


def test_confirm_row_error_does_not_roll_back_other_rows():
    """FR-534: ошибка строки не откатывает остальные — строки до и после неё загружены, итог считает ошибку; строка
    с пустой хвостовой ячейкой загружена без значения."""
    lst = ls.make_list()
    when = ls.add_field(lst, "date", "Дата ввода")
    author = ls.as_editor(lst)
    job = _job(lst, author.user, ["Название", "Дата ввода"],
               [["А", "01.02.2026"], ["Б", "не дата"], ["В", "03.02.2026"], ["Г", None]],
               {"0": "title", "1": when.code})
    assert _imports().counts(job) == {"will_create": 3, "will_update": 0, "error": 1}

    result = _imports().confirm(lst, job.pk, author.user)

    assert result == {"created": 3, "updated": 0, "errors": 1}
    assert _titles(lst) == {"А", "В", "Г"}
    by_title = {record.title: record for record in _records(lst)}
    assert by_title["А"].values[when.code] == "2026-02-01"
    assert by_title["В"].values[when.code] == "2026-02-03"
    assert when.code not in by_title["Г"].values


def test_confirm_unique_constraint_violation_becomes_row_error_others_loaded(monkeypatch):
    """FR-534: предварительная проверка обойдена подменой `unique.find_owners` (решение 28) — ограничение базы даёт
    ошибку строки; строки до и после загружены, владелец значения не тронут, дубля нет."""
    lst, uniq, room, author = _world()
    job = _job(lst, author.user, HEADER,
               [["Один", "SN-0", "1"], ["Два", "SN-1", "2"], ["Три", "SN-2", "3"]], _mapping(uniq, room))
    assert _imports().counts(job)["will_create"] == 3
    owner = ls.make_record(lst, "Владелец", {uniq.code: "SN-1"}, created_by=author.user)
    monkeypatch.setattr("apps.lists.services.unique.find_owners", lambda *args, **kwargs: {})

    result = _imports().confirm(lst, job.pk, author.user)

    assert result == {"created": 2, "updated": 0, "errors": 1}
    assert _titles(lst) == {"Владелец", "Один", "Три"}
    taken = _unique_values(lst).filter(field=uniq, norm_value="sn-1")
    assert [row.record_id for row in taken] == [owner.pk]
    assert _unique_values(lst).filter(field=uniq).count() == 3
    owner.refresh_from_db()
    assert owner.title == "Владелец" and owner.values[uniq.code] == "SN-1"


# --- 4. история -----------------------------------------------------------------------------------------

def test_confirm_new_records_get_history_created_by_import():
    """FR-534, FR-522: новая запись — событие `created` с `via_import` и автором-загрузившим; в истории — «Создана
    загрузкой»."""
    lst, uniq, room, author = _world()
    job = _job(lst, author.user, HEADER, [["Принтер", "SN-1", "101"], ["Сканер", "SN-2", "102"]],
               _mapping(uniq, room))

    _imports().confirm(lst, job.pk, author.user)

    for record in _records(lst):
        rows = list(_changes(record))
        assert [(row.kind, row.via_import, row.actor_id) for row in rows] == [("created", True, author.user.pk)]
        response = author.client.get(f"/lists/{lst.pk}/{record.number}/history/")
        assert response.status_code == 200
        items = q(tree(response), "li", data_change_kind="created")
        assert len(items) == 1
        assert _has_attr(items[0], "data-via-import")
        assert "Создана загрузкой" in items[0].text()


def test_confirm_updated_records_get_history_changed_by_import():
    """FR-534, FR-522: обновлённая запись — событие `imported` с изменённым полем «было → стало» (строки показа);
    события `created` нет; в истории — «Изменено загрузкой»."""
    lst, uniq, room, author = _world()
    record = ls.make_record(lst, "Старое", {uniq.code: "SN-9", room.code: "101"}, created_by=author.user)
    job = _job(lst, author.user, HEADER, [["Новое", "SN-9", "202"]], _mapping(uniq, room))
    assert _imports().counts(job)["will_update"] == 1

    result = _imports().confirm(lst, job.pk, author.user)

    assert result == {"created": 0, "updated": 1, "errors": 0}
    record.refresh_from_db()
    assert record.title == "Новое" and record.values[room.code] == "202"
    rows = list(_changes(record))
    assert [(row.kind, row.actor_id) for row in rows] == [("imported", author.user.pk)]
    entries = [entry for entry in rows[0].changes if entry.get("field") == room.code]
    assert len(entries) == 1
    assert (entries[0]["name"], entries[0]["before"], entries[0]["after"]) == ("Кабинет", "101", "202")
    response = author.client.get(f"/lists/{lst.pk}/{record.number}/history/")
    items = q(tree(response), "li", data_change_kind="imported")
    assert len(items) == 1 and "Изменено загрузкой" in items[0].text()
    change = q(items[0], None, data_change_field=room.code)
    assert len(change) == 1
    assert q(change[0], None, data_field="before")[0].text() == "101"
    assert q(change[0], None, data_field="after")[0].text() == "202"


def test_confirm_new_records_numbered_from_next_number():
    """FR-534: новые записи получают номера из `next_number` списка, счётчик растёт на число созданных."""
    lst, uniq, room, author = _world()
    support.model("lists.List").objects.filter(pk=lst.pk).update(next_number=5)
    lst.refresh_from_db()
    job = _job(lst, author.user, HEADER, [["Принтер", "SN-1", "101"], ["Сканер", "SN-2", "102"]],
               _mapping(uniq, room))

    _imports().confirm(lst, job.pk, author.user)

    assert sorted(_records(lst).values_list("number", flat=True)) == [5, 6]
    lst.refresh_from_db()
    assert lst.next_number == 7


def test_confirm_refreshes_vectors_once_in_batch_at_end():
    """FR-534: вектор поиска пересчитывается пакетно в конце — один запрос `unnest(` после всех вставок записей;
    у созданных и обновлённых записей вектор по новому названию."""
    lst, uniq, room, author = _world()
    ls.make_record(lst, "Старое-один", {uniq.code: "SN-1"}, created_by=author.user)
    ls.make_record(lst, "Старое-два", {uniq.code: "SN-2"}, created_by=author.user)
    job = _job(lst, author.user, HEADER,
               [["Alphacore", "SN-1", "1"], ["Betacore", "SN-2", "2"], ["Gammacore", "SN-3", "3"],
                ["Deltacore", "SN-4", "4"], ["Epsiloncore", "SN-5", "5"]], _mapping(uniq, room))
    assert _imports().counts(job) == {"will_create": 3, "will_update": 2, "error": 0}
    Record = support.model("lists.Record")
    assert not Record.objects.filter(search=SearchQuery("alphacore", config="russian")).exists()

    sqls = ls.sql_of(_imports().confirm, lst, job.pk, author.user)

    batches = [number for number, sql in enumerate(sqls)
               if "unnest(" in sql.lower() and "update" in sql.lower()]
    assert len(batches) == 1, f"пакетный пересчёт один, найдено {len(batches)}"
    inserts = [number for number, sql in enumerate(sqls)
               if sql.lstrip().upper().startswith('INSERT INTO "LISTS_RECORD"')]
    assert inserts and batches[0] > max(inserts), "пересчёт — в конце, после вставок записей"
    for word in ("alphacore", "betacore", "gammacore", "deltacore", "epsiloncore"):
        assert Record.objects.filter(list=lst, search=SearchQuery(word, config="russian")).count() == 1, word


def test_reupload_same_file_updates_same_records_creates_none():
    """FR-534: повторная загрузка того же файла при сопоставленном уникальном поле обновляет те же записи и не создаёт
    новых; события `created` не прибавляются."""
    lst, uniq, room, author = _world()
    rows = [["Принтер", "SN-1", "101"], ["Сканер", "SN-2", "102"]]
    first = _job(lst, author.user, HEADER, rows, _mapping(uniq, room))
    assert _imports().confirm(lst, first.pk, author.user) == {"created": 2, "updated": 0, "errors": 0}
    before = dict(_records(lst).values_list("title", "pk"))
    lst.refresh_from_db()
    next_number = lst.next_number

    again = _job(lst, author.user, HEADER, rows, _mapping(uniq, room))
    assert _imports().counts(again) == {"will_create": 0, "will_update": 2, "error": 0}
    result = _imports().confirm(lst, again.pk, author.user)

    assert result == {"created": 0, "updated": 2, "errors": 0}
    assert dict(_records(lst).values_list("title", "pk")) == before
    lst.refresh_from_db()
    assert lst.next_number == next_number
    assert _all_changes(lst).filter(kind="created").count() == 2


# --- 5. итог --------------------------------------------------------------------------------------------

def test_confirm_clears_rows_keeps_summary_and_sets_done():
    """FR-534: после подтверждения `rows` пусты, `summary` — три числа, `state="done"`; возвращённый итог совпадает."""
    lst, uniq, room, author = _world()
    ls.make_record(lst, "Старое", {uniq.code: "SN-1"}, created_by=author.user)
    job = _job(lst, author.user, HEADER,
               [["Новое", "SN-1", "1"], ["Сканер", "SN-2", "2"], [None, "SN-3", "3"]], _mapping(uniq, room))
    assert len(job.rows) == 3

    result = _imports().confirm(lst, job.pk, author.user)

    assert result == {"created": 1, "updated": 1, "errors": 1}
    job.refresh_from_db()
    assert job.state == "done"
    assert job.rows == []
    assert job.summary == {"created": 1, "updated": 1, "errors": 1}


def test_confirm_shows_summary_created_updated_with_errors_error_rows_not_loaded(clock):
    """FR-534: страница итога — «Создано: N, обновлено: M, с ошибками: K», числа по зацепкам, время завершения по
    Москве; строка с ошибкой не загружена."""
    clock.set(MOMENT)
    lst, uniq, room, author = _world()
    ls.make_record(lst, "Старое", {uniq.code: "SN-1"}, created_by=author.user)
    job = _job(lst, author.user, HEADER,
               [["Новое", "SN-1", "1"], ["Сканер", "SN-2", "2"], [None, "SN-3", "3"]], _mapping(uniq, room))

    response = author.client.post(_confirm_url(lst, job))

    assert response.status_code == 200
    root = tree(response)
    assert q(root, None, data_lists_page="summary")
    notice = q(root, None, data_lists_notice="import-summary")
    assert len(notice) == 1
    shown = {name: q(notice[0], None, data_import_summary=name) for name in ("created", "updated", "errors")}
    assert all(len(found) == 1 for found in shown.values())
    assert [_digits(shown[name][0]) for name in ("created", "updated", "errors")] == ["1", "1", "1"]
    assert "Создано: 1, обновлено: 1, с ошибками: 1" in notice[0].text()
    moments = q(notice[0], "time", data_at=True) or q(root, "time", data_at=True)
    assert len(moments) == 1
    assert moments[0].text() == "11.03.2026 00:30" == ls.moscow_text(MOMENT)
    assert datetime.fromisoformat(moments[0].attrs["datetime"]) == MOMENT
    assert _titles(lst) == {"Новое", "Сканер"}
    assert not _unique_values(lst).filter(norm_value="sn-3").exists()


def test_confirm_writes_audit_import_done_with_counts_and_file_name_without_values():
    """FR-542: итог загрузки — одна запись журнала `lists.import_done` (тип `lists.import`, id задания): список, файл,
    три числа; значений ячеек нет."""
    lst, uniq, room, author = _world()
    job = _job(lst, author.user, HEADER,
               [["Модем-777", "SN-UNIQ-777", "кабинет-777"], ["Сканер", "SN-2", "2"], [None, "SN-3", "3"]],
               _mapping(uniq, room), name=FILE)

    _imports().confirm(lst, job.pk, author.user)

    events = _audit("lists.import_done", job)
    assert len(events) == 1
    event = events[0]
    assert event.actor == author.user
    assert event.object_type == "lists.import"
    assert set(event.after) == {"list", "file", "created", "updated", "errors"}
    assert event.after["list"] == lst.name and event.after["file"] == FILE
    assert [str(event.after[name]) for name in ("created", "updated", "errors")] == ["2", "0", "1"]
    dump = json.dumps([event.before, event.after], ensure_ascii=False)
    for value in ("Модем-777", "SN-UNIQ-777", "кабинет-777"):
        assert value not in dump


def test_confirm_500_rows_in_one_confirmation():
    """FR-533, FR-534: 500 строк загружаются одним подтверждением — 500 записей с номерами подряд, историей и
    уникальными значениями."""
    lst, uniq, room, author = _world()
    rows = [[f"Узел {n}", f"SN-{n:04d}", str(n)] for n in range(500)]
    job = _job(lst, author.user, HEADER, rows, _mapping(uniq, room))
    assert _imports().counts(job)["will_create"] == 500

    result = _imports().confirm(lst, job.pk, author.user)

    assert result == {"created": 500, "updated": 0, "errors": 0}
    assert sorted(_records(lst).values_list("number", flat=True)) == list(range(1, 501))
    assert _all_changes(lst).filter(kind="created", via_import=True).count() == 500
    assert _unique_values(lst).filter(field=uniq).count() == 500
    lst.refresh_from_db()
    assert lst.next_number == 501


# --- 6. образец -----------------------------------------------------------------------------------------

def test_sample_for_editor_reader_403_non_reader_404():
    """FR-535, FR-514: образец — правящему 200 (XLSX, имя «<список> — образец.xlsx»), читающему 403, не читающему 404;
    управляющему 200."""
    lst, uniq, room, editor = _world()
    url = f"/lists/{lst.pk}/sample/"

    response = editor.client.get(url)
    assert response.status_code == 200
    assert "spreadsheetml" in response["Content-Type"]
    disposition = response["Content-Disposition"]
    assert disposition.startswith("attachment") and "filename*=UTF-8''" in disposition
    assert unquote(disposition.split("filename*=UTF-8''", 1)[1]) == f"{lst.name} — образец.xlsx"
    assert ls.read_xlsx(response.content)[0][0] == "Название"

    assert ls.as_reader(lst).client.get(url).status_code == 403
    assert support.logged_in(support.make_employee()).get(url).status_code == 404
    assert ls.as_manager().client.get(url).status_code == 200


def test_sample_first_sheet_title_and_visible_fields_except_file_and_hidden():
    """FR-535: первый лист — одна строка заголовков: «Название» и видимые поля по порядку, без полей «файл» и
    скрытых."""
    lst = ls.make_list()
    ls.add_field(lst, "string", "Серийный номер")
    ls.add_field(lst, "file", "Скан договора")
    ls.add_field(lst, "string", "Служебное", hidden=True)
    ls.add_field(lst, "number", "Количество")

    content = _imports().sample(lst)

    rows = ls.read_xlsx(content)
    assert [cell for cell in rows[0] if cell is not None] == ["Название", "Серийный номер", "Количество"]
    assert len(rows) == 1, "на первом листе только заголовки"
    assert _book(content).sheetnames[1] == "Подсказки"


def test_sample_hints_sheet_type_required_unique_choices_and_formats():
    """FR-535: лист «Подсказки» — по строке на заголовок первого листа: тип, обязательное, уникальное, значения без
    архивных, формат даты, даты и времени и денег."""
    lst = ls.make_list()
    ls.add_field(lst, "string", "Серийный номер", required=True, unique=True)
    state = ls.add_field(lst, "choice", "Состояние", choices=["Новый", "Списан", "Старый"])
    ls.add_field(lst, "file", "Скан договора")
    ls.add_field(lst, "date", "Дата покупки")
    ls.add_field(lst, "datetime", "Время ввода")
    ls.add_field(lst, "money", "Цена")
    ls.add_field(lst, "number", "Количество")
    ls.add_field(lst, "bool", "Включено")
    support.model("lists.Choice").objects.filter(pk=ls.choice_id(state, "Старый")).update(archived=True)

    book = _book(_imports().sample(lst))

    assert book.sheetnames[1] == "Подсказки"
    first = _sheet_rows(book.worksheets[0], 8)[0]
    hints = _sheet_rows(book.worksheets[1], 6)
    assert hints[0] == ["Поле", "Тип", "Обязательное", "Уникальное", "Значения", "Формат"]
    assert [row[0] for row in hints[1:]] == [cell for cell in first if cell != ""]
    assert hints[1:] == [
        ["Название", "Строка", "Да", "Нет", "", ""],
        ["Серийный номер", "Строка", "Да", "Да", "", ""],
        ["Состояние", "Список значений", "Нет", "Нет", "Новый; Списан", ""],
        ["Дата покупки", "Дата", "Нет", "Нет", "", "ДД.ММ.ГГГГ"],
        ["Время ввода", "Дата и время", "Нет", "Нет", "", "ДД.ММ.ГГГГ ЧЧ:ММ"],
        ["Цена", "Деньги", "Нет", "Нет", "", "1234,50"],
        ["Количество", "Число", "Нет", "Нет", "", ""],
        ["Включено", "Да/нет", "Нет", "Нет", "", ""],
    ]


def test_sample_headers_and_hints_escaped_against_formulas():
    """FR-535: заголовки и подсказки, начинающиеся с `=`, `+`, `-`, `@`, получают апостроф впереди и не становятся
    формулами — на обоих листах."""
    lst = ls.make_list()
    ls.add_field(lst, "string", "=SUM(A1)")
    ls.add_field(lst, "string", "+7 поле")
    ls.add_field(lst, "string", "-минус")
    kind = ls.add_field(lst, "choice", "@упом", choices=["=1+1", "ок"])
    content = _imports().sample(lst)

    book = _book(content)
    first = [cell for cell in book.worksheets[0][1]]
    assert [cell.value for cell in first] == ["Название", "'=SUM(A1)", "'+7 поле", "'-минус", "'@упом"]
    assert all(cell.data_type != "f" for sheet in book.worksheets for row in sheet.iter_rows() for cell in row)
    hints = _sheet_rows(book.worksheets[1], 6)
    assert [row[0] for row in hints[1:]] == ["Название", "'=SUM(A1)", "'+7 поле", "'-минус", "'@упом"]
    assert hints[-1][4] == "'=1+1; ок"
    assert ls.read_back(content).header == ["Название", "'=SUM(A1)", "'+7 поле", "'-минус", "'@упом"]
    assert kind.name == "@упом"


def test_upload_reads_only_first_sheet_of_sample():
    """FR-535, FR-533: загрузка читает только первый лист образца — лист «Подсказки» в строки не попадает; заполненный
    образец создаёт запись."""
    lst, uniq, room, editor = _world()
    book = _book(_imports().sample(lst))
    book.worksheets[0].append(["Из образца", "SN-S1", "305"])
    out = io.BytesIO()
    book.save(out)

    imp = _imports()
    job = imp.upload(lst, editor.user, name="образец.xlsx", content=out.getvalue())
    assert list(job.header) == HEADER
    assert len(job.rows) == 1
    dump = json.dumps(job.rows, ensure_ascii=False)
    for hint in ("Подсказки", "Обязательное", "Уникальное", "Строка"):
        assert hint not in dump

    job = imp.preview(job, _mapping(uniq, room), editor.user)
    assert job.rows[0]["cells"] == ["Из образца", "SN-S1", "305"]
    assert imp.counts(job) == {"will_create": 1, "will_update": 0, "error": 0}
    assert imp.confirm(lst, job.pk, editor.user) == {"created": 1, "updated": 0, "errors": 0}
    record = _records(lst).get()
    assert record.title == "Из образца" and record.values[uniq.code] == "SN-S1"
