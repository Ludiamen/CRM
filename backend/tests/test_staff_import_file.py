"""T-2110: приём и разбор файла сотрудников (FR-136…138, FR-141, NFR-108). Сценарии 5 (п. 9, 10, 11) спеки."""
import csv
import datetime
import io
import secrets
import zipfile
from html.parser import HTMLParser

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from tests import support, support_org
from tests.support import model

pytestmark = pytest.mark.django_db

MAX_BYTES = 1_048_576
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
VOID_TAGS = ("area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "source", "track", "wbr")
OWN_KEYS = {"email", "last_name", "first_name", "middle_name", "position", "phone", "internal_phone",
            "workplace", "unit_path", "head_raw"}


# --- помощники этого файла --------------------------------------------------------------------

def importer():
    """Администратор с обоими правами, нужными для загрузки (FR-136)."""
    return support.make_with_permissions("accounts.manage", "org.manage")


def service():
    from apps.org.services import staff_import
    return staff_import


def org_error():
    from apps.org.services import OrgError
    return OrgError


def upload(client, content: bytes, name: str = "people.csv", content_type: str = "application/octet-stream"):
    return client.post("/structure/import/", {"file": SimpleUploadedFile(name, content, content_type=content_type)})


class _Page(HTMLParser):
    """Теги страницы: (тег, атрибуты); тексты `li` внутри блока с данным атрибутом данных."""

    def __init__(self, block: str):
        super().__init__()
        self.block, self.tags, self.stack, self.items, self._open = block, [], [], [], None
        self._block_depth = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        self.tags.append((tag, attrs))
        if tag in VOID_TAGS:
            return
        self.stack.append(tag)
        if self.block in attrs and self._block_depth is None:
            self._block_depth = len(self.stack)
        if tag == "li" and self._block_depth is not None:
            self._open = []

    def handle_endtag(self, tag):
        if tag == "li" and self._open is not None:
            self.items.append("".join(self._open).strip())
            self._open = None
        if self.stack and self.stack[-1] == tag:
            if self._block_depth == len(self.stack):
                self._block_depth = None
            self.stack.pop()

    def handle_data(self, data):
        if self._open is not None:
            self._open.append(data)


def page(html: str, block: str = "data-import-not-loaded") -> _Page:
    parser = _Page(block)
    parser.feed(html)
    return parser


def has(parsed: _Page, tag: str, **want) -> bool:
    """Есть тег с атрибутами (`data_form` читается как `data-form`)."""
    return any(t == tag and all(a.get(k.replace("_", "-")) == v for k, v in want.items()) for t, a in parsed.tags)


def alert_shown(response) -> bool:
    return response.status_code == 200 and any(a.get("role") == "alert" for _, a in page(response.content.decode()).tags)


def tail_after(text: str, marker: str) -> str:
    return text[text.index(marker):]


def rows_of(count: int, **over) -> list[dict]:
    return [support_org.staff_row(**over) for _ in range(count)]


def xlsx_with_cached_formula(cached: str) -> bytes:
    """Книга, где в B2 формула `=1+1` с сохранённым значением `cached` (как после пересчёта в Excel)."""
    import re
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.append(list(support_org.STAFF_HEADERS))
    sheet.append([support_org.staff_row()["Почта"], "=1+1", "Имя", "Должность"])
    raw = io.BytesIO()
    book.save(raw)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                text, count = re.subn(r'<c r="B2"[^>]*>.*?</c>',
                                      f'<c r="B2" t="str"><f>1+1</f><v>{cached}</v></c>', text, flags=re.S)
                assert count == 1, "помощник теста не нашёл ячейку B2"
                data = text.encode("utf-8")
            target.writestr(info.filename, data)
    return out.getvalue()


def exact_size_csv(total: int) -> bytes:
    """Годный CSV ровно из `total` байт: 1000 строк, лишние байты — в значении «Должность» строк."""
    rows = rows_of(1000, position="")
    base = len(support_org.staff_csv(rows))
    each, rest = divmod(total - base, 1000)
    for i, row in enumerate(rows):
        row["Должность"] = "x" * (each + (rest if i == 999 else 0))
    data = support_org.staff_csv(rows)
    assert len(data) == total, "помощник теста собрал не тот размер"
    return data


def nothing_created(before_users: int):
    assert model("org.StaffImport").objects.count() == 0, "отказ: запись загрузки не создаётся"
    assert model("core.User").objects.count() == before_users, "отказ: сотрудники не создаются"


def refused(content: bytes, code: str, *, name: str = "people.csv", must_mention: tuple = ()):
    """Служба отказывает кодом `code`, страница — 200 с `[role=alert]`, ничего не создано."""
    with pytest.raises(org_error()) as raised:
        service().parse_file(content)
    assert raised.value.code == code, f"ожидали {code!r}, получили {raised.value.code!r}"
    for word in must_mention:
        assert word in str(raised.value), f"текст отказа должен называть {word!r}"
    admin = importer()
    users = model("core.User").objects.count()
    response = upload(support.logged_in(admin), content, name)
    assert alert_shown(response), f"{code}: ожидали 200 с [role=alert], получили {response.status_code}"
    nothing_created(users)


# --- тесты --------------------------------------------------------------------------------------

def test_import_and_samples_need_both_permissions():
    """FR-136: загрузка и образцы требуют оба права, `accounts.manage` и `org.manage`; кнопка в списке — тоже."""
    only_accounts = support.make_with_permissions("accounts.manage", "staff.view")
    only_org = support.make_with_permissions("org.manage", "staff.view")
    nobody = support.make_with_permissions("staff.view")
    both = importer()
    paths = ("/structure/import/", "/structure/import/sample.csv", "/structure/import/sample.xlsx")
    anonymous = support.make_client()
    for path in paths:
        assert anonymous.get(path).status_code == 302 and "/login/" in anonymous.get(path)["Location"], path
    for user in (only_accounts, only_org, nobody):
        client = support.logged_in(user)
        for path in paths:
            assert client.get(path).status_code == 403, f"{path}: одного права мало"
        before = model("core.User").objects.count()
        assert upload(client, support_org.staff_csv([support_org.staff_row()])).status_code == 403
        assert model("org.StaffImport").objects.count() == 0 and model("core.User").objects.count() == before
    client = support.logged_in(both)
    form = client.get("/structure/import/")
    assert form.status_code == 200
    html = form.content.decode()
    parsed = page(html)
    assert has(parsed, "form", data_form="import-upload")
    assert any(t == "input" and a.get("type") == "file" and a.get("name") == "file" for t, a in parsed.tags)
    sample_csv = client.get("/structure/import/sample.csv")
    assert sample_csv.status_code == 200
    assert sample_csv["Content-Type"] == "text/csv; charset=utf-8"
    assert sample_csv["Content-Disposition"].startswith("attachment")
    sample_xlsx = client.get("/structure/import/sample.xlsx")
    assert sample_xlsx.status_code == 200
    assert sample_xlsx["Content-Type"] == XLSX_TYPE
    assert sample_xlsx["Content-Disposition"].startswith("attachment")
    # кнопка «Загрузить из таблицы» в списке — только при обоих правах
    for user, shown in ((both, True), (only_accounts, False), (only_org, False)):
        listing = support.logged_in(user).get("/staff/")
        assert listing.status_code == 200
        buttons = [a for t, a in page(listing.content.decode()).tags
                   if t == "a" and a.get("data-action") == "staff-import"]
        assert bool(buttons) is shown, f"кнопка загрузки: ожидали {shown}"
        if shown:
            assert buttons[0]["href"] == "/structure/import/"
            assert "Загрузить из таблицы" in listing.content.decode()


def test_csv_semicolon_comma_bom_and_xlsx_accepted():
    """FR-136: CSV с «;» и «,», UTF-8 с BOM и без, и XLSX принимаются; формат определяется по содержимому."""
    rows = [support_org.staff_row(last_name="Иванов", first_name="Пётр", unit_path="ИТ-отдел / Сеть", head="да"),
            support_org.staff_row(last_name="Петрова", position="Аналитик", head="нет")]
    variants = {
        "csv;bom": (support_org.staff_csv(rows, delimiter=";", bom=True), "csv"),
        "csv;nobom": (support_org.staff_csv(rows, delimiter=";", bom=False), "csv"),
        "csv,bom": (support_org.staff_csv(rows, delimiter=",", bom=True), "csv"),
        "csv,nobom": (support_org.staff_csv(rows, delimiter=",", bom=False), "csv"),
        "xlsx": (support_org.staff_xlsx(rows), "xlsx"),
    }
    for label, (content, expected_format) in variants.items():
        parsed = service().parse_file(content)
        assert parsed.file_format == expected_format, label
        assert len(parsed.rows) == 2, label
        for source, got in zip(rows, parsed.rows):
            assert set(got) == OWN_KEYS, f"{label}: ключи строки {sorted(got)}"
            assert got["email"] == source["Почта"], label
            assert (got["last_name"], got["first_name"], got["middle_name"], got["position"]) == (
                source["Фамилия"], source["Имя"], source["Отчество"], source["Должность"]), label
            assert (got["phone"], got["internal_phone"], got["workplace"], got["unit_path"], got["head_raw"]) == (
                source["Телефон"], source["Внутренний телефон"], source["Место работы"],
                source["Подразделение"], source["Руководитель"]), label
        assert parsed.not_loaded == [], label
    # через страницу: файл принят, запись загрузки создана с форматом из содержимого
    admin = importer()
    client = support.logged_in(admin)
    for label, (content, expected_format) in variants.items():
        name = "people.xlsx" if expected_format == "xlsx" else "people.csv"
        response = upload(client, content, name)
        assert response.status_code == 302, f"{label}: {response.status_code}"
        item = model("org.StaffImport").objects.order_by("-id").first()
        assert response["Location"].endswith(f"/structure/import/{item.pk}/"), label
        assert item.file_format == expected_format and item.state == "preview" and item.file_name == name, label
        assert len(item.rows) == 2 and item.rows[0]["email"] == rows[0]["Почта"].lower(), label


def test_samples_csv_and_xlsx_headers_and_two_fake_rows():
    """FR-136: образцы — заголовки колонок по порядку и две вымышленные строки; CSV — UTF-8 с BOM и «;»."""
    staff_import = service()
    headers = list(staff_import.REQUIRED_COLUMNS + staff_import.OPTIONAL_COLUMNS)
    assert headers == ["Почта", "Фамилия", "Имя", "Должность", "Отчество", "Телефон", "Внутренний телефон",
                       "Место работы", "Подразделение", "Руководитель"]
    raw = staff_import.sample_csv()
    assert raw.startswith(b"\xef\xbb\xbf"), "образец CSV — UTF-8 с BOM"
    lines = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
    assert lines[0] == headers
    assert len(lines) == 3, "заголовок и две строки"
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(staff_import.sample_xlsx()), data_only=True).worksheets[0]
    table = [["" if cell is None else str(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    table = [row for row in table if any(row)]
    assert table[0] == headers and len(table) == 3, "XLSX: заголовок и две строки"
    for name, data in (("csv", lines[1:]), ("xlsx", table[1:])):
        for row in data:
            assert all(row[:4]), f"{name}: обязательные поля вымышленных строк заполнены"
            assert "@" in row[0]
        assert data[0][0].lower() != data[1][0].lower(), f"{name}: почты строк разные"
    # образцы разбираются тем же разбором, строки годятся для загрузки
    for content, fmt in ((raw, "csv"), (staff_import.sample_xlsx(), "xlsx")):
        parsed = staff_import.parse_file(content)
        assert parsed.file_format == fmt and len(parsed.rows) == 2 and parsed.not_loaded == []
    # страницы отдают те же файлы
    client = support.logged_in(importer())
    assert client.get("/structure/import/sample.csv").content == raw
    from_page = load_workbook(io.BytesIO(client.get("/structure/import/sample.xlsx").content)).worksheets[0]
    assert [c.value for c in next(from_page.iter_rows())] == headers


def test_headers_case_space_and_yo_insensitive_required_columns():
    """FR-137: заголовки без учёта регистра, пробелов и «е/ё»; порядок колонок любой; нужны Почта, Фамилия, Имя, Должность."""
    row = support_org.staff_row(last_name="Сидоров", first_name="Фёдор", position="Инженер", phone="+7 900 111-22-33",
                                workplace="Кабинет 2", internal_phone="202", unit_path="ИТ-отдел", head="да")
    messy = ["  ПОЧТА ", "фамилия", "И м Я", "ДОЛЖНОСТЬ", "Отчество", "Тёлефон", "внутреннийтелефон",
             "Место  работы", "Подразделёние", "РУКОВОДИТЕЛЬ"]
    plain = list(support_org.STAFF_HEADERS)
    values = [row[h] for h in plain]
    for build in (lambda: support_org.staff_csv([values], headers=messy),
                  lambda: support_org.staff_csv([values], headers=messy, delimiter=","),
                  lambda: support_org.staff_xlsx([values], headers=messy)):
        got = service().parse_file(build()).rows[0]
        assert (got["email"], got["last_name"], got["first_name"], got["position"]) == (
            row["Почта"], "Сидоров", "Фёдор", "Инженер")
        assert (got["middle_name"], got["phone"], got["internal_phone"], got["workplace"], got["unit_path"],
                got["head_raw"]) == (row["Отчество"], "+7 900 111-22-33", "202", "Кабинет 2", "ИТ-отдел", "да")
    # другой порядок колонок
    reverse = plain[::-1]
    got = service().parse_file(support_org.staff_csv([[row[h] for h in reverse]], headers=reverse)).rows[0]
    assert (got["email"], got["last_name"], got["position"], got["head_raw"]) == (
        row["Почта"], "Сидоров", "Инженер", "да")
    # только обязательные колонки — достаточно
    required = list(service().REQUIRED_COLUMNS)
    only = service().parse_file(support_org.staff_csv([[row[h] for h in required]], headers=required))
    assert only.rows[0]["email"] == row["Почта"] and only.not_loaded == []
    # без любой из обязательных колонок отказ получает весь файл
    for missing in required:
        kept = [h for h in plain if h != missing]
        refused(support_org.staff_csv([[row[h] for h in kept]], headers=kept), "missing_columns",
                must_mention=(missing,))


def test_extra_columns_listed_as_not_loaded():
    """FR-138: лишние колонки не читаются, а предпросмотр называет их «не загружаются»."""
    token = secrets.token_hex(6)
    headers = list(support_org.STAFF_HEADERS) + ["Роль", "Дата рождения"]
    row = support_org.staff_row()
    values = [row[h] for h in support_org.STAFF_HEADERS] + [f"роль-{token}", f"дата-{token}"]
    for content in (support_org.staff_csv([values], headers=headers), support_org.staff_xlsx([values], headers=headers)):
        parsed = service().parse_file(content)
        assert parsed.not_loaded == ["Роль", "Дата рождения"]
        assert len(parsed.rows) == 1 and set(parsed.rows[0]) == OWN_KEYS
        assert token not in "".join(parsed.rows[0].values()), "значения лишних колонок не читаются"
    admin = importer()
    client = support.logged_in(admin)
    response = upload(client, support_org.staff_csv([values], headers=headers))
    assert response.status_code == 302
    item = model("org.StaffImport").objects.get()
    assert item.summary["not_loaded"] == ["Роль", "Дата рождения"]
    shown = client.get(response["Location"])
    assert shown.status_code == 200
    html = shown.content.decode()
    assert len(page(html).items) == 2 and all(
        any(name in text for text in page(html).items) for name in ("Роль", "Дата рождения"))
    assert "не загружаются" in html
    assert token not in html, "значения лишних колонок на странице не показываются"
    # без лишних колонок блока нет
    clean = upload(client, support_org.staff_csv([row]))
    clean_html = client.get(clean["Location"]).content.decode()
    assert page(clean_html).items == [] and model("org.StaffImport").objects.order_by("-id").first().summary[
        "not_loaded"] == []


def test_refused_over_1mb_over_1000_rows_not_csv_or_xlsx_not_utf8_missing_columns():
    """FR-141, NFR-108: отказ без создания чего-либо — больше 1 МБ, больше 1000 строк, не CSV и не XLSX, не UTF-8,
    без обязательных колонок; ровно 1 МБ и ровно 1000 строк принимаются."""
    # граница размера: ровно 1 048 576 принимается, на байт больше — too_big (до разбора, любое содержимое)
    exact = exact_size_csv(MAX_BYTES)
    assert len(service().parse_file(exact).rows) == 1000, "ровно 1 МБ и ровно 1000 строк принимаются"
    refused(exact_size_csv(MAX_BYTES + 1), "too_big")
    refused(b"\x00" * (MAX_BYTES + 1), "too_big")
    # больше 1000 строк (CSV и XLSX)
    refused(support_org.staff_csv(rows_of(1001)), "too_many_rows")
    refused(support_org.staff_xlsx(rows_of(1001)), "too_many_rows", name="people.xlsx")
    assert len(service().parse_file(support_org.staff_xlsx(rows_of(1000))).rows) == 1000
    # не CSV и не XLSX: архив ZIP без частей книги; изображение
    plain_zip = io.BytesIO()
    with zipfile.ZipFile(plain_zip, "w") as archive:
        archive.writestr("readme.txt", "это не книга")
    refused(plain_zip.getvalue(), "bad_format", name="people.xlsx")
    refused(b"\x89PNG\r\n\x1a\n\xff\xfe\xfd" * 50, "not_utf8", name="photo.png")
    # ZIP с `[Content_Types].xml`, но без `xl/workbook.xml` — тоже не книга
    no_book = io.BytesIO()
    with zipfile.ZipFile(no_book, "w") as archive:
        archive.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
        archive.writestr("xl/worksheets/sheet1.xml", "<worksheet/>")
    refused(no_book.getvalue(), "bad_format", name="people.xlsx")
    # порядок проверок: колонки раньше числа строк
    no_email = [h for h in support_org.STAFF_HEADERS if h != "Почта"]
    many_rows = [[support_org.staff_row()[h] for h in no_email] for _ in range(1001)]
    refused(support_org.staff_csv(many_rows, headers=no_email), "missing_columns", must_mention=("Почта",))
    refused(support_org.staff_xlsx(many_rows, headers=no_email), "missing_columns", name="people.xlsx",
            must_mention=("Почта",))
    # не UTF-8: Windows-1251; нулевой байт в данных
    cp1251 = support_org.staff_csv([support_org.staff_row(last_name="Иванов")], bom=False).decode("utf-8").encode(
        "cp1251")
    refused(cp1251, "not_utf8")
    with_nul = support_org.staff_csv([support_org.staff_row(workplace="Каб\x00инет")])
    refused(with_nul, "not_utf8")
    # без обязательных колонок: в тексте названы все недостающие
    headers = ["Фамилия", "Должность", "Отчество"]
    refused(support_org.staff_csv([["Иванов", "Инженер", "Иванович"]], headers=headers), "missing_columns",
            must_mention=("Почта", "Имя"))
    refused(support_org.staff_xlsx([["Иванов", "Инженер", "Иванович"]], headers=headers), "missing_columns",
            name="people.xlsx", must_mention=("Почта", "Имя"))
    # только заголовок со сломанным набором: «Email» вместо «Почта»
    wrong = [("Email" if h == "Почта" else h) for h in support_org.STAFF_HEADERS]
    refused(support_org.staff_csv([[support_org.staff_row()[h] for h in support_org.STAFF_HEADERS]], headers=wrong),
            "missing_columns", must_mention=("Почта",))


def test_format_by_content_not_by_name():
    """NFR-108: формат определяет содержимое, а не имя и не заявленный тип файла."""
    admin = importer()
    client = support.logged_in(admin)
    csv_bytes = support_org.staff_csv([support_org.staff_row()])
    xlsx_bytes = support_org.staff_xlsx([support_org.staff_row()])
    cases = [
        (csv_bytes, "people.xlsx", XLSX_TYPE, "csv"),
        (csv_bytes, "people.xls", "application/vnd.ms-excel", "csv"),
        (xlsx_bytes, "people.csv", "text/csv", "xlsx"),
        (xlsx_bytes, "people", "application/octet-stream", "xlsx"),
    ]
    for content, name, content_type, expected in cases:
        response = upload(client, content, name, content_type)
        assert response.status_code == 302, f"{name}/{content_type}: {response.status_code}"
        item = model("org.StaffImport").objects.order_by("-id").first()
        assert item.file_format == expected, f"{name}: формат по содержимому — {expected}, а не {item.file_format}"
        assert item.file_name == name
    # имя «.xlsx» у не-книги и «.csv» у ZIP не помогают
    fake = io.BytesIO()
    with zipfile.ZipFile(fake, "w") as archive:
        archive.writestr("a.csv", "Почта;Фамилия")
    users = model("core.User").objects.count()
    done = model("org.StaffImport").objects.count()
    for content, name in ((fake.getvalue(), "people.csv"), (fake.getvalue(), "people.xlsx")):
        assert alert_shown(upload(client, content, name, XLSX_TYPE)), name
    assert model("org.StaffImport").objects.count() == done and model("core.User").objects.count() == users


def test_xlsx_zip_bomb_and_xml_bomb_refused():
    """NFR-108: XLSX с больше чем 100 частями или больше 20 МБ в распакованном виде — `zip_bomb`; XML-бомба — `xml_bomb`."""
    bomb = support_org.xlsx_zip_bomb()
    assert len(bomb) < MAX_BYTES, "бомба сама в пределе размера: отказ должен прийти от проверки оглавления"
    refused(bomb, "zip_bomb", name="people.xlsx")
    many = support_org.xlsx_many_parts()
    assert len(many) < MAX_BYTES
    refused(many, "zip_bomb", name="people.xlsx")
    refused(support_org.xlsx_xml_bomb(), "xml_bomb", name="people.xlsx")
    # порядок проверок: оглавление (`zip_bomb`) раньше разбора (`xml_bomb`)
    both = io.BytesIO(support_org.xlsx_xml_bomb())
    with zipfile.ZipFile(both, "a") as archive:
        for number in range(101):
            archive.writestr(f"xl/extra/part{number}.xml", "<a/>")
    assert len(both.getvalue()) < MAX_BYTES
    refused(both.getvalue(), "zip_bomb", name="people.xlsx")
    # обычная книга при этом принимается
    assert service().parse_file(support_org.staff_xlsx([support_org.staff_row()])).file_format == "xlsx"


def test_xlsx_formulas_not_evaluated_numbers_and_dates_as_text():
    """NFR-108: формулы XLSX не вычисляются; числа и даты из ячеек приводятся к тексту (целое — без «.0» и экспоненты)."""
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.append(list(support_org.STAFF_HEADERS))
    email = support_org.staff_row()["Почта"]
    sheet.append([email, "=1+1", "Имя", 123, "Отчество", 79001234567, 101, datetime.date(2026, 3, 5),
                  datetime.datetime(2027, 1, 2, 10, 30), "нет"])
    sheet.append([f"={chr(34)}{email}{chr(34)}", "Ф2", "И2", "007", "", 12345678901234, 5.0, 1.5, "", ""])
    out = io.BytesIO()
    book.save(out)
    parsed = service().parse_file(out.getvalue())
    assert parsed.file_format == "xlsx" and len(parsed.rows) == 2
    first, second = parsed.rows
    for row in parsed.rows:
        assert set(row) == OWN_KEYS and all(isinstance(value, str) for value in row.values()), (
            "все значения — текст, не числа и не даты")
    assert first["last_name"] == "", "формула =1+1 без сохранённого значения — пустая строка, не «2» и не текст формулы"
    assert second["email"] == "", "формула-строка без сохранённого значения — пустая строка"
    assert first["position"] == "123"
    assert first["phone"] == "79001234567", "большое целое — без экспоненты и «.0»"
    assert first["internal_phone"] == "101"
    assert first["workplace"] == "05.03.2026", "дата — текстом ДД.ММ.ГГГГ"
    assert first["unit_path"] == "02.01.2027", "дата-время — ДД.ММ.ГГГГ, время отбрасывается"
    assert second["position"] == "007", "текст с ведущими нулями остаётся как есть"
    assert second["phone"] == "12345678901234" and "e" not in second["phone"].lower()
    assert second["internal_phone"] == "5", "5.0 — целое число без «.0»"
    assert second["workplace"] == "1.5"
    assert first["head_raw"] == "нет"
    # формула с сохранённым значением: читается значение текстом, текста формулы нет
    cached = service().parse_file(xlsx_with_cached_formula("Сидоров")).rows[0]
    assert cached["last_name"] == "Сидоров" and "=" not in cached["last_name"]


def test_file_not_stored():
    """FR-141: файл разбирается в памяти и нигде не сохраняется — ни на диске, ни в базе, ни в журнале."""
    import os
    from django.conf import settings
    token = secrets.token_hex(8)
    marker = f"секрет-{token}"
    headers = list(support_org.STAFF_HEADERS) + ["Заметка"]
    row = support_org.staff_row()
    values = [row[h] for h in support_org.STAFF_HEADERS] + [marker]
    content = support_org.staff_csv([values], headers=headers)
    roots = [p for p in (getattr(settings, "PROTECTED_ROOT", None), getattr(settings, "MEDIA_ROOT", None)) if p]
    assert roots, "в настройках нет каталогов хранения — проверка диска была бы пустой"

    def listing():
        return sorted(os.path.join(d, f) for root in roots if os.path.isdir(root) for d, _, fs in os.walk(root)
                      for f in fs)
    before = listing()
    admin = importer()
    client = support.logged_in(admin)
    response = upload(client, content, "people.csv")
    assert response.status_code == 302
    item = model("org.StaffImport").objects.get()
    assert listing() == before, "на диске после загрузки новых файлов нет"
    for field in item._meta.concrete_fields:
        kind = field.get_internal_type()
        assert kind not in ("FileField", "ImageField", "BinaryField"), f"{field.name}: поле для файла не нужно"
        assert marker not in str(getattr(item, field.attname)), f"{field.name}: значение лишней колонки не хранится"
    assert item.file_name == "people.csv"
    assert model("core.StoredFile").objects.count() == 0, "файл загрузки не регистрируется как хранимый"
    for event in model("core.AuditEvent").objects.all():
        assert marker not in f"{event.before}{event.after}", "содержимое файла не попадает в журнал"
    # то же после просмотра предпросмотра
    assert marker not in client.get(response["Location"]).content.decode()
