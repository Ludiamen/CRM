"""T-6148: сценарий 9 — выгрузка в Excel, загрузка из таблицы, образец (FR-513, FR-514, FR-522, FR-532, FR-533,
FR-534, FR-535, FR-542, FR-546, NFR-507).

Общее для файла (шаблон задания e2e/template.md, CONTRACT-006.md):
- управляющий — администратор круга `admin_page`; читающие и правящие — `helpers_lists.make_list_reader` /
  `make_list_editor`; каждый тест заводит свои списки с уникальными названиями, шесть стартовых не трогает; общих
  настроек не меняет; сброса стенда, команд стенда и `stand_shell` в файле нет;
- файлы для загрузки собирает `helpers_sd.make_table` (CSV с «;» и XLSX), выгруженные файлы читает `openpyxl`
  (как `helpers_sd`); даты и числа в файлах для загрузки пишутся строками;
- строки предпросмотра нумеруются с 2 (шапка — строка 1; К11); причина строки — «<Поле>: <текст>» (12.3);
- п. 4 сценария («больше 10 000 записей в отборе») в приёмке не проходится — серверный T-6117 (решение 19);
- табуляция и возврат каретки в начале значения через форму не заводятся надёжно (обрезка, нормализация `\\r` в
  `\\n`, К12), поэтому п. 3 здесь проходит `=`, `+`, `-`, `@` у значений и заголовков (вопрос к Opus в отчёте);
- тексты — литералы раздела 12 контракта.
"""
import re
from datetime import date, datetime, timedelta, timezone
from urllib.parse import quote

from playwright.sync_api import expect

from tests import factories
from tests.e2e import helpers_lists, helpers_org, helpers_sd
from tests.e2e.helpers_lists import (
    add_field, create_list, create_record, open_add_record, open_record, open_settings, pick,
    set_choices, status_of, submit_record, unique_list_name, unique_word)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv"
SAVE = "button[type=submit]:not([data-action])"   # главная кнопка формы (15.4, п. 1)
MOSCOW = timezone(timedelta(hours=3))             # Москва: UTC+3 без перехода на летнее время
PDF = b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog >>\nendobj\ntrailer\n<< /Root 1 0 R >>\n%%EOF\n"
LONG = 120_000                                    # ожидание медленных шагов (500 строк), мс


# --- Помощники файла -------------------------------------------------------------------------------------------

def _mime(fmt: str) -> str:
    return XLSX_MIME if fmt == "xlsx" else CSV_MIME


def _digits(locator) -> int:
    """Первое целое число в тексте элемента (счётчики, «Найдено: N»)."""
    return int(re.search(r"\d+", locator.inner_text().replace(" ", "").replace(" ", "")).group())


def _found(page, list_id, query: str = "") -> int:
    """«Найдено: N» таблицы списка (`[data-lists-found]`) при параметрах `query` (строка запроса без «?»)."""
    page.goto(f"/lists/{list_id}/" + (f"?{query}" if query else ""))
    expect(page.locator("[data-lists-found]")).to_be_visible()
    return _digits(page.locator("[data-lists-found]"))


def _record_number(page, list_id, title: str) -> int:
    """Номер записи по точному названию: таблица с отбором `f.title`, ссылка `a[data-field="title"]`."""
    page.goto(f"/lists/{list_id}/?f.title={quote(title)}")
    link = page.locator('a[data-field="title"]').filter(has_text=re.compile(rf"^\s*{re.escape(title)}\s*$"))
    expect(link).to_have_count(1)
    return helpers_lists.record_number_from_url(link.get_attribute("href"))


def _value(page, code: str):
    """Значение видимого поля на карточке записи."""
    return page.locator(f'[data-lists-value="{code}"]')


def _grid_headers(page) -> list:
    """Заголовки таблицы: пары (код, название) `th[data-sort]` по порядку."""
    return page.locator("#lists-grid th[data-sort]").evaluate_all(
        "els => els.map(e => [e.getAttribute('data-sort'), ((e.querySelector('a') || e).textContent || '').trim()])")


def _grid_rows(page) -> list:
    """Строки таблицы по порядку: `{"number", "title", "cells": {код: текст}}`."""
    return page.locator("#lists-grid tr[data-record-number]").evaluate_all(
        """rows => rows.map(r => ({number: r.getAttribute('data-record-number'),
            title: ((r.querySelector('a[data-field="title"]') || {}).textContent || '').trim(),
            cells: Object.fromEntries(Array.from(r.querySelectorAll('td[data-col]')).map(
                td => [td.getAttribute('data-col'), (td.textContent || '').trim()]))}))""")


def _set_columns(page, list_id, wanted=None) -> None:
    """Столбцы таблицы читающего (FR-518): диалог «Выбрать столбцы», отмечены коды `wanted` (`None` — все); «Название»
    системное. Ждёт, пока таблица покажет ровно эти столбцы."""
    page.goto(f"/lists/{list_id}/")
    expect(page.locator("#lists-grid")).to_be_visible()
    page.locator('button[data-action="columns"]').click()  # «Выбрать столбцы» (CONTRACT-006, 15.5 п. 6)
    dialog = page.locator("dialog[data-lists-columns]")
    expect(dialog).to_be_visible()
    codes = []
    for box in dialog.locator('input[name="columns"][data-lists-column-choice]').all():
        code = box.get_attribute("data-lists-column-choice")
        if code == "title":
            continue
        codes.append(code)
        box.set_checked(wanted is None or code in wanted)
    # сохранение — POST `/lists/<id>/columns/` (10.2); ждём его ответ, затем таблицу
    with page.expect_response(lambda r: r.url.split("?")[0].endswith(f"/lists/{list_id}/columns/")
                              and r.request.method == "POST"):
        dialog.locator(SAVE).click()
    # выбор сохранён (`UserView`): таблица открывается заново, чтобы не читать прежний DOM и не ловить переход по 302
    page.goto(f"/lists/{list_id}/")
    expect(page.locator("#lists-grid")).to_be_visible()
    # столбец ищется по ячейкам `td[data-col]` (11.3): `th[data-sort]` есть только у сортируемых столбцов, у поля
    # «файл» его нет (решение Opus, review-decisions-1, E2); у всех вызывающих тестов в таблице есть строки
    for code in codes:
        cells = page.locator(f'#lists-grid td[data-col="{code}"]')
        if wanted is None or code in wanted:
            expect(cells.first).to_be_attached()
        else:
            expect(cells).to_have_count(0)


def _read_book(download):
    """Книга выгруженного файла (с форматами ячеек)."""
    import openpyxl

    return openpyxl.load_workbook(download.path())


def _sheet_values(sheet) -> list:
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _download(page, locator):
    """Нажимает ссылку-выгрузку и возвращает скачанный файл."""
    with page.expect_download() as info:
        locator.click()
    return info.value


def _open_import(page, list_id):
    """`/lists/<id>/import/`: ждёт страницу и возвращает локатор формы выбора файла."""
    page.goto(f"/lists/{list_id}/import/")
    expect(page.locator('[data-lists-page="import"]')).to_be_visible()
    form = page.locator('[data-lists-page="import"] form:has(input[type=file][name=file])')
    expect(form).to_have_count(1)
    return form


def _send_file(form, content: bytes, name: str, mime: str) -> None:
    form.locator('input[type=file][name="file"]').set_input_files({"name": name, "mimeType": mime, "buffer": content})
    form.locator(SAVE).click()


def _to_mapping(page, list_id, content: bytes, name: str, mime: str, *, timeout: int = 30_000) -> str:
    """Выбор файла и «Загрузить»: страница сопоставления. Возвращает номер задания из адреса."""
    _send_file(_open_import(page, list_id), content, name, mime)
    expect(page.locator('[data-lists-page="mapping"]')).to_be_visible(timeout=timeout)
    return re.search(r"/import/(\d+)/", page.url).group(1)


def _mapping_form(page):
    return page.locator('[data-lists-page="mapping"] form:has(select[data-lists-map-column])')


def _check(page, *, timeout: int = 30_000) -> None:
    """«Проверить» на сопоставлении: страница предпросмотра."""
    _mapping_form(page).locator(SAVE).click()
    expect(page.locator('[data-lists-page="preview"]')).to_be_visible(timeout=timeout)


def _confirm(page, *, timeout: int = 30_000) -> None:
    """Подтверждение на предпросмотре: страница итога."""
    page.locator('button[data-action="import-confirm"]').click()
    expect(page.locator('[data-lists-page="summary"]')).to_be_visible(timeout=timeout)


def _preview_rows(page) -> dict:
    """Строки предпросмотра: номер строки файла -> `{"result", "reason", "number"}`."""
    rows = page.locator("tr[data-import-row]").evaluate_all(
        """rows => rows.map(r => ({row: r.getAttribute('data-import-row'), result: r.getAttribute('data-result'),
            reason: ((r.querySelector('[data-field="reason"]') || {}).textContent || '').trim(),
            number: ((r.querySelector('[data-field="number"]') || {}).textContent || '').trim()}))""")
    return {int(r["row"]): r for r in rows}


def _counts(page) -> dict:
    return {kind: _digits(page.locator(f'[data-import-count="{kind}"]')) for kind in ("will_create", "will_update", "error")}


def _summary(page) -> dict:
    return {kind: _digits(page.locator(f'[data-import-summary="{kind}"]')) for kind in ("created", "updated", "errors")}


def _import(page, list_id, rows: list, fmt: str, name: str, *, mapping: dict | None = None, timeout: int = 30_000):
    """Весь путь загрузки: файл, сопоставление (поправки `mapping`: номер столбца -> код), «Проверить», предпросмотр.
    Возвращает номер задания; страница остаётся на предпросмотре."""
    job = _to_mapping(page, list_id, helpers_sd.make_table(rows, fmt), name, _mime(fmt), timeout=timeout)
    for column, code in (mapping or {}).items():
        _mapping_form(page).locator(f'select[name="map.{column}"]').select_option(value=code)
    _check(page, timeout=timeout)
    return job


def _equipment(page, label: str = "Оборудование"):
    """Список с полями «Инвентарный номер» (строка, уникальное) и «Модель» (строка); возвращает
    `(номер списка, код инвентарного номера, код модели)`."""
    list_id = create_list(page, unique_list_name(label))
    inv = add_field(page, list_id, "Инвентарный номер", "string", unique=True)
    model = add_field(page, list_id, "Модель", "string")
    return list_id, inv, model


# --- Тесты -----------------------------------------------------------------------------------------------------

def test_l09_export_downloads_xlsx_with_field_headers_and_rows_order_and_columns_as_in_filtered_table(
        admin_page, employee):
    """Сценарий 9, п. 1: читающий выбрал столбцы, отбор и сортировку; «Выгрузить в Excel» сразу скачивает XLSX:
    заголовки — названия полей, строки, их порядок и столбцы — как в таблице (FR-532)."""
    page = admin_page
    tag = unique_word()
    name = unique_list_name("Выгрузка")
    list_id = create_list(page, name)
    inv = add_field(page, list_id, "Инвентарный номер", "string")
    state = add_field(page, list_id, "Состояние", "choice")
    qty = add_field(page, list_id, "Количество", "number")
    new_id, _working_id, _written_off_id = set_choices(page, list_id, state, ["Новое", "Рабочее", "Списано"])
    plan = (("Ноутбук", "Новое"), ("Монитор", "Рабочее"), ("Принтер", "Новое"), ("Сканер", "Списано"))
    for index, (word, status) in enumerate(plan, start=1):
        create_record(page, list_id, f"{word} {tag}", {inv: f"INV-{tag}-{index}", state: status, qty: str(index)})
    reader = helpers_lists.make_list_reader(employee, page, list_id)
    rp = reader.page
    # п. 1: своя выборка читающего — без столбца «Количество», отбор «Новое», сортировка по названию по убыванию
    _set_columns(rp, list_id, {inv, state})
    assert _found(rp, list_id, f"sort=title&dir=desc&f.{state}={new_id}") == 2
    rows = _grid_rows(rp)
    # заголовки выгрузки — названия полей (план «Выгрузка») в порядке столбцов таблицы; текст ссылки заголовка не
    # берётся: у текущего столбца сортировки в нём может быть знак порядка (11.3 задаёт только `[data-dir]`)
    names = {"title": "Название", inv: "Инвентарный номер", state: "Состояние", qty: "Количество"}
    codes = [code for code, _text in _grid_headers(rp) if code != "number"]
    assert qty not in codes, "столбец «Количество» не скрыт"
    headers = [names[code] for code in codes]
    assert [r["title"] for r in rows] == [f"Принтер {tag}", f"Ноутбук {tag}"], "порядок таблицы не по убыванию названия"
    download = _download(rp, rp.locator('a[data-action="export"]'))
    assert download.suggested_filename == f"{name}.xlsx"
    values = _sheet_values(_read_book(download).worksheets[0])
    # заголовки — названия полей в порядке таблицы
    assert [str(v) for v in values[0]] == headers
    # строки — те же, в том же порядке, ячейки — как на экране
    title_at = headers.index("Название")
    assert [row[title_at] for row in values[1:]] == [r["title"] for r in rows]
    inv_at = headers.index("Инвентарный номер")
    assert [row[inv_at] for row in values[1:]] == [r["cells"][inv] for r in rows]
    exported = " ".join(str(c) for row in values for c in row if c is not None)
    for word in ("Монитор", "Сканер"):
        assert f"{word} {tag}" not in exported, f"в выгрузке запись «{word}» вне отбора"


def test_l09_exported_cells_have_dates_full_names_unit_and_record_titles_file_names_yes_no_and_money_with_currency(
        admin_page, employee):
    """Сценарий 9, п. 2: даты — ДД.ММ.ГГГГ, сотрудник — ФИО, подразделение — название, запись — название, файлы —
    имена, да/нет — «Да»/«Нет», деньги — число с валютой в формате ячейки (FR-532)."""
    page = admin_page
    tag = unique_word()
    full_name = f"Выгрузкин{tag} Пётр Сергеевич"
    employee(full_name=full_name)
    unit_name = helpers_org.unique_name("Отдел выгрузки")
    unit_id = helpers_org.create_unit(page, unit_name)
    target_id = create_list(page, unique_list_name("Справочник"))
    ref_title = f"Модель {tag}"
    create_record(page, target_id, ref_title)
    list_id = create_list(page, unique_list_name("Типы"))
    f_date = add_field(page, list_id, "Куплено", "date")
    f_time = add_field(page, list_id, "Введено", "datetime")
    f_person = add_field(page, list_id, "Ответственный", "person")
    f_unit = add_field(page, list_id, "Подразделение", "unit")
    f_record = add_field(page, list_id, "Образец", "record", target=target_id)
    f_weight = add_field(page, list_id, "Вес", "number")
    f_price = add_field(page, list_id, "Цена", "money", currency="RUB")
    f_dollars = add_field(page, list_id, "Цена в долларах", "money", currency="USD")
    f_ok = add_field(page, list_id, "Исправно", "bool")
    f_state = add_field(page, list_id, "Состояние", "choice")
    f_scan = add_field(page, list_id, "Скан", "file")
    set_choices(page, list_id, f_state, ["Новое", "Рабочее"])
    title_a, title_b = f"Устройство А {tag}", f"Устройство Б {tag}"
    open_add_record(page, list_id)
    helpers_lists.fill_record(page, {"title": title_a, f_date: "25.10.2026", f_time: "25.10.2026 14:30",
                                     f_weight: "12,5", f_price: "1234,50", f_dollars: "99,90", f_ok: True,
                                     f_state: "Рабочее"})
    pick(page, f_person, tag)
    pick(page, f_unit, unit_name, pick_id=unit_id)
    pick(page, f_record, tag)
    number_a = submit_record(page)
    assert number_a, helpers_lists.refusal_codes(page)
    create_record(page, list_id, title_b)    # у второй записи да/нет не отмечено, остальное пусто
    file_name = f"scan-{tag}.pdf"
    file_id, _seen = helpers_lists.attach_file(page, list_id, number_a, f_scan, PDF, file_name, mime="application/pdf")
    assert helpers_lists.wait_file_clean(page, list_id, number_a, file_id) == "clean"
    _set_columns(page, list_id)    # все столбцы
    download = _download(page, page.locator('a[data-action="export"]'))
    sheet = _read_book(download).worksheets[0]
    rows = list(sheet.iter_rows())
    head = [str(c.value) for c in rows[0]]
    by_title = {r[head.index("Название")].value: r for r in rows[1:]}
    assert {title_a, title_b} <= set(by_title)
    a = by_title[title_a]

    def cell(row, name):
        return row[head.index(name)]

    # п. 2: дата — дата с форматом ДД.ММ.ГГГГ
    got = cell(a, "Куплено")
    assert (got.value.date() if isinstance(got.value, datetime) else got.value) == date(2026, 10, 25)
    assert got.number_format.upper() == "DD.MM.YYYY"
    # п. 2: дата и время — текст ДД.ММ.ГГГГ ЧЧ:ММ по Москве
    assert cell(a, "Введено").value == "25.10.2026 14:30"
    # п. 2: сотрудник — ФИО, подразделение и запись — названия, файл — имя
    assert cell(a, "Ответственный").value == full_name
    assert cell(a, "Подразделение").value == unit_name
    assert cell(a, "Образец").value == ref_title
    assert cell(a, "Скан").value == file_name
    assert cell(a, "Состояние").value == "Рабочее"
    # п. 2: да/нет — «Да» / «Нет»
    assert cell(a, "Исправно").value == "Да"
    assert cell(by_title[title_b], "Исправно").value == "Нет"
    # п. 2: деньги и число — числовые ячейки, валюта — в формате
    assert float(cell(a, "Вес").value) == 12.5
    for name, amount, sign in (("Цена", 1234.5, "₽"), ("Цена в долларах", 99.9, "$")):
        money = cell(a, name)
        assert isinstance(money.value, (int, float)), f"«{name}» не числовая ячейка"
        assert float(money.value) == amount
        assert sign in money.number_format, f"в формате «{name}» нет знака {sign}: {money.number_format}"


def test_l09_exported_text_cells_starting_with_formula_chars_get_leading_apostrophe(admin_page):
    """Сценарий 9, п. 3: текст, начинающийся с `=`, `+`, `-`, `@` (значения и заголовки), выгружается с апострофом
    впереди; обычный текст, числа и даты не трогаются (FR-532)."""
    page = admin_page
    tag = unique_word()
    list_id = create_list(page, unique_list_name("Формулы"))
    f_text = add_field(page, list_id, "=Формула", "string")
    f_num = add_field(page, list_id, "-Остаток", "number")
    # п. 3: четыре начала формулы и обычный текст в названии записи и в строковом поле
    samples = (f"=1+1 {tag}", f"+7 999 {tag}", f"-минус {tag}", f"@имя {tag}", f"обычный {tag}")
    for value in samples:
        create_record(page, list_id, value, {f_text: value, f_num: "-5"})
    _set_columns(page, list_id)
    download = _download(page, page.locator('a[data-action="export"]'))
    sheet = _read_book(download).worksheets[0]
    rows = list(sheet.iter_rows())
    head = [c.value for c in rows[0]]
    # заголовки, начинающиеся с «=» и «-», тоже защищены
    assert "'=Формула" in head and "'-Остаток" in head, f"заголовки не защищены: {head}"
    assert "Название" in head, "обычный заголовок «Название» изменён"
    by_title = {}
    for row in rows[1:]:
        by_title[row[head.index("Название")].value] = row
    for value in samples[:4]:
        row = by_title[f"'{value}"]
        assert row[head.index("'=Формула")].value == f"'{value}", f"строковое поле «{value}» не защищено"
    plain = by_title[samples[4]]
    assert plain[head.index("'=Формула")].value == samples[4], "обычный текст получил апостроф"
    # число «-5» остаётся числом: апостроф ставится только текстовым ячейкам
    number = plain[head.index("'-Остаток")]
    assert isinstance(number.value, (int, float)) and number.value == -5


def test_l09_import_of_xlsx_and_csv_shows_mapping_of_columns_to_fields_by_names_editable_by_hand(admin_page):
    """Сценарий 9, п. 5: правящий нажимает «Загрузить из таблицы», выбирает XLSX или CSV; видит сопоставление
    столбцов с полями по названиям (без учёта регистра и крайних пробелов) и поправляет его вручную (FR-533)."""
    page = admin_page
    list_id, inv, model = _equipment(page)
    stock = add_field(page, list_id, "Склад", "string")
    tag = unique_word()
    for fmt in ("xlsx", "csv"):
        title, number, kind, place = f"Позиция {fmt} {tag}", f"{fmt}-{tag}", f"Модель-{fmt}", f"Склад-{fmt}"
        rows = [["название", "  Инвентарный номер ", "МОДЕЛЬ", "Заводской"], [title, number, kind, place]]
        # п. 5: кнопка «Загрузить из таблицы» на странице таблицы ведёт на страницу загрузки
        page.goto(f"/lists/{list_id}/")
        # у пустого списка ссылка может быть и в панели, и в пустом состоянии (план «Таблица записей»)
        page.locator('a[data-action="import"]').first.click()
        expect(page.locator('[data-lists-page="import"]')).to_be_visible()
        form = page.locator('[data-lists-page="import"] form:has(input[type=file][name=file])')
        _send_file(form, helpers_sd.make_table(rows, fmt), f"import-{tag}.{fmt}", _mime(fmt))
        expect(page.locator('[data-lists-page="mapping"]')).to_be_visible()
        # п. 5: столбцы сопоставлены по названиям; «Заводской» не совпал ни с одним полем
        selects = _mapping_form(page).locator("select[data-lists-map-column]")
        expect(selects).to_have_count(4)
        expected = {0: "title", 1: inv, 2: model, 3: ""}
        for column, code in expected.items():
            assert _mapping_form(page).locator(f'select[name="map.{column}"]').input_value() == code, \
                f"{fmt}: столбец {column} сопоставлен не с «{code}»"
        # п. 5: сопоставление правится вручную — «Заводской» → «Склад»
        _mapping_form(page).locator('select[name="map.3"]').select_option(value=stock)
        _check(page)
        assert _counts(page) == {"will_create": 1, "will_update": 0, "error": 0}
        _confirm(page)
        # правка вручную сработала: значения лежат в тех полях, которые выбраны
        number_ = _record_number(page, list_id, title)
        open_record(page, list_id, number_)
        expect(_value(page, stock)).to_have_text(place)
        expect(_value(page, model)).to_have_text(kind)
        expect(_value(page, inv)).to_have_text(number)


def test_l09_file_over_5mb_over_5000_rows_or_other_format_is_refused_with_reason_and_nothing_created(admin_page):
    """Сценарий 9, п. 6: файл больше 5 МБ, больше 5 000 строк или другого формата отклоняется с причиной у поля
    «файл»; страница остаётся страницей загрузки, ничего не создано (FR-533, FR-546)."""
    page = admin_page
    list_id, _inv, _model = _equipment(page)
    before = _found(page, list_id)
    head = "Название;Инвентарный номер\r\n"
    big = (head + "".join(f"Позиция {i};{'x' * 1480}\r\n" for i in range(4000))).encode("utf-8")
    assert 5 * 1024 * 1024 < len(big) < 6 * 1024 * 1024, "проверочный файл вне вилки: больше 5 МБ и меньше 6 МБ"
    many = helpers_sd.make_table([["Название", "Инвентарный номер"]] + [[f"П {i}", f"N{i}"] for i in range(5001)], "csv")
    other = b"%PDF-1.4\n\x00\x01\x02 not a table \xff\xfe\n"
    cases = (("big.csv", CSV_MIME, big, "import_too_big", "Файл больше 5 МБ"),
             ("many.csv", CSV_MIME, many, "import_too_many_rows", "Строк больше 5 000"),
             ("doc.pdf", "application/pdf", other, "import_bad_format", "Формат не поддерживается"))
    for name, mime, content, code, text in cases:
        form = _open_import(page, list_id)
        _send_file(form, content, name, mime)
        refusal = page.locator(f'[data-error="file"][data-code="{code}"]')
        expect(refusal).to_have_count(1, timeout=60_000)
        expect(refusal).to_contain_text(text)
        expect(page.locator('[data-lists-page="import"]')).to_be_visible()
        assert "/import/" in page.url and not re.search(r"/import/\d+/", page.url), "перешли на сопоставление"
    # п. 6: ничего не создано
    assert _found(page, list_id) == before == 0


def test_l09_preview_marks_rows_will_be_updated_by_unique_value_will_be_created_or_error_with_reason_and_row_number(
        admin_page):
    """Сценарий 9, п. 7: сопоставлено уникальное поле — строка с имеющимся значением «будет обновлена», иначе «будет
    создана»; строка с ошибкой — «с ошибкой» с причиной; номер строки — в строке предпросмотра (FR-534)."""
    page = admin_page
    list_id, inv, _model = _equipment(page)
    tag = unique_word()
    existing = create_record(page, list_id, f"Имеется {tag}", {inv: f"E-{tag}"})
    before = _found(page, list_id)
    rows = [["Название", "Инвентарный номер"],
            [f"Обновление {tag}", f"E-{tag}"],      # строка 2: значение уже есть
            [f"Новая {tag}", f"N-{tag}"],           # строка 3: значения нет
            ["", f"M-{tag}"],                       # строка 4: нет названия у создаваемой записи
            [f"Повтор {tag}", f"N-{tag}"]]          # строка 5: значение повторяет строку 3
    _import(page, list_id, rows, "xlsx", f"preview-{tag}.xlsx")
    preview = _preview_rows(page)
    assert set(preview) == {2, 3, 4, 5}, f"номера строк предпросмотра: {sorted(preview)}"
    # п. 7: «будет обновлена» (с номером записи), «будет создана»
    assert preview[2]["result"] == "will_update" and str(existing) in preview[2]["number"]
    assert preview[3]["result"] == "will_create"
    # п. 7: «с ошибкой» и причина
    assert preview[4]["result"] == "error" and "Не указано название" in preview[4]["reason"]
    assert preview[5]["result"] == "error" and "Значение повторяется в строке 3" in preview[5]["reason"]
    assert _counts(page) == {"will_create": 1, "will_update": 1, "error": 2}
    # до подтверждения ничего не меняется
    assert _found(page, list_id) == before


def test_l09_preview_marks_unknown_choice_ambiguous_full_name_unknown_email_and_bad_date_as_row_errors_with_reasons(
        admin_page, employee):
    """Сценарий 9, п. 8: неизвестное значение списка, неоднозначное ФИО, неизвестная почта сотрудника и дата не в
    ДД.ММ.ГГГГ — строки с ошибкой и причиной «<Поле>: <текст>» (FR-533)."""
    page = admin_page
    tag = unique_word()
    twin = f"Двойников{tag} Иван Иванович"
    employee(full_name=twin)
    employee(full_name=twin)
    known = employee(full_name=f"Известнов{tag} Пётр Петрович")
    ghost = factories.new_email("ghost")
    list_id = create_list(page, unique_list_name("Ошибки"))
    state = add_field(page, list_id, "Состояние", "choice")
    add_field(page, list_id, "Ответственный", "person")
    add_field(page, list_id, "Дата ввода", "date")
    set_choices(page, list_id, state, ["Новое", "Рабочее"])
    rows = [["Название", "Состояние", "Ответственный", "Дата ввода"],
            [f"Верная {tag}", "Новое", known.email, "25.10.2026"],      # строка 2
            [f"Значение {tag}", "Несуществующее", "", ""],              # строка 3
            [f"Тёзка {tag}", "", twin, ""],                             # строка 4
            [f"Почта {tag}", "", ghost, ""],                            # строка 5
            [f"Дата {tag}", "", "", "завтра"]]                          # строка 6
    _import(page, list_id, rows, "csv", f"errors-{tag}.csv")
    preview = _preview_rows(page)
    assert preview[2]["result"] == "will_create", preview[2]
    expected = {3: ("Состояние", "Не найдено: Несуществующее"), 4: ("Ответственный", f"Неоднозначно: {twin}"),
                5: ("Ответственный", f"Не найдено: {ghost}"), 6: ("Дата ввода", "Введите дату ДД.ММ.ГГГГ")}
    for row, (field, text) in expected.items():
        assert preview[row]["result"] == "error", f"строка {row}: {preview[row]}"
        assert field in preview[row]["reason"] and text in preview[row]["reason"], \
            f"строка {row}: причина «{preview[row]['reason']}»"
    assert _counts(page) == {"will_create": 1, "will_update": 0, "error": 4}


def test_l09_confirm_loads_valid_rows_skips_error_rows_shows_summary_and_journal_has_counts_without_values(
        admin_page):
    """Сценарий 9, п. 9: подтверждение создаёт и обновляет строки без ошибок, строки с ошибками не загружает; итог
    «Создано: N, обновлено: M, с ошибками: K» с временем по Москве; в журнале итог без значений (FR-534, FR-542)."""
    page = admin_page
    list_id, inv, model = _equipment(page)
    add_field(page, list_id, "Дата ввода", "date")
    tag = unique_word()
    existing_title = f"Имеется {tag}"
    existing = create_record(page, list_id, existing_title, {inv: f"E-{tag}", model: "старая"})
    before = _found(page, list_id)
    file_name = f"confirm-{tag}.xlsx"
    new1, new2, bad = f"Новая-1 {tag}", f"Новая-2 {tag}", f"Ошибка {tag}"
    rows = [["Название", "Инвентарный номер", "Модель", "Дата ввода"],
            [new1, f"N1-{tag}", f"модель-секрет-1-{tag}", ""],
            [new2, f"N2-{tag}", f"модель-секрет-2-{tag}", "01.11.2026"],
            [existing_title, f"E-{tag}", f"обновлено-{tag}", ""],
            [bad, f"N3-{tag}", f"модель-секрет-3-{tag}", "завтра"]]
    job = _import(page, list_id, rows, "xlsx", file_name)
    assert _counts(page) == {"will_create": 2, "will_update": 1, "error": 1}
    _confirm(page)
    # п. 9: итог — числа и время завершения по Москве
    notice = page.locator('[data-lists-notice="import-summary"]')
    expect(notice).to_contain_text("Создано: 2, обновлено: 1, с ошибками: 1")
    assert _summary(page) == {"created": 2, "updated": 1, "errors": 1}
    stamp = notice.locator("time[data-at]")
    expect(stamp).to_have_count(1)
    moment = datetime.fromisoformat(stamp.get_attribute("datetime").replace("Z", "+00:00"))
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=MOSCOW)
    assert abs(datetime.now(timezone.utc) - moment) < timedelta(minutes=10), "время итога не близко к текущему"
    assert stamp.inner_text().strip() == moment.astimezone(MOSCOW).strftime("%d.%m.%Y %H:%M"), \
        "время итога не по Москве"
    # п. 9: строки без ошибок загружены, строка с ошибкой — нет
    assert _found(page, list_id) == before + 2
    assert _found(page, list_id, f"f.title={quote(bad)}") == 0
    for title in (new1, new2):
        assert _found(page, list_id, f"f.title={quote(title)}") == 1
    open_record(page, list_id, existing)
    expect(_value(page, model)).to_have_text(f"обновлено-{tag}")
    # п. 9: журнал — одна строка итога: имя файла и числа, без значений из файла
    entries = helpers_lists.audit_rows(page, "lists.import_done", job)
    assert len(entries) == 1, f"строк журнала итога загрузки: {len(entries)}"
    text = entries[0]
    assert file_name in text
    for label, count in (("создано", 2), ("обновлено", 1), ("с ошибками", 1)):
        assert re.search(rf"{label}\D{{0,5}}{count}\b", text), f"в журнале нет «{label}: {count}»: {text}"
    for leaked in (new1, new2, bad, f"N1-{tag}", f"модель-секрет-1-{tag}", f"обновлено-{tag}"):
        assert leaked not in text, f"в журнале значение из файла: {leaked}"


def test_l09_employee_without_edit_right_has_no_import_button_and_direct_import_url_is_refused(
        admin_page, employee):
    """Сценарий 9, п. 10: у сотрудника без права правки нет кнопки «Загрузить из таблицы», прямое обращение к
    загрузке и образцу — отказ (403 у читающего, 404 у не читающего); у правящего кнопка есть (FR-513, FR-514)."""
    page = admin_page
    list_id, _inv, _model = _equipment(page)
    create_record(page, list_id, f"Запись {unique_word()}")
    reader = helpers_lists.make_list_reader(employee, page, list_id)
    editor = helpers_lists.make_list_editor(employee, page, list_id)
    stranger = employee(role=helpers_lists.make_list_role(page)[1])
    # п. 10: правящий видит кнопку и открывает загрузку
    editor.page.goto(f"/lists/{list_id}/")
    expect(editor.page.locator('a[data-action="import"]')).to_have_count(1)
    assert status_of(editor.page, f"/lists/{list_id}/import/") == 200
    assert status_of(editor.page, f"/lists/{list_id}/sample/") == 200
    # п. 10: читающий — кнопки нет (выгрузка есть), прямое обращение — 403
    reader.page.goto(f"/lists/{list_id}/")
    expect(reader.page.locator('[data-lists-page="grid"]')).to_be_visible()
    expect(reader.page.locator('a[data-action="export"]')).to_have_count(1)
    expect(reader.page.locator('a[data-action="import"]')).to_have_count(0)
    assert status_of(reader.page, f"/lists/{list_id}/import/") == 403
    assert status_of(reader.page, f"/lists/{list_id}/sample/") == 403
    assert helpers_lists.post_status(reader.page, f"/lists/{list_id}/import/") == 403
    # без прав на список — тот же отказ, что у несуществующего (404)
    for path in (f"/lists/{list_id}/import/", f"/lists/{list_id}/sample/"):
        assert status_of(stranger.page, path) == 404
    assert helpers_lists.post_status(stranger.page, f"/lists/{list_id}/import/") == 404


def test_l09_sample_xlsx_has_headers_of_visible_fields_except_file_and_hints_sheet(admin_page):
    """Сценарий 9, п. 11: «Скачать образец» отдаёт XLSX: первый лист — «Название» и видимые поля кроме файлов (скрытое
    поле не входит), второй лист «Подсказки» — по строке на каждый заголовок (FR-535)."""
    page = admin_page
    name = unique_list_name("Образец")
    list_id = create_list(page, name)
    add_field(page, list_id, "Инвентарный номер", "string", required=True, unique=True)
    f_state = add_field(page, list_id, "Состояние", "choice")
    add_field(page, list_id, "Дата ввода", "date")
    add_field(page, list_id, "Время сдачи", "datetime")
    add_field(page, list_id, "Цена", "money", currency="RUB")
    add_field(page, list_id, "Скан договора", "file")
    f_hidden = add_field(page, list_id, "Скрытое", "string")
    set_choices(page, list_id, f_state, ["Новое", "Рабочее"])
    # скрытое поле: «Скрыть» на вкладке «Поля»
    open_settings(page, list_id, "fields")
    row = page.locator(f'[data-lists-fields-grid] tr[data-field-code="{f_hidden}"]')
    row.locator('button[data-action="field-hide"]').click()
    # `data-hidden` — на ячейке названия строки (15.5 п. 7)
    expect(page.locator(f'[data-lists-fields-grid] tr[data-field-code="{f_hidden}"] [data-field="name"]'
                        '[data-hidden="true"]')).to_be_attached()
    # п. 11: «Скачать образец» — `a[data-action="sample"]` на странице загрузки (15.5 п. 6)
    page.goto(f"/lists/{list_id}/import/")
    link = page.locator('[data-lists-page="import"] a[data-action="sample"]')
    expect(link).to_have_count(1)
    expect(link).to_have_attribute("href", f"/lists/{list_id}/sample/")
    expect(link).to_contain_text("Скачать образец")
    download = _download(page, link)
    assert download.suggested_filename == f"{name} — образец.xlsx"
    book = _read_book(download)
    assert book.sheetnames[1] == "Подсказки"
    first = _sheet_values(book.worksheets[0])
    assert first[0] == ["Название", "Инвентарный номер", "Состояние", "Дата ввода", "Время сдачи", "Цена"]
    # строк данных образец может не иметь или иметь строку-пример: план этого не задаёт
    # п. 11: лист подсказок — по строке на каждый заголовок (12.4)
    hints = [["" if c is None else str(c) for c in r] for r in _sheet_values(book.worksheets[1])]
    assert hints[0] == ["Поле", "Тип", "Обязательное", "Уникальное", "Значения", "Формат"]
    assert hints[1:] == [
        ["Название", "Строка", "Да", "Нет", "", ""],
        ["Инвентарный номер", "Строка", "Да", "Да", "", ""],
        ["Состояние", "Список значений", "Нет", "Нет", "Новое; Рабочее", ""],
        ["Дата ввода", "Дата", "Нет", "Нет", "", "ДД.ММ.ГГГГ"],
        ["Время сдачи", "Дата и время", "Нет", "Нет", "", "ДД.ММ.ГГГГ ЧЧ:ММ"],
        ["Цена", "Деньги", "Нет", "Нет", "", "1234,50"]]


def test_l09_file_field_is_not_offered_in_mapping_and_preview_uploads_no_files(admin_page):
    """Сценарий 9, п. 12: в сопоставлении нет поля «файл»; даже принудительно выбранный столбец не загружается —
    после подтверждения у записи нет файлов (FR-533)."""
    page = admin_page
    list_id, inv, _model = _equipment(page)
    scan = add_field(page, list_id, "Скан договора", "file")
    tag = unique_word()
    title = f"С файлом {tag}"
    rows = [["Название", "Инвентарный номер", "Скан договора"], [title, f"S-{tag}", f"scan-{tag}.pdf"]]
    _to_mapping(page, list_id, helpers_sd.make_table(rows, "xlsx"), f"scan-{tag}.xlsx", XLSX_MIME)
    # п. 12: поля «файл» среди вариантов нет, столбец «Скан договора» не сопоставлен
    select = _mapping_form(page).locator('select[name="map.2"]')
    options = select.locator("option").evaluate_all("els => els.map(e => [e.value, e.textContent.trim()])")
    assert scan not in [value for value, _text in options], "поле «файл» предложено в сопоставлении"
    assert "Скан договора" not in [text for _value_, text in options], "поле «файл» названо среди вариантов"
    assert select.input_value() == ""
    # принудительно выбранный код поля «файл» отбрасывается как «не загружать» (5.7)
    select.evaluate("""(el, code) => { const o = document.createElement('option'); o.value = code;
        o.textContent = 'принудительно'; el.add(o); el.value = code; }""", scan)
    _check(page)
    assert _preview_rows(page)[2]["result"] == "will_create"
    _confirm(page)
    number = _record_number(page, list_id, title)
    open_record(page, list_id, number)
    expect(_value(page, inv)).to_have_text(f"S-{tag}")
    expect(page.locator(f'[data-lists-files="{scan}"] li[data-file-id]')).to_have_count(0)


def test_l09_record_updated_by_import_shows_changed_by_import_in_history(admin_page):
    """Сценарий 9, п. 13: запись, обновлённая загрузкой, в истории — «Изменено загрузкой»; созданная загрузкой —
    «Создана загрузкой» (FR-522, FR-534)."""
    page = admin_page
    list_id, inv, model = _equipment(page)
    tag = unique_word()
    title = f"Обновляемая {tag}"
    number = create_record(page, list_id, title, {inv: f"U-{tag}", model: "старая"})
    new_title = f"Созданная {tag}"
    rows = [["Название", "Инвентарный номер", "Модель"], [title, f"U-{tag}", f"новая-{tag}"],
            [new_title, f"C-{tag}", "создана"]]
    _import(page, list_id, rows, "csv", f"history-{tag}.csv")
    assert _counts(page) == {"will_create": 1, "will_update": 1, "error": 0}
    _confirm(page)
    # п. 13: у обновлённой записи — «Изменено загрузкой» с изменением поля
    open_record(page, list_id, number)
    expect(_value(page, model)).to_have_text(f"новая-{tag}")
    imported = page.locator('[data-lists-history] li[data-change-kind="imported"]')
    expect(imported).to_have_count(1)
    expect(imported).to_contain_text("Изменено загрузкой")
    expect(imported.locator(f'[data-change-field="{model}"]')).to_be_visible()
    # у созданной загрузкой записи — «Создана загрузкой»
    open_record(page, list_id, _record_number(page, list_id, new_title))
    created = page.locator('[data-lists-history] li[data-change-kind="created"][data-via-import], '
                           '[data-lists-history] li[data-change-kind="created"] [data-via-import]').first
    expect(created).to_be_attached()
    expect(page.locator('[data-lists-history] li[data-change-kind="created"]')).to_contain_text("Создана загрузкой")


def test_l09_row_matching_unique_value_of_archived_record_is_error_record_in_archive(admin_page):
    """Сценарий 9, п. 14: строка с уникальным значением записи из архива — «с ошибкой» «Запись в архиве»; запись
    не создаётся и не обновляется (FR-534)."""
    page = admin_page
    list_id, inv, model = _equipment(page)
    tag = unique_word()
    number = create_record(page, list_id, f"Архивная {tag}", {inv: f"A-{tag}", model: "до архива"})
    helpers_lists.archive_record(page, list_id, number)
    rows = [["Название", "Инвентарный номер", "Модель"], [f"Из файла {tag}", f"A-{tag}", "из файла"]]
    _import(page, list_id, rows, "xlsx", f"archive-{tag}.xlsx")
    row = _preview_rows(page)[2]
    assert row["result"] == "error" and "Запись в архиве" in row["reason"], row
    assert _counts(page) == {"will_create": 0, "will_update": 0, "error": 1}
    # запись в архиве не изменена, новой нет (до подтверждения ничего не меняется)
    open_record(page, list_id, number)
    expect(_value(page, model)).to_have_text("до архива")
    assert _found(page, list_id) == 0


def test_l09_list_of_500_rows_loads_in_one_pass_preview_confirm_summary(admin_page):
    """Сценарий 9, п. 5, 7, 9 (запас 500 строк): файл в 500 строк проходит предпросмотр, подтверждение и итог за один
    проход, все записи созданы (FR-533, FR-534)."""
    page = admin_page
    list_id, inv, _model = _equipment(page)
    tag = unique_word()
    rows = [["Название", "Инвентарный номер", "Модель"]] + [
        [f"Позиция {i} {tag}", f"{tag}-{i}", f"Модель {i % 7}"] for i in range(500)]
    _import(page, list_id, rows, "csv", f"rows500-{tag}.csv", timeout=LONG)
    assert _counts(page) == {"will_create": 500, "will_update": 0, "error": 0}
    _confirm(page, timeout=LONG)
    assert _summary(page) == {"created": 500, "updated": 0, "errors": 0}
    expect(page.locator('[data-lists-notice="import-summary"]')).to_contain_text("Создано: 500, обновлено: 0, с ошибками: 0")
    assert _found(page, list_id) == 500
    assert _found(page, list_id, f"f.{inv}={tag}-499") == 1


def test_l09_import_mapping_preview_and_summary_pages_do_not_widen_page_at_360px(
        browser, browser_context_args, admin_page):
    """Сценарий 9, NFR-507: на 360 px загрузка, сопоставление, предпросмотр и итог не дают горизонтальной прокрутки
    страницы."""
    page = admin_page
    list_id, _inv, _model = _equipment(page)
    tag = unique_word()
    long_title = "Очень" + "длинноеназвание" * 6 + tag
    rows = [["Название", "Инвентарный номер", "Модель"], [long_title, f"W-{tag}", "модель"],
            [f"Короткая {tag}", f"W2-{tag}", "модель-2"], ["", f"W3-{tag}", "без названия"]]
    context, narrow = helpers_lists.open_narrow(browser, browser_context_args, page, f"/lists/{list_id}/import/",
                                                "import")
    try:
        assert helpers_lists.page_fits(narrow), "горизонтальная прокрутка на странице загрузки"
        form = narrow.locator('[data-lists-page="import"] form:has(input[type=file][name=file])')
        _send_file(form, helpers_sd.make_table(rows, "xlsx"), f"narrow-{tag}.xlsx", XLSX_MIME)
        expect(narrow.locator('[data-lists-page="mapping"]')).to_be_visible()
        assert helpers_lists.page_fits(narrow), "горизонтальная прокрутка на сопоставлении"
        _check(narrow)
        assert helpers_lists.page_fits(narrow), "горизонтальная прокрутка на предпросмотре"
        _confirm(narrow)
        assert helpers_lists.page_fits(narrow), "горизонтальная прокрутка на итоге"
    finally:
        context.close()
