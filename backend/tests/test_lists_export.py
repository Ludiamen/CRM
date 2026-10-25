"""T-6117: выгрузка списка в Excel — набор и столбцы как в таблице, типы ячеек, защита от формул, предел 10 000,
имя файла, скорость (FR-532, NFR-503). Сценарий 9 («Выгрузка в Excel…») спеки, пункты 1-4.

Выгрузку берёт настоящий маршрут `/lists/<id>/export/` (раздел 10.2 контракта); книгу читают `openpyxl`
(ячейки: тип, формат) и помощники `support_lists.read_xlsx`/`read_back`. Код среза на уровне модуля не
импортируется. Смотрит сотрудник с правом `lists.manage` (читает все списки); читатель без доступа к целевому
списку — в тестах нечитаемой ссылки и прав. Значения кладутся прямой записью в форме хранения раздела 2.3.
"""
import html as html_lib
import io
import re
import time
import zipfile
from decimal import Decimal
from urllib.parse import quote, unquote

import pytest
from django.utils import timezone

from tests import support, support_org
from tests import support_lists as lp

pytestmark = pytest.mark.django_db

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TITLE_HEADER = "Название"
UNAVAILABLE = "запись недоступного списка"


# --- помощники только этого файла ------------------------------------------------------------------------

def _env():
    """Управляющий (читает все списки) и пустой список его авторства."""
    manager = lp.as_manager()
    return manager, lp.make_list(created_by=manager.user)


def _use_columns(user, lst, columns):
    """Сохранённый выбор столбцов сотрудника (`UserView`) прямой записью."""
    return support.model("lists.UserView").objects.create(user=user, list=lst, columns=list(columns))


def _export(session, lst, params=None):
    return session.client.get(f"/lists/{lst.pk}/export/", params or {})


def _file(session, lst, params=None):
    """Успешная выгрузка: ответ с книгой XLSX."""
    response = _export(session, lst, params)
    assert response.status_code == 200, f"выгрузка: ожидался 200, получен {response.status_code}"
    assert response["Content-Type"].split(";")[0] == XLSX_TYPE, f"тип ответа: {response['Content-Type']!r}"
    return response


def _sheet(response):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(response.content)).worksheets[0]


def _cells(response) -> tuple[list, list[list]]:
    """Заголовок и строки первого листа."""
    table = lp.read_xlsx(response.content)
    return table[0], table[1:]


def _column(response, header: str) -> list:
    """Значения столбца по заголовку, строка за строкой."""
    head, rows = _cells(response)
    assert header in head, f"в выгрузке нет столбца {header!r}: {head!r}"
    index = head.index(header)
    return [row[index] for row in rows]


def _cell_of(response, row: int, header: str):
    """Ячейка `openpyxl`: `row` — номер строки данных с 1 (под заголовком)."""
    sheet = _sheet(response)
    head = [c.value for c in sheet[1]]
    assert header in head, f"в выгрузке нет столбца {header!r}: {head!r}"
    return sheet.cell(row=row + 1, column=head.index(header) + 1)


def _number(cell) -> Decimal:
    assert cell.data_type == "n", f"ячейка должна быть числовой, а тип {cell.data_type!r}, значение {cell.value!r}"
    return Decimal(str(cell.value))


def _bulk(lst, count, values_of, *, author):
    """Массовые записи `bulk_create` (решение 22 задач); `values_of(i)` — `values` записи номер `i`."""
    Record = support.model("lists.Record")
    now = timezone.now()
    rows = [Record(list=lst, number=i, title=f"Запись {i}", values=values_of(i), archived=False, version=1,
                   created_at=now, created_by=author, updated_at=now, updated_by=author)
            for i in range(1, count + 1)]
    Record.objects.bulk_create(rows, batch_size=1000)
    lst.next_number = count + 1
    lst.save(update_fields=["next_number"])


def _grid_numbers(session, lst, params=None) -> list[int]:
    """Номера записей строк таблицы `/lists/<id>/` по порядку."""
    response = session.client.get(f"/lists/{lst.pk}/", params or {})
    assert response.status_code == 200
    return [int(n) for n in re.findall(r'data-record-number="(\d+)"', response.content.decode())]


def _alerts(response, code: str) -> list[str]:
    """Теги страницы с `data-code=<код>` и `role="alert"`."""
    text = response.content.decode()
    return [tag for tag in re.findall(r"<[^>]*>", text)
            if f'data-code="{code}"' in tag and 'role="alert"' in tag]


# --- набор и столбцы (FR-532, FR-509) ----------------------------------------------------------------------

def test_export_uses_grid_preset_filters_query_and_sort_same_rows_and_order():
    """FR-532: выгрузка берёт отбор, фильтр поля, поиск и сортировку таблицы: строки и их порядок те же, что у
    таблицы с теми же параметрами (и такие, как ждёт спека), а запись из архива — только при отборе «Архив»."""
    manager, lst = _env()
    lp.add_field(lst, "string", "Статус", code="st", show_in_table=True)
    lp.add_field(lst, "number", "Очередь", code="qn", show_in_table=True)
    alpha = lp.make_record(lst, "Alpha zorblax", {"st": "open", "qn": "3"})
    beta = lp.make_record(lst, "Beta", {"st": "open", "qn": "1"})
    gamma = lp.make_record(lst, "Gamma zorblax", {"st": "closed", "qn": "2"})
    delta = lp.make_record(lst, "Delta", {"st": "open", "qn": "5"}, archived=True)
    epsilon = lp.make_record(lst, "Epsilon zorblax", {"st": "open", "qn": "4"})
    by_title = {r.title: r.number for r in (alpha, beta, gamma, delta, epsilon)}
    cases = [
        ({}, ["Alpha zorblax", "Beta", "Gamma zorblax", "Epsilon zorblax"]),
        ({"f.st": "OPEN", "sort": "qn", "dir": "desc"}, ["Epsilon zorblax", "Alpha zorblax", "Beta"]),
        ({"q": "zorblax"}, ["Alpha zorblax", "Gamma zorblax", "Epsilon zorblax"]),
        ({"preset": "archive"}, ["Delta"]),
        ({"f.st": "open", "q": "zorblax", "sort": "title", "dir": "desc"}, ["Epsilon zorblax", "Alpha zorblax"]),
    ]
    for params, expected in cases:
        titles = _column(_file(manager, lst, params), TITLE_HEADER)
        assert titles == expected, f"параметры {params}: в выгрузке {titles}, ждали {expected}"
        assert [by_title[t] for t in titles] == _grid_numbers(manager, lst, params), (
            f"параметры {params}: выгрузка и таблица показывают разные записи или порядок")


def test_export_columns_from_user_view_in_grid_order_headers_are_field_names():
    """FR-532: столбцы и их порядок — выбор сотрудника (`UserView`), без выбора — «Название» и поля «показывать в
    таблице» по порядку; заголовки — названия полей; значения под своими заголовками; чужой код столбца молча
    отбрасывается."""
    manager, lst = _env()
    lp.add_field(lst, "string", "Первое", code="a", show_in_table=True)
    lp.add_field(lst, "string", "Второе", code="b")
    lp.add_field(lst, "string", "Третье", code="c", show_in_table=True)
    lp.make_record(lst, "Запись один", {"a": "А1", "b": "Б1", "c": "В1"})
    lp.make_record(lst, "Запись два", {"a": "А2", "b": "Б2", "c": "В2"})
    _use_columns(manager.user, lst, ["title", "c", "zz", "b", "a"])
    head, rows = _cells(_file(manager, lst))
    assert head == [TITLE_HEADER, "Третье", "Второе", "Первое"]
    assert rows == [["Запись один", "В1", "Б1", "А1"], ["Запись два", "В2", "Б2", "А2"]]
    other = lp.as_manager()
    head, rows = _cells(_file(other, lst))
    assert head == [TITLE_HEADER, "Первое", "Третье"]
    assert rows == [["Запись один", "А1", "В1"], ["Запись два", "А2", "В2"]]


def test_export_hidden_field_not_exported_even_if_in_user_view():
    """FR-509, FR-532: скрытое поле не выгружается, даже если код стоит в сохранённом выборе столбцов и у поля
    «показывать в таблице»; его значения в файле не встречаются нигде."""
    manager, lst = _env()
    lp.add_field(lst, "string", "Видимое", code="shown", show_in_table=True)
    lp.add_field(lst, "string", "Скрытое", code="hid", show_in_table=True, hidden=True)
    lp.make_record(lst, "Запись", {"shown": "видно", "hid": "тайная-метка-zq"})
    _use_columns(manager.user, lst, ["title", "hid", "shown"])
    for params in ({}, {"sort": "hid"}):
        response = _file(manager, lst, params)
        head, rows = _cells(response)
        assert head == [TITLE_HEADER, "Видимое"]
        assert all("тайная-метка-zq" not in str(value) for row in rows for value in row)
    assert "тайная-метка-zq" not in " ".join(str(c.value) for row in _sheet(response).iter_rows() for c in row)


# --- типы ячеек (FR-532) -----------------------------------------------------------------------------------

def test_export_number_is_numeric_decimal_cell():
    """FR-532: число — числовая ячейка со значением как у `Decimal` (не текст, без апострофа), в том числе
    отрицательное и дробное."""
    manager, lst = _env()
    lp.add_field(lst, "number", "Количество", code="amount", show_in_table=True)
    for title, value in (("a", "12.5"), ("b", "100"), ("c", "-3"), ("d", "0.125")):
        lp.make_record(lst, title, {"amount": value})
    response = _file(manager, lst)
    got = [_number(_cell_of(response, row, "Количество")) for row in range(1, 5)]
    assert got == [Decimal("12.5"), Decimal("100"), Decimal("-3"), Decimal("0.125")]


def test_export_money_is_numeric_cell_with_currency_format_rub_usd_eur():
    """FR-532: деньги — числовая ячейка (не текст) с форматом валюты поля: `#,##0.00 "₽"`, `"$"`, `"€"`."""
    manager, lst = _env()
    lp.add_field(lst, "money", "Рубли", code="rub", show_in_table=True, currency="RUB")
    lp.add_field(lst, "money", "Доллары", code="usd", show_in_table=True, currency="USD")
    lp.add_field(lst, "money", "Евро", code="eur", show_in_table=True, currency="EUR")
    lp.make_record(lst, "Платёж", {"rub": "1234.50", "usd": "10.00", "eur": "0.50"})
    response = _file(manager, lst)
    for header, value, mark in (("Рубли", "1234.5", "₽"), ("Доллары", "10", "$"), ("Евро", "0.5", "€")):
        cell = _cell_of(response, 1, header)
        assert _number(cell) == Decimal(value)
        assert cell.number_format == f'#,##0.00 "{mark}"', f"{header}: формат {cell.number_format!r}"


def test_export_date_is_date_cell_with_dd_mm_yyyy_format():
    """FR-532: дата — ячейка даты (не текст) с форматом `DD.MM.YYYY`; пустая дата — пустая ячейка."""
    manager, lst = _env()
    lp.add_field(lst, "date", "Покупка", code="bought", show_in_table=True)
    lp.make_record(lst, "С датой", {"bought": "2026-10-25"})
    lp.make_record(lst, "Без даты", {})
    response = _file(manager, lst)
    cell = _cell_of(response, 1, "Покупка")
    assert cell.is_date, f"ячейка не дата: тип {cell.data_type!r}, значение {cell.value!r}"
    assert cell.value.date().isoformat() == "2026-10-25"
    assert cell.number_format == "DD.MM.YYYY"
    assert _cell_of(response, 2, "Покупка").value in (None, "")


def test_export_datetime_is_moscow_text_dd_mm_yyyy_hh_mm():
    """FR-532: дата и время — текст «ДД.ММ.ГГГГ ЧЧ:ММ» по Москве (хранится UTC): 21:30 UTC 25.10 — это 00:30 26.10,
    07:05 UTC — 10:05; ячейка текстовая (так файл загружается обратно)."""
    manager, lst = _env()
    lp.add_field(lst, "datetime", "Когда", code="moment", show_in_table=True)
    lp.make_record(lst, "Ночь", {"moment": "2026-10-25T21:30:00Z"})
    lp.make_record(lst, "Утро", {"moment": "2026-01-15T07:05:00Z"})
    response = _file(manager, lst)
    for row, text in ((1, "26.10.2026 00:30"), (2, "15.01.2026 10:05")):
        cell = _cell_of(response, row, "Когда")
        assert cell.data_type == "s" and cell.value == text, f"строка {row}: {cell.data_type!r} {cell.value!r}"


def test_export_bool_employee_unit_and_record_shown_as_text():
    """FR-532: да/нет — «Да»; сотрудник — ФИО; подразделение — название узла (не путь); запись другого списка —
    её название; всё текстом."""
    manager, lst = _env()
    target = lp.make_list(created_by=manager.user)
    target_record = lp.make_record(target, "Принтер из кабинета 5")
    lp.add_field(lst, "bool", "Исправно", code="ok", show_in_table=True)
    lp.add_field(lst, "person", "Ответственный", code="who", show_in_table=True)
    lp.add_field(lst, "unit", "Отдел", code="dep", show_in_table=True)
    lp.add_field(lst, "record", "Устройство", code="dev", show_in_table=True, target=target)
    person = support.make_employee(full_name="Воробьёв Игорь Петрович")
    unit = support_org.make_unit("Отдел связи")
    lp.make_record(lst, "С значениями", {"ok": True, "who": [person.pk], "dep": [unit.pk], "dev": [target_record.pk]})
    lp.make_record(lst, "Без значений", {})
    response = _file(manager, lst)
    assert _cell_of(response, 1, "Исправно").value == "Да"
    assert _cell_of(response, 2, "Исправно").value in ("Нет", None, "")
    assert _cell_of(response, 1, "Ответственный").value == "Воробьёв Игорь Петрович"
    assert _cell_of(response, 1, "Отдел").value == "Отдел связи"
    assert _cell_of(response, 1, "Устройство").value == "Принтер из кабинета 5"
    for header in ("Ответственный", "Отдел", "Устройство"):
        assert _cell_of(response, 2, header).value in (None, "")
        assert _cell_of(response, 1, header).data_type == "s"


def test_export_file_names_and_multiple_values_joined_with_semicolon():
    """FR-532: имена файлов и несколько значений (список значений, сотрудники, подразделения, записи) — через
    «; » в одной ячейке; состав — все значения, ничего лишнего."""
    manager, lst = _env()
    target = lp.make_list(created_by=manager.user)
    first_target = lp.make_record(target, "Цель один")
    second_target = lp.make_record(target, "Цель два")
    files = lp.add_field(lst, "file", "Файлы", code="docs", multiple=True, show_in_table=True)
    lp.add_field(lst, "choice", "Метки", code="tags", multiple=True, show_in_table=True, choices=("Красная", "Синяя"))
    lp.add_field(lst, "person", "Люди", code="people", multiple=True, show_in_table=True)
    lp.add_field(lst, "unit", "Отделы", code="deps", multiple=True, show_in_table=True)
    lp.add_field(lst, "record", "Цели", code="goals", multiple=True, show_in_table=True, target=target)
    tags = support.model("lists.Field").objects.get(list=lst, code="tags")
    people = [support.make_employee(full_name="Борисов Борис"), support.make_employee(full_name="Андреев Андрей")]
    deps = [support_org.make_unit("Сеть"), support_org.make_unit("Склад")]
    record = lp.make_record(lst, "Много значений", {
        "tags": [lp.choice_id(tags, "Красная"), lp.choice_id(tags, "Синяя")], "people": [p.pk for p in people],
        "deps": [d.pk for d in deps], "goals": [first_target.pk, second_target.pk]})
    for name in ("акт.pdf", "счёт.pdf", "фото.png"):
        lp.attach_file(record, files, name)
    response = _file(manager, lst)
    expected = {"Файлы": {"акт.pdf", "счёт.pdf", "фото.png"}, "Метки": {"Красная", "Синяя"},
                "Люди": {"Борисов Борис", "Андреев Андрей"}, "Отделы": {"Сеть", "Склад"},
                "Цели": {"Цель один", "Цель два"}}
    for header, names in expected.items():
        text = _cell_of(response, 1, header).value
        assert isinstance(text, str) and text.count("; ") == len(names) - 1, f"{header}: {text!r}"
        assert set(text.split("; ")) == names, f"{header}: {text!r}"


# --- защита от формул (FR-532) -----------------------------------------------------------------------------

FORMULA_TEXTS = ["=1+1", "+7 999", "-5", "@x", "\tA", "\rB"]


def _after_reading(text: str) -> str:
    """Апостроф впереди; `\\r` после чтения книги — `\\n` (факт T-6005: XML нормализует возврат каретки)."""
    return "'" + text.replace("\r", "\n")


def test_export_formula_prefixes_in_values_get_apostrophe_string_cell():
    """FR-532: текст, начинающийся с `=`, `+`, `-`, `@`, табуляции или возврата каретки, — в строковой ячейке с
    апострофом впереди (название записи и поле «строка»); обычный текст и текст с апострофом не меняются."""
    manager, lst = _env()
    lp.add_field(lst, "string", "Заметка", code="note", show_in_table=True)
    for text in FORMULA_TEXTS:
        lp.make_record(lst, text, {"note": text})
    lp.make_record(lst, "Обычная", {"note": "Принтер"})
    lp.make_record(lst, "'уже апостроф", {"note": "'уже апостроф"})
    response = _file(manager, lst)
    for row, text in enumerate(FORMULA_TEXTS, start=1):
        for header in (TITLE_HEADER, "Заметка"):
            cell = _cell_of(response, row, header)
            assert cell.data_type == "s", f"{text!r}/{header}: тип {cell.data_type!r}"
            assert cell.value == _after_reading(text), f"{text!r}/{header}: {cell.value!r}"
    for header in (TITLE_HEADER, "Заметка"):
        assert _cell_of(response, 7, header).value in ("Обычная", "Принтер")
        assert _cell_of(response, 8, header).value == "'уже апостроф"


def test_export_formula_prefixes_in_headers_get_apostrophe():
    """FR-532: защита действует и на заголовки: названия полей с опасным первым знаком — с апострофом, обычные —
    без него."""
    manager, lst = _env()
    names = ["=SUM(A1)", "+плюс", "-минус", "@имя", "Обычное"]
    for index, name in enumerate(names):
        lp.add_field(lst, "string", name, code=f"h{index}", show_in_table=True)
    lp.make_record(lst, "Запись", {})
    head, _ = _cells(_file(manager, lst))
    assert head == [TITLE_HEADER, "'=SUM(A1)", "'+плюс", "'-минус", "'@имя", "Обычное"]
    sheet = _sheet(_file(manager, lst))
    assert all(c.data_type == "s" for c in sheet[1])


def test_export_negative_numbers_money_and_dates_not_escaped():
    """FR-532: числа, деньги и даты формул не образуют: отрицательное число и сумма остаются числовыми ячейками
    без апострофа, дата — ячейкой даты."""
    manager, lst = _env()
    lp.add_field(lst, "number", "Остаток", code="rest", show_in_table=True)
    lp.add_field(lst, "money", "Долг", code="debt", show_in_table=True)
    lp.add_field(lst, "date", "Срок", code="until", show_in_table=True)
    lp.make_record(lst, "Минус", {"rest": "-5", "debt": "-100.50", "until": "2026-01-05"})
    response = _file(manager, lst)
    assert _number(_cell_of(response, 1, "Остаток")) == Decimal("-5")
    assert _number(_cell_of(response, 1, "Долг")) == Decimal("-100.5")
    date_cell = _cell_of(response, 1, "Срок")
    assert date_cell.is_date and date_cell.value.date().isoformat() == "2026-01-05"
    for header in ("Остаток", "Долг", "Срок"):
        assert not str(_cell_of(response, 1, header).value).startswith("'")


def test_export_workbook_has_no_formula_cells():
    """FR-532: в книге нет ни одной формулы: ни ячеек типа `f`, ни тегов `<f>` в XML листа — при значениях,
    начинающихся с `=`, `+`, `-`, `@`, и в названиях полей."""
    manager, lst = _env()
    lp.add_field(lst, "string", "=Поле", code="fx", show_in_table=True)
    lp.add_field(lst, "number", "Число", code="num", show_in_table=True)
    for text in FORMULA_TEXTS:
        lp.make_record(lst, "=" + text, {"fx": text, "num": "-1"})
    response = _file(manager, lst)
    sheet = _sheet(response)
    assert all(c.data_type != "f" for row in sheet.iter_rows() for c in row)
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        sheets = [archive.read(n).decode() for n in archive.namelist() if n.startswith("xl/worksheets/")]
    assert sheets and all("<f>" not in xml and "<f " not in xml for xml in sheets)


# --- загрузка выгрузки обратно (FR-532, FR-533) --------------------------------------------------------------

def test_export_uploaded_back_through_import_preview_has_no_row_errors_for_number_money_1234_5_date_datetime_bool_unit_employee_and_apostrophe_values():
    """FR-532, FR-533: выгруженный файл загружается обратно без ошибок строк: число, деньги «1234.5» из числовой
    ячейки, дата, дата и время, да/нет, подразделение, сотрудник, список значений, запись другого списка и значения
    с апострофом (его загрузка снимает)."""
    from apps.lists.services import imports

    manager, lst = _env()
    target = lp.make_list(created_by=manager.user)
    target_record = lp.make_record(target, "Цель выгрузки")
    lp.add_field(lst, "string", "Заметка", code="note", show_in_table=True)
    lp.add_field(lst, "number", "Количество", code="amount", show_in_table=True)
    lp.add_field(lst, "money", "Сумма", code="sum", show_in_table=True)
    lp.add_field(lst, "date", "Дата покупки", code="bought", show_in_table=True)
    lp.add_field(lst, "datetime", "Момент", code="moment", show_in_table=True)
    lp.add_field(lst, "bool", "Исправно", code="ok", show_in_table=True)
    lp.add_field(lst, "unit", "Отдел", code="dep", show_in_table=True)
    lp.add_field(lst, "person", "Ответственный", code="who", show_in_table=True)
    lp.add_field(lst, "choice", "Вид", code="kind", show_in_table=True, choices=("Первый", "Второй"))
    lp.add_field(lst, "record", "Связь", code="link", show_in_table=True, target=target)
    kind = support.model("lists.Field").objects.get(list=lst, code="kind")
    person = support.make_employee(full_name="Кузнецов Уникальный Сергеевич")
    unit = support_org.make_unit("Отдел обратной загрузки")
    full = {"note": "=1+1", "amount": "12.5", "sum": "1234.50", "bought": "2026-10-25",
            "moment": "2026-10-25T11:30:00Z", "ok": True, "dep": [unit.pk], "who": [person.pk],
            "kind": [lp.choice_id(kind, "Второй")], "link": [target_record.pk]}
    lp.make_record(lst, "=формула", full)
    lp.make_record(lst, "+7 999", {"note": "-5", "amount": "-3", "sum": "0.50"})
    lp.make_record(lst, "Пустая", {})
    response = _file(manager, lst)
    job = imports.upload(lst, manager.user, name="выгрузка.xlsx", content=response.content)
    job = imports.preview(job, imports.auto_mapping(lst, job.header), manager.user)
    job.refresh_from_db()
    bad = [(row["row"], row["reason"]) for row in job.rows if row["result"] == "error"]
    assert not bad, f"строки с ошибками при загрузке выгрузки обратно: {bad}"
    assert imports.counts(job) == {"will_create": 3, "will_update": 0, "error": 0}


def test_export_record_field_value_pointing_to_unreadable_list_is_shown_as_unavailable_list_record_without_title():
    """FR-527, FR-532: ссылка на запись списка, которого смотрящий не читает, — «запись недоступного списка» без
    названия записи (решение 57 задач); читающий оба списка видит название; названия закрытой записи нет в файле."""
    manager, lst = _env()
    target = lp.make_list(created_by=manager.user)
    hidden_record = lp.make_record(target, "Закрытое-название-zq")
    lp.add_field(lst, "record", "Устройство", code="dev", show_in_table=True, target=target)
    lp.make_record(lst, "Запись", {"dev": [hidden_record.pk]})
    reader = lp.as_reader(lst)
    response = _file(reader, lst)
    cell = _cell_of(response, 1, "Устройство")
    assert cell.data_type == "s" and cell.value.casefold() == UNAVAILABLE
    assert all("Закрытое-название-zq" not in str(c.value) for row in _sheet(response).iter_rows() for c in row)
    assert b"Closed" not in response.content and "Закрытое-название-zq".encode() not in response.content
    assert _cell_of(_file(manager, lst), 1, "Устройство").value == "Закрытое-название-zq"


# --- предел 10 000 (FR-532, FR-546) --------------------------------------------------------------------------

def test_export_over_10000_records_refused_with_message_on_grid_page():
    """FR-532, FR-546: 10 001 запись в отборе — не файл, а 200 со страницей таблицы и сообщением
    `[role=alert][data-code="export_too_many"]` «Уточните отбор: записей больше 10 000»; служба даёт отказ
    `ListsError("export_too_many")`."""
    from django.http import QueryDict

    from apps.lists.services import export
    from apps.lists.services.errors import ListsError

    manager, lst = _env()
    _bulk(lst, 10001, lambda i: {}, author=manager.user)
    response = _export(manager, lst)
    assert response.status_code == 200
    assert response["Content-Type"].split(";")[0] != XLSX_TYPE
    assert "attachment" not in response.get("Content-Disposition", "")
    assert len(_alerts(response, "export_too_many")) == 1, "нет сообщения `export_too_many` с role=alert"
    assert "Уточните отбор: записей больше 10 000" in html_lib.unescape(response.content.decode())
    assert 'data-lists-page="grid"' in response.content.decode()
    with pytest.raises(ListsError) as raised:
        export.build(lst, manager.user, QueryDict())
    assert raised.value.code == "export_too_many"


def test_export_exactly_10000_records_returns_file():
    """FR-532: ровно 10 000 записей — предел включительно: приходит файл, в нём заголовок и 10 000 строк."""
    manager, lst = _env()
    lp.add_field(lst, "number", "Номер по порядку", code="seq", show_in_table=True)
    _bulk(lst, 10000, lambda i: {"seq": str(i)}, author=manager.user)
    response = _file(manager, lst)
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode()
    assert len(re.findall(r"<row[ >]", xml)) == 10001
    sheet = _sheet(response)
    assert sheet.max_row == 10001
    assert sheet.cell(row=2, column=1).value == "Запись 1"
    assert sheet.cell(row=10001, column=1).value == "Запись 10000"


def test_export_limit_counts_filtered_set_not_whole_list():
    """FR-532: предел считается по отбору, а не по всему списку: в списке 10 003 записи — без отбора отказ, с
    фильтром на три записи — файл с тремя строками."""
    manager, lst = _env()
    lp.add_field(lst, "string", "Метка", code="mark", show_in_table=True)
    _bulk(lst, 10003, lambda i: {"mark": "needle" if i in (7, 500, 10003) else "plain"}, author=manager.user)
    refused = _export(manager, lst)
    assert refused.status_code == 200 and len(_alerts(refused, "export_too_many")) == 1
    response = _file(manager, lst, {"f.mark": "needle"})
    assert _column(response, TITLE_HEADER) == ["Запись 7", "Запись 500", "Запись 10003"]


# --- имя файла, права, журнал (FR-532, FR-514, FR-542) --------------------------------------------------------

def test_export_xlsx_attachment_with_cyrillic_filename_star():
    """FR-532: ответ — вложение XLSX; имя — `filename*=UTF-8''<название списка в процентной записи>.xlsx`, кириллица
    только закодированной (заголовок ASCII)."""
    manager = lp.as_manager()
    name = "Парк оборудования 2026"
    lst = lp.make_list(name, created_by=manager.user)
    lp.make_record(lst, "Запись")
    response = _file(manager, lst)
    header = response["Content-Disposition"]
    header.encode("ascii")
    assert header.startswith("attachment;")
    found = re.search(r"filename\*=UTF-8''([^;]+)", header)
    assert found, f"нет `filename*=UTF-8''…`: {header!r}"
    assert "%D0" in found.group(1)
    assert unquote(found.group(1)) == name + ".xlsx"


def test_export_for_non_reader_returns_404_same_as_missing_list():
    """FR-532, FR-514: не читающий список (в том числе читающий другой список) получает 404, неотличимый от ответа на
    несуществующий список; читающий — 200 с файлом."""
    manager, lst = _env()
    lp.make_record(lst, "Секретное-название-zq")
    other = lp.make_list(created_by=manager.user)
    outsider = lp.as_reader(other)
    stranger = support.logged_in(support.make_employee())
    # Ответы разных сотрудников различаются меню оболочки, поэтому закрытый список сравнивается с несуществующим
    # в пределах одного клиента (CONTRACT-006, раздел 15.3, К-10).
    for client in (outsider.client, stranger):
        missing = client.get("/lists/987654321/export/")
        closed = client.get(f"/lists/{lst.pk}/export/")
        assert missing.status_code == 404 and closed.status_code == 404
        assert "Секретное-название-zq" not in closed.content.decode()
        assert support.same_failure(closed, missing), "ответ 404 на закрытый список отличим от ответа на несуществующий"
    assert _file(lp.as_reader(lst), lst).status_code == 200


def test_export_not_written_to_audit_log():
    """FR-542: выгрузка (успешная и отказанная по пределу) не пишет в журнал портала и не создаёт событий истории
    записей."""
    manager, lst = _env()
    _bulk(lst, 3, lambda i: {}, author=manager.user)
    Event, Change = support.model("core.AuditEvent"), support.model("lists.Change")
    before = (Event.objects.count(), Change.objects.count())
    _file(manager, lst)
    big = lp.make_list(created_by=manager.user)
    _bulk(big, 10001, lambda i: {}, author=manager.user)
    assert len(_alerts(_export(manager, big), "export_too_many")) == 1
    assert (Event.objects.count(), Change.objects.count()) == before
    assert not Event.objects.filter(action__startswith="lists.").filter(action__icontains="export").exists()


# --- скорость (NFR-503) -----------------------------------------------------------------------------------

def test_export_10000_rows_of_equipment_fields_within_10_seconds():
    """NFR-503: 10 000 записей × 15 столбцов (поля оборудования всех основных типов) выгружаются службой за не более
    10 с (лучший из трёх запусков; данные `bulk_create` в замер не входят, решение 22 задач)."""
    from django.http import QueryDict

    from apps.lists.services import export

    manager = lp.as_manager()
    target = lp.make_list(created_by=manager.user)
    devices = [lp.make_record(target, f"Устройство {i}") for i in range(1, 4)]
    lst = lp.make_list("Оборудование для замера", created_by=manager.user)
    kinds = lp.add_field(lst, "choice", "Вид", code="kind", choices=("Компьютер", "Ноутбук", "Монитор"))
    lp.add_field(lst, "string", "Инвентарный номер", code="inv")
    lp.add_field(lst, "string", "Серийный номер", code="ser")
    lp.add_field(lst, "string", "Модель", code="mod")
    lp.add_field(lst, "choice", "Состояние", code="state", choices=("В работе", "На складе"))
    lp.add_field(lst, "person", "За кем закреплено", code="who")
    lp.add_field(lst, "unit", "Подразделение", code="dep")
    lp.add_field(lst, "string", "Место", code="place")
    lp.add_field(lst, "date", "Дата покупки", code="bought")
    lp.add_field(lst, "date", "Гарантия до", code="warranty")
    lp.add_field(lst, "money", "Стоимость", code="cost")
    lp.add_field(lst, "text", "Примечание", code="note")
    lp.add_field(lst, "record", "Связанное", code="dev", target=target)
    lp.add_field(lst, "string", "Прочее", code="extra")
    codes = ["kind", "inv", "ser", "mod", "state", "who", "dep", "place", "bought", "warranty", "cost", "note", "dev",
             "extra"]
    _use_columns(manager.user, lst, ["title"] + codes)
    state = support.model("lists.Field").objects.get(list=lst, code="state")
    people = [support.make_employee().pk for _ in range(3)]
    units = [support_org.make_unit().pk for _ in range(3)]
    kind_ids = [lp.choice_id(kinds, t) for t in ("Компьютер", "Ноутбук", "Монитор")]
    state_ids = [lp.choice_id(state, t) for t in ("В работе", "На складе")]
    _bulk(lst, 10000, lambda i: {
        "kind": [kind_ids[i % 3]], "inv": f"ИН-{i:05d}", "ser": f"SN{i * 7919:09d}", "mod": f"Модель {i % 50}",
        "state": [state_ids[i % 2]], "who": [people[i % 3]], "dep": [units[i % 3]], "place": f"Кабинет {i % 40}",
        "bought": "2025-03-14", "warranty": "2028-03-14", "cost": f"{i}.50", "note": "Строка примечания\nвторая",
        "dev": [devices[i % 3].pk], "extra": f"Прочее {i}"}, author=manager.user)
    best, response = None, None
    for _ in range(3):
        started = time.perf_counter()
        response = export.build(lst, manager.user, QueryDict())
        content = response.content if hasattr(response, "content") else b"".join(response.streaming_content)
        spent = time.perf_counter() - started
        best = spent if best is None else min(best, spent)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode()
    assert len(re.findall(r"<row[ >]", xml)) == 10001
    assert best <= 10, f"выгрузка 10 000 × 15: лучший из трёх запусков {best:.2f} с, бюджет 10 с"
