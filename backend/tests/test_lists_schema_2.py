"""T-6104: изменение схемы списка с данными — поле с данными, замок схемы, пустое поле, скрыть и показать, значения
списка значений (FR-509, FR-510, FR-511, FR-515, FR-516, FR-542, FR-546). Сценарии 14 и 15 спеки среза 006.

Код среза 006 импортируется внутри тестов; модели берутся через `support.model`. Разметку разбирает `html.parser`,
элементы ищутся по зацепкам `data-*` (CONTRACT-006, раздел 11). Отказы службы — `ListsError` с `code` и `field`
(раздел 5.1), отказы формы — 200 с `[data-error][data-code]` (раздел 10.1). Перечни мест V1–V20 и A1–A9 — раздел 6."""
import io
import re
from datetime import timedelta
from html.parser import HTMLParser

import pytest
from django.contrib.postgres.search import SearchQuery
from django.http import QueryDict
from django.utils import timezone

from tests import factories, support
from tests import support_lists as lists
from tests.support_notify import run_jobs

pytestmark = pytest.mark.django_db

_VOID = {"input", "br", "img", "meta", "link", "hr", "area", "base", "col", "embed", "source", "track", "wbr"}


# --- разбор разметки и вызовы (только для этого файла) -----------------------------------------------------

class _El:
    def __init__(self, tag, attrs, parent):
        self.tag, self.attrs, self.parent, self.raw = tag, attrs, parent, ""

    @property
    def text(self) -> str:
        return " ".join(self.raw.split())


class _Doc(HTMLParser):
    def __init__(self, html: str):
        super().__init__(convert_charrefs=True)
        self.elements, self._stack = [], []
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        el = _El(tag, {name: (value if value is not None else "") for name, value in attrs},
                 self._stack[-1] if self._stack else None)
        self.elements.append(el)
        if tag not in _VOID:
            self._stack.append(el)

    def handle_startendtag(self, tag, attrs):
        self.elements.append(_El(tag, {name: (value if value is not None else "") for name, value in attrs},
                                 self._stack[-1] if self._stack else None))

    def handle_endtag(self, tag):
        for i in range(len(self._stack) - 1, -1, -1):
            if self._stack[i].tag == tag:
                del self._stack[i:]
                break

    def handle_data(self, data):
        for el in self._stack:
            el.raw += data


def _doc(response) -> _Doc:
    return _Doc(response.content.decode("utf-8"))


def _all(doc, tag=None, attrs=None, within=None) -> list:
    """Элементы с тегом `tag` и атрибутами `attrs` (значение `True` — достаточно наличия); `within` — предок."""
    found = []
    for el in doc.elements:
        if tag and el.tag != tag:
            continue
        if any((name not in el.attrs) if value is True else (el.attrs.get(name) != value)
               for name, value in (attrs or {}).items()):
            continue
        if within is not None and not _inside(el, within):
            continue
        found.append(el)
    return found


def _inside(el, ancestor) -> bool:
    node = el
    while node is not None:
        if node is ancestor:
            return True
        node = node.parent
    return False


def _row(doc, code):
    rows = _all(doc, "tr", {"data-field-code": code})
    return rows[0] if rows else None


def _buttons(doc, row) -> set:
    return {el.attrs.get("data-action") for el in doc.elements
            if el.tag in ("button", "a") and "data-action" in el.attrs and _inside(el, row)}


def _body(response) -> str:
    return response.content.decode("utf-8")


def _raises(call, code, field=None):
    from apps.lists.services.errors import ListsError

    with pytest.raises(ListsError) as caught:
        call()
    assert caught.value.code == code, (caught.value.code, code)
    if field is not None:
        assert caught.value.field == field, (caught.value.field, field)
    return caught.value


def _schema():
    from apps.lists.services import schema

    return schema


def _events(action: str, obj=None) -> list:
    query = support.model("core.AuditEvent").objects.filter(action=action)
    if obj is not None:
        query = query.filter(object_id=str(obj.pk))
    return list(query.order_by("id"))


def _flag(value) -> bool:
    """`hidden` в журнале — булево JSON (ответ 57, CONTRACT-006, 8.1, 15.3), не строка."""
    assert value is None or isinstance(value, bool), f"hidden — булево JSON, получено {value!r}"
    return value is True


def _world(*, with_manager: bool = True):
    manager = lists.as_manager()
    return manager, lists.make_list(created_by=manager.user)


def _fields_url(lst) -> str:
    return f"/lists/{lst.pk}/settings/fields/"


def _field_url(lst, field) -> str:
    return f"/lists/{lst.pk}/settings/fields/{field.code}/"


def _fresh(obj):
    return type(obj).objects.get(pk=obj.pk)


def _change(record, kind, code, name, before, after, actor):
    return support.model("lists.Change").objects.create(
        record=record, kind=kind, actor=actor, at=timezone.now(),
        changes=[{"field": code, "name": name, "before": before, "after": after}])


def _vector_matches(record, word: str) -> bool:
    Record = support.model("lists.Record")
    return Record.objects.filter(pk=record.pk, search=SearchQuery(word, config="russian")).exists()


# --- поле с данными и пустое поле (FR-509) ------------------------------------------------------------------

@pytest.mark.parametrize("case", ["live_value", "archived_value", "history_only", "file_removed_in_history_only",
                                  "file_only"])
def test_field_with_data_in_record_archived_record_history_only_removed_file_in_history_or_file_only_is_not_deleted_or_retyped(case):
    """FR-509: поле с данными — значение в записи, в записи архива, только в истории, только удалённый файл в
    истории или только файл — не удаляется (`field_has_data`) и не меняет тип (`field_locked` у поля `type`);
    поле без данных рядом помечено пустым; отказ не пишет журнал."""
    schema = _schema()
    manager, lst = _world()
    is_file = case.startswith("file")
    field = lists.add_field(lst, "file" if is_file else "string", "Проверяемое", multiple=is_file)
    empty = lists.add_field(lst, "string", "Пустое")
    values = {field.code: "значение"} if case in ("live_value", "archived_value") else {}
    record = lists.make_record(lst, "Запись", values, archived=(case == "archived_value"))
    if case == "history_only":
        _change(record, "changed", field.code, field.name, "значение", "", manager.user)
    if case == "file_removed_in_history_only":
        _change(record, "file_removed", field.code, field.name, "акт.pdf", "", manager.user)
    if case == "file_only":
        lists.attach_file(record, field)
    assert schema.has_data(field) is True
    assert schema.has_data(empty) is False
    _raises(lambda: schema.delete_field(field, actor=manager.user), "field_has_data")
    new_type = "number" if not is_file else "string"
    _raises(lambda: schema.update_field(field, actor=manager.user, type=new_type), "field_locked", "type")
    field.refresh_from_db()
    assert field.type == ("file" if is_file else "string")
    assert support.model("lists.Field").objects.filter(pk=field.pk).exists()
    assert not _events("lists.field_deleted") and not _events("lists.field_changed")


def test_fields_tab_shows_hide_without_delete_and_type_for_field_with_data():
    """FR-509: вкладка «Поля» — у поля с данными есть «Скрыть», нет «Удалить»; у скрытого вместо «Скрыть» —
    «Показать»; у пустого есть «Удалить» с диалогом подтверждения; у системного «Названия» кнопок нет; в форме поля
    с данными нет выбора `type`, в форме пустого он есть."""
    manager, lst = _world()
    with_data = lists.add_field(lst, "string", "С данными")
    hidden = lists.add_field(lst, "string", "Скрытое", hidden=True)
    empty = lists.add_field(lst, "number", "Пустое")
    lists.make_record(lst, "Запись", {with_data.code: "значение", hidden.code: "тайна"})
    response = manager.client.get(_fields_url(lst))
    assert response.status_code == 200
    doc = _doc(response)
    row = _row(doc, with_data.code)
    assert row is not None
    assert any(el.attrs.get("data-has-data") == "true" for el in doc.elements if _inside(el, row))
    assert "field-hide" in _buttons(doc, row) and "field-delete" not in _buttons(doc, row)
    assert "Строка" in " ".join(el.text for el in doc.elements if el.attrs.get("data-field") == "type"
                                and _inside(el, row))
    hidden_row = _row(doc, hidden.code)
    assert any(el.attrs.get("data-hidden") == "true" for el in doc.elements if _inside(el, hidden_row))
    assert "field-show" in _buttons(doc, hidden_row) and "field-hide" not in _buttons(doc, hidden_row)
    assert "field-delete" not in _buttons(doc, hidden_row)
    empty_row = _row(doc, empty.code)
    assert any(el.attrs.get("data-has-data") == "false" for el in doc.elements if _inside(el, empty_row))
    assert "field-delete" in _buttons(doc, empty_row)
    assert _all(doc, "dialog", {"data-lists-confirm": True})
    title_row = _all(doc, "tr", {"data-field-code": "title", "data-system": True})
    assert title_row and not _buttons(doc, title_row[0])
    form_data = _doc(manager.client.get(_field_url(lst, with_data)))
    assert _all(form_data, "form", {"data-lists-field-form": True})
    assert not _all(form_data, None, {"name": "type"})
    form_empty = _doc(manager.client.get(_field_url(lst, empty)))
    assert _all(form_empty, None, {"name": "type"})


def test_empty_field_deletable_despite_same_code_with_data_in_other_list_and_bool_without_ticks_is_empty():
    """FR-509: код `f1` с данными в другом списке не делает пустое поле списка «полем с данными»; поле «да/нет», у
    записей которого нет галочек (ключа нет), пусто и удаляется; с галочкой — нет."""
    schema = _schema()
    manager, lst = _world()
    other = lists.make_list(created_by=manager.user)
    # Коды не задаются вручную: явный `code` не двигает `next_field_no` (support_lists.add_field), а в двух свежих
    # списках автоматические коды первых полей и так совпадают.
    mine = lists.add_field(lst, "string", "Моё")
    foreign = lists.add_field(other, "string", "Чужое")
    assert mine.code == foreign.code == "f1"
    foreign_record = lists.make_record(other, "Чужая запись", {foreign.code: "данные"})
    _change(foreign_record, "changed", "f1", "Чужое", "а", "б", manager.user)
    lists.make_record(lst, "Запись без значения")
    assert schema.has_data(mine) is False and schema.has_data(foreign) is True
    flag = lists.add_field(lst, "bool", "Флажок")
    ticked = lists.add_field(lst, "bool", "С галочкой")
    lists.make_record(lst, "Без галочек", {})
    lists.make_record(lst, "С галочкой", {ticked.code: True})
    assert schema.has_data(flag) is False and schema.has_data(ticked) is True
    schema.delete_field(mine, actor=manager.user)
    schema.delete_field(flag, actor=manager.user)
    Field = support.model("lists.Field")
    assert not Field.objects.filter(pk__in=[mine.pk, flag.pk]).exists()
    assert Field.objects.filter(pk=foreign.pk).exists()
    _raises(lambda: schema.delete_field(ticked, actor=manager.user), "field_has_data")


def test_empty_field_deleted_after_confirmation_writes_field_deleted():
    """FR-509, FR-542: пустое поле удаляется после подтверждения — POST на `…/delete/` даёт 302 на вкладку «Поля»,
    строки `Field` нет, в журнале одна запись `lists.field_deleted` с `before` из `list`, `field`, `type`."""
    manager, lst = _world()
    lst.name = "Список для удаления поля"
    lst.save(update_fields=["name"])
    field = lists.add_field(lst, "string", "Лишнее")
    keep = lists.add_field(lst, "number", "Остаётся")
    page = _doc(manager.client.get(_fields_url(lst)))
    assert "field-delete" in _buttons(page, _row(page, field.code))
    assert _all(page, "dialog", {"data-lists-confirm": True})
    pk = field.pk
    response = manager.client.post(_field_url(lst, field) + "delete/")
    assert response.status_code == 302 and response["Location"].endswith(_fields_url(lst))
    Field = support.model("lists.Field")
    assert not Field.objects.filter(pk=pk).exists() and Field.objects.filter(pk=keep.pk).exists()
    events = [e for e in _events("lists.field_deleted") if e.object_id == str(pk)]
    assert len(events) == 1
    event = events[0]
    assert event.object_type == "lists.field"
    assert set(event.before) == {"list", "field", "type"}, event.before
    assert event.before["list"] == "Список для удаления поля" and event.before["field"] == "Лишнее"
    assert event.before["type"] == "Строка"


def test_retype_empty_field_resets_foreign_properties_default_thresholds_currency_target_and_choices():
    """FR-509, FR-542: смена типа пустого поля сбрасывает свойства, недоступные новому типу (целевой список и
    «несколько» у записи, «уникальное» и «обязательное» у «да/нет», умолчание и значения списка значений, срок и
    пороги, валюту); журнал `lists.field_changed` показывает тип «было / стало»."""
    schema = _schema()
    manager, lst = _world()
    target = lists.make_list(created_by=manager.user)
    Choice = support.model("lists.Choice")

    record_field = lists.add_field(lst, "record", "Ссылка", target=target, multiple=True)
    schema.update_field(record_field, actor=manager.user, type="string")
    record_field.refresh_from_db()
    assert record_field.type == "string" and record_field.target_id is None and record_field.multiple is False

    string_field = lists.add_field(lst, "string", "Уникальное", unique=True, required=True)
    schema.update_field(string_field, actor=manager.user, type="bool")
    string_field.refresh_from_db()
    assert string_field.type == "bool" and string_field.unique is False and string_field.required is False

    choice_field = lists.add_field(lst, "choice", "Выбор", choices=["А", "Б"])
    choice_field.default = [lists.choice_id(choice_field, "А")]
    choice_field.save(update_fields=["default"])
    schema.update_field(choice_field, actor=manager.user, type="text")
    choice_field.refresh_from_db()
    assert choice_field.type == "text" and choice_field.default is None
    assert not Choice.objects.filter(field=choice_field).exists()

    date_field = lists.add_field(lst, "date", "Срок", deadline=True, thresholds=[30, 7], default={"today": True})
    schema.update_field(date_field, actor=manager.user, type="datetime")
    date_field.refresh_from_db()
    assert date_field.type == "datetime" and date_field.default is None
    assert not date_field.settings.get("deadline") and not date_field.settings.get("thresholds")

    money_field = lists.add_field(lst, "money", "Цена", currency="USD", default="10.00")
    schema.update_field(money_field, actor=manager.user, type="text")
    money_field.refresh_from_db()
    assert money_field.type == "text" and money_field.default is None
    assert not money_field.settings.get("currency")

    event = [e for e in _events("lists.field_changed", date_field)][-1]
    assert event.before["type"] == "Дата" and event.after["type"] == "Дата и время"
    assert (event.after.get("field") or event.before.get("field")) == "Срок"


def test_field_with_data_allows_name_hint_order_show_in_table_default_thresholds_and_enabling_multiple():
    """FR-509: у поля с данными меняются название, подсказка, порядок, «показывать в таблице», значение по умолчанию,
    пороги срока и включается «несколько значений» (значения записей при этом не переписываются)."""
    schema = _schema()
    manager, lst = _world()
    string_field = lists.add_field(lst, "string", "Строка")
    date_field = lists.add_field(lst, "date", "Срок", deadline=True, thresholds=[30])
    person_field = lists.add_field(lst, "person", "Человек")
    employee = support.make_employee()
    record = lists.make_record(lst, "Запись", {string_field.code: "данные", date_field.code: "2030-01-15",
                                               person_field.code: [employee.pk]})
    for field in (string_field, date_field, person_field):
        assert schema.has_data(field) is True
    schema.update_field(string_field, actor=manager.user, name="Новое название", hint="Новая подсказка",
                        show_in_table=True, default="умолчание")
    string_field.refresh_from_db()
    assert (string_field.name, string_field.hint, string_field.show_in_table, string_field.default) == (
        "Новое название", "Новая подсказка", True, "умолчание")
    schema.update_field(date_field, actor=manager.user, thresholds=(60, 14))
    date_field.refresh_from_db()
    assert date_field.settings["thresholds"] == [60, 14] and date_field.settings.get("deadline") is True
    schema.update_field(person_field, actor=manager.user, multiple=True)
    person_field.refresh_from_db()
    assert person_field.multiple is True
    Field = support.model("lists.Field")
    before = list(Field.objects.filter(list=lst).order_by("order").values_list("code", flat=True))
    schema.move_field(string_field, "down", actor=manager.user)
    after = list(Field.objects.filter(list=lst).order_by("order").values_list("code", flat=True))
    assert after == [before[1], before[0], before[2]]
    record.refresh_from_db()
    assert record.values == {string_field.code: "данные", date_field.code: "2030-01-15",
                             person_field.code: [employee.pk]}
    assert schema.has_data(string_field) is True


@pytest.mark.parametrize("case", ["person_live_two", "person_archived_two", "file_live_two", "file_archived_two"])
def test_disabling_multiple_rejected_if_any_record_including_archived_has_two_values_or_two_files(case):
    """FR-509: «несколько значений» выключается, только если ни у одной записи (в том числе в архиве) нет двух
    значений или двух файлов; отказ — `multiple_in_use` у поля `multiple`; когда остаётся по одному, выключение идёт."""
    schema = _schema()
    manager, lst = _world()
    archived = case.endswith("archived_two")
    is_file = case.startswith("file")
    field = lists.add_field(lst, "file" if is_file else "person", "Несколько", multiple=True)
    lists.make_record(lst, "Первая запись")
    first, second = support.make_employee(), support.make_employee()
    if is_file:
        record = lists.make_record(lst, "Вторая запись", {}, archived=archived)
        one = lists.attach_file(record, field, "один.pdf")
        two = lists.attach_file(record, field, "два.pdf")
    else:
        record = lists.make_record(lst, "Вторая запись", {field.code: [first.pk, second.pk]}, archived=archived)
    _raises(lambda: schema.update_field(field, actor=manager.user, multiple=False), "multiple_in_use", "multiple")
    field.refresh_from_db()
    assert field.multiple is True
    if is_file:
        two.delete()
    else:
        support.model("lists.Record").objects.filter(pk=record.pk).update(values={field.code: [first.pk]})
    schema.update_field(field, actor=manager.user, multiple=False)
    field.refresh_from_db()
    assert field.multiple is False


def test_type_currency_and_target_of_field_with_data_rejected():
    """FR-509, FR-511: у поля с данными не меняются тип, валюта и целевой список (служба и форма): отказ
    `field_locked`, у формы — 200 с `[data-error=…][data-code="field_locked"]` и текстом; поле не изменилось."""
    schema = _schema()
    manager, lst = _world()
    target = lists.make_list(created_by=manager.user)
    other_target = lists.make_list(created_by=manager.user)
    number = lists.add_field(lst, "number", "Число")
    money = lists.add_field(lst, "money", "Деньги", currency="RUB")
    link = lists.add_field(lst, "record", "Ссылка", target=target)
    target_record = lists.make_record(target, "Целевая запись")
    lists.make_record(lst, "Запись", {number.code: "5", money.code: "10.00", link.code: [target_record.pk]})
    _raises(lambda: schema.update_field(number, actor=manager.user, type="string"), "field_locked", "type")
    _raises(lambda: schema.update_field(money, actor=manager.user, currency="USD"), "field_locked", "currency")
    # Целевой список службой — объект `List` (5.3, К-13); код отказа — `field="target"` (решение Opus, S1-Q2).
    _raises(lambda: schema.update_field(link, actor=manager.user, target=other_target), "field_locked", "target")
    # Присланный `type` у поля с данными в форме — 200 с `[data-error="type"][data-code="field_locked"]` (10.2).
    response = manager.client.post(_field_url(lst, number), {"name": number.name, "type": "string"})
    assert response.status_code == 200
    assert _all(_doc(response), None, {"data-error": "type", "data-code": "field_locked"})
    response = manager.client.post(_field_url(lst, money), {"name": money.name, "currency": "USD"})
    assert response.status_code == 200
    doc = _doc(response)
    assert _all(doc, None, {"data-error": "currency", "data-code": "field_locked"})
    assert "У поля с данными это не меняется" in _body(response)
    response = manager.client.post(_field_url(lst, link), {"name": link.name, "target": str(other_target.pk)})
    assert response.status_code == 200
    assert _all(_doc(response), None, {"data-error": "target", "data-code": "field_locked"})
    for field in (number, money, link):
        field.refresh_from_db()
    assert number.type == "number" and money.settings.get("currency") == "RUB" and link.target_id == target.pk
    assert not _events("lists.field_changed")


def test_enabling_unique_with_repeats_lists_first_10_number_and_value_counting_archived():
    """FR-509, FR-546: включение «Уникального» при повторах — `unique_repeats` у поля `unique`, в тексте ровно 10
    пар «№N — значение» через «; » при 12 нарушениях; повтор с записью в архиве учитывается (и у числа «100» в
    каноническом виде хранения, раздел 2.3); строки `UniqueValue` не создаются, признак не включён, журнала нет."""
    schema = _schema()
    manager, lst = _world()
    field = lists.add_field(lst, "string", "Повторяется")
    numbers = {}
    for index in range(1, 13):
        for _ in range(2):
            record = lists.make_record(lst, f"З{index}", {field.code: f"дубль{index}"})
            numbers[record.number] = f"дубль{index}"
    error = _raises(lambda: schema.update_field(field, actor=manager.user, unique=True), "unique_repeats", "unique")
    assert error.message.startswith("Значения повторяются: ")
    items = error.message.split(": ", 1)[1].split("; ")
    assert len(items) == 10
    for item in items:
        match = re.fullmatch(r"№(\d+) — (.+)", item)
        assert match, item
        assert numbers[int(match.group(1))] == match.group(2)
    field.refresh_from_db()
    assert field.unique is False
    assert not support.model("lists.UniqueValue").objects.filter(field=field).exists()
    assert not _events("lists.field_changed")

    second = lists.make_list(created_by=manager.user)
    archived_field = lists.add_field(second, "string", "С архивом")
    lists.make_record(second, "Живая", {archived_field.code: "тот же"})
    archived = lists.make_record(second, "В архиве", {archived_field.code: "тот же"}, archived=True)
    error = _raises(lambda: schema.update_field(archived_field, actor=manager.user, unique=True), "unique_repeats",
                    "unique")
    assert f"№{archived.number}" in error.message
    response = manager.client.post(_field_url(second, archived_field), {"name": archived_field.name, "unique": "on"})
    assert response.status_code == 200
    assert _all(_doc(response), None, {"data-error": "unique", "data-code": "unique_repeats"})
    assert "Значения повторяются" in _body(response)

    third = lists.make_list(created_by=manager.user)
    number_field = lists.add_field(third, "number", "Число")
    lists.make_record(third, "Живая", {number_field.code: "100"})
    # Р3: число в `values` хранится канонически («100»); равенство «100» и «100.0» проверяет T-6107.
    archived_number = lists.make_record(third, "В архиве", {number_field.code: "100"}, archived=True)
    error = _raises(lambda: schema.update_field(number_field, actor=manager.user, unique=True), "unique_repeats",
                    "unique")
    assert f"№{archived_number.number}" in error.message


def test_enabling_unique_without_repeats_creates_unique_values_and_disabling_erases_them():
    """FR-509: включение «Уникального» без повторов создаёт строки `UniqueValue` по всем записям (в том числе в
    архиве) в нормализации раздела 2.3 («ИН-1» → «ин-1», число — каноническая строка «12.5»), пустые значения строк не дают;
    выключение стирает строки поля."""
    schema = _schema()
    manager, lst = _world()
    field = lists.add_field(lst, "string", "Номер")
    number_field = lists.add_field(lst, "number", "Количество")
    # Р3: число в `values` хранится канонически («12.5»), `unique.normalize` принимает форму хранения (О3).
    lists.make_record(lst, "Первая", {field.code: "ИН-1", number_field.code: "12.5"})
    lists.make_record(lst, "Вторая", {field.code: " Б-2 "}, archived=True)
    lists.make_record(lst, "Третья", {number_field.code: "7"})
    UniqueValue = support.model("lists.UniqueValue")
    assert not UniqueValue.objects.filter(list=lst).exists()
    schema.update_field(field, actor=manager.user, unique=True)
    schema.update_field(number_field, actor=manager.user, unique=True)
    assert set(UniqueValue.objects.filter(field=field).values_list("norm_value", flat=True)) == {"ин-1", "б-2"}
    assert set(UniqueValue.objects.filter(field=number_field).values_list("norm_value", flat=True)) == {"12.5", "7"}
    field.refresh_from_db()
    assert field.unique is True
    schema.update_field(field, actor=manager.user, unique=False)
    assert not UniqueValue.objects.filter(field=field).exists()
    assert UniqueValue.objects.filter(field=number_field).count() == 2


def test_enabling_required_with_empty_values_lists_first_10_records_counting_archived():
    """FR-509, FR-546: включение «Обязательного» при пустых значениях — `required_missing` у поля `required`, в
    тексте ровно 10 номеров «№N» через «, » при 13 пустых; пустая запись в архиве учитывается; когда заполнены
    все (и архивные), включение проходит."""
    schema = _schema()
    manager, lst = _world()
    field = lists.add_field(lst, "string", "Пустое")
    empties = [lists.make_record(lst, f"З{i}", {}) for i in range(12)]
    empties.append(lists.make_record(lst, "В архиве", {}, archived=True))
    error = _raises(lambda: schema.update_field(field, actor=manager.user, required=True), "required_missing",
                    "required")
    assert error.message.startswith("Не заполнено у записей: ")
    items = error.message.split(": ", 1)[1].split(", ")
    assert len(items) == 10 and len(set(items)) == 10
    valid = {f"№{record.number}" for record in empties}
    assert all(item in valid for item in items), items
    field.refresh_from_db()
    assert field.required is False and not _events("lists.field_changed")

    second = lists.make_list(created_by=manager.user)
    archived_field = lists.add_field(second, "string", "Для архива")
    lists.make_record(second, "Живая", {archived_field.code: "есть"})
    archived = lists.make_record(second, "В архиве", {}, archived=True)
    error = _raises(lambda: schema.update_field(archived_field, actor=manager.user, required=True),
                    "required_missing", "required")
    assert error.message == f"Не заполнено у записей: №{archived.number}"
    response = manager.client.post(_field_url(second, archived_field), {"name": archived_field.name, "required": "on"})
    assert response.status_code == 200
    assert _all(_doc(response), None, {"data-error": "required", "data-code": "required_missing"})
    support.model("lists.Record").objects.filter(pk=archived.pk).update(values={archived_field.code: "заполнено"})
    schema.update_field(archived_field, actor=manager.user, required=True)
    archived_field.refresh_from_db()
    assert archived_field.required is True


def test_schema_and_access_changes_take_exclusive_list_lock_and_record_save_shared_lock_before_schema_read():
    """FR-509: каждое изменение схемы и прав берёт исключительный замок списка (И1) раньше чтения `lists_field`;
    сохранение записи берёт разделяемый замок того же списка, затем `FOR UPDATE` записи, затем читает схему."""
    schema = _schema()
    from apps.lists.services import records

    manager, lst = _world()
    user = manager.user
    gone = lists.make_list(created_by=user)
    string_field = lists.add_field(lst, "string", "Строка")
    choice_field = lists.add_field(lst, "choice", "Выбор", choices=["А", "Б"])
    empty = lists.add_field(lst, "number", "Пустое")
    role = factories.make_role()

    def reversed_codes():
        found = support.model("lists.Field").objects.filter(list=lst).order_by("-order")
        return list(found.values_list("code", flat=True))

    calls = [
        ("add_field", lst, lambda: schema.add_field(lst, name="Новое поле", type="string", actor=user)),
        ("update_field", lst, lambda: schema.update_field(string_field, actor=user, hint="подсказка")),
        ("hide_field", lst, lambda: schema.hide_field(string_field, actor=user)),
        ("show_field", lst, lambda: schema.show_field(string_field, actor=user)),
        ("delete_field", lst, lambda: schema.delete_field(empty, actor=user)),
        ("move_field", lst, lambda: schema.move_field(string_field, "down", actor=user)),
        ("set_order", lst, lambda codes: schema.set_order(lst, codes, actor=user)),
        ("save_choices", lst, lambda: schema.save_choices(
            choice_field, [(lists.choice_id(choice_field, "А"), "А"), (None, "В")], actor=user)),
        ("save_access", lst, lambda: schema.save_access(lst, readers=[("role", role.pk)], editors=[], actor=user)),
        ("delete_list", gone, lambda: schema.delete_list(gone, actor=user)),
    ]
    field_table = f'"{support.model("lists.Field")._meta.db_table}"'
    for name, target, call in calls:
        # новый порядок читается до перехвата SQL: иначе чтение полей самим тестом стоит раньше замка
        codes = reversed_codes() if name == "set_order" else None
        with lists.capture_sql() as sqls:
            call(codes) if name == "set_order" else call()
        locks = lists.list_locks(sqls)
        assert lists.ListLock("exclusive", lists.list_no(target.pk)) in locks, (name, locks)
        assert all(lock.kind == "exclusive" for lock in locks), (name, locks)
        lock_at = next(i for i, sql in enumerate(sqls) if "pg_advisory_xact_lock(" in sql)
        field_at = next((i for i, sql in enumerate(sqls) if field_table in sql), None)
        if field_at is not None:
            assert lock_at < field_at, (name, lock_at, field_at)

    record = lists.make_record(lst, "Запись", {})
    sqls = lists.sql_of(records.save, lst, record.number, {"title": "Новое название"}, base_version=1, actor=user)
    locks = lists.list_locks(sqls)
    assert locks and locks[0] == lists.ListLock("shared", lists.list_no(lst.pk)), locks
    lock_at = next(i for i, sql in enumerate(sqls) if "pg_advisory_xact_lock_shared" in sql)
    field_at = next(i for i, sql in enumerate(sqls) if field_table in sql)
    update_sql = lists.assert_for_update(sqls, support.model("lists.Record")._meta.db_table)
    assert lock_at < sqls.index(update_sql) and lock_at < field_at


# --- скрыть и показать (FR-509, FR-516) --------------------------------------------------------------------

def test_hide_and_show_set_flag_write_field_hidden_and_field_shown_and_recompute_vectors(
        clock, django_capture_on_commit_callbacks):
    """FR-509: «Скрыть» и «Показать» ставят признак `hidden`, пишут `lists.field_hidden` и `lists.field_shown`;
    у поля «строка» ставят `vectors_stale_since` (страница показывает «Поиск по списку обновляется») и после
    фиксации задание `refresh_vectors(список, поле)`, после задания слово только из скрытого поля не находит запись, а
    после показа находит; у поля «число» признака и задания нет."""
    import django_rq

    schema = _schema()
    manager, lst = _world()
    hidden = lists.add_field(lst, "string", "Скрываемое")
    number = lists.add_field(lst, "number", "Число")
    word = "хризантема"
    record = lists.make_record(lst, "Запись", {hidden.code: f"цветок {word}", number.code: "5"})
    assert _vector_matches(record, word)
    queue = django_rq.get_queue("maintenance")

    def jobs_for(field):
        return [job for job in queue.jobs if tuple(job.args) == (lst.pk, field.pk)]

    with django_capture_on_commit_callbacks(execute=True):
        schema.hide_field(hidden, actor=manager.user)
    hidden.refresh_from_db()
    lst.refresh_from_db()
    assert hidden.hidden is True and lst.vectors_stale_since == timezone.now()
    assert len(jobs_for(hidden)) == 1
    page = _doc(manager.client.get(f"/lists/{lst.pk}/"))
    notice = _all(page, None, {"data-lists-notice": "vectors-stale"})
    assert notice and "Поиск по списку обновляется" in notice[0].text
    with django_capture_on_commit_callbacks(execute=True):
        assert run_jobs("maintenance") >= 1
    lst.refresh_from_db()
    assert lst.vectors_stale_since is None
    assert not _vector_matches(record, word)
    assert not _all(_doc(manager.client.get(f"/lists/{lst.pk}/")), None, {"data-lists-notice": "vectors-stale"})
    hidden_events = _events("lists.field_hidden", hidden)
    assert len(hidden_events) == 1 and hidden_events[0].object_type == "lists.field"
    assert _flag(hidden_events[0].after["hidden"]) and not _flag(hidden_events[0].before.get("hidden"))

    with django_capture_on_commit_callbacks(execute=True):
        schema.show_field(hidden, actor=manager.user)
    hidden.refresh_from_db()
    lst.refresh_from_db()
    assert hidden.hidden is False and lst.vectors_stale_since == timezone.now()
    assert len(jobs_for(hidden)) == 1
    with django_capture_on_commit_callbacks(execute=True):
        run_jobs("maintenance")
    lst.refresh_from_db()
    assert lst.vectors_stale_since is None and _vector_matches(record, word)
    shown = _events("lists.field_shown", hidden)
    assert len(shown) == 1 and not _flag(shown[0].after["hidden"]) and _flag(shown[0].before.get("hidden"))

    with django_capture_on_commit_callbacks(execute=True):
        schema.hide_field(number, actor=manager.user)
    lst.refresh_from_db()
    assert lst.vectors_stale_since is None and not jobs_for(number)
    assert len(_events("lists.field_hidden", number)) == 1


def test_hidden_field_values_absent_from_record_card_and_show_returns_them():
    """FR-509, FR-516: значения скрытого поля в карточке записи не видны (решение владельца 2026-10-25), после
    «Показать» возвращаются вместе со значениями."""
    schema = _schema()
    manager, lst = _world()
    visible = lists.add_field(lst, "string", "Видимое")
    hidden = lists.add_field(lst, "string", "Скрываемое")
    record = lists.make_record(lst, "Запись", {visible.code: "видимая-метка", hidden.code: "скрытая-метка"})
    reader = lists.as_reader(lst)
    url = f"/lists/{lst.pk}/{record.number}/"
    page = reader.client.get(url)
    assert _all(_doc(page), None, {"data-lists-value": hidden.code}) and "скрытая-метка" in _body(page)
    schema.hide_field(hidden, actor=manager.user)
    page = reader.client.get(url)
    doc = _doc(page)
    assert not _all(doc, None, {"data-lists-value": hidden.code}) and "скрытая-метка" not in _body(page)
    assert _all(doc, None, {"data-lists-value": visible.code}) and "видимая-метка" in _body(page)
    schema.show_field(hidden, actor=manager.user)
    page = reader.client.get(url)
    shown = _all(_doc(page), None, {"data-lists-value": hidden.code})
    assert shown and "скрытая-метка" in shown[0].text


def test_hidden_field_history_rows_remain_and_field_returns_after_show():
    """FR-509: после «Скрыть» строки истории по полю остаются в базе и в `…/history/`; после «Показать» поле
    возвращается в карточку."""
    schema = _schema()
    manager, lst = _world()
    field = lists.add_field(lst, "string", "Скрываемое")
    record = lists.make_record(lst, "Запись", {field.code: "новое"})
    _change(record, "changed", field.code, field.name, "старое", "новое", manager.user)
    Change = support.model("lists.Change")
    reader = lists.as_reader(lst)
    history = f"/lists/{lst.pk}/{record.number}/history/"
    schema.hide_field(field, actor=manager.user)
    assert Change.objects.filter(record=record).count() == 1
    doc = _doc(reader.client.get(history))
    assert _all(doc, None, {"data-change-field": field.code})
    assert not _all(_doc(reader.client.get(f"/lists/{lst.pk}/{record.number}/")), None,
                    {"data-lists-value": field.code})
    schema.show_field(field, actor=manager.user)
    assert Change.objects.filter(record=record).count() == 1
    assert _all(_doc(reader.client.get(f"/lists/{lst.pk}/{record.number}/")), None, {"data-lists-value": field.code})
    assert _all(_doc(reader.client.get(history)), None, {"data-change-field": field.code})


# --- перечни мест V1–V20 и A1–A9 (раздел 6, решение 27 задач) -----------------------------------------------

def _pair(lst, type, visible_name="Видимое", hidden_name="Скрытое", **extra):
    visible = lists.add_field(lst, type, visible_name, **extra)
    hidden = lists.add_field(lst, type, hidden_name, hidden=True, **extra)
    return visible, hidden


def _v1():
    manager, lst = _world()
    visible = lists.add_field(lst, "string", "Видимое")
    hidden = lists.add_field(lst, "string", "Скрытое", hidden=True, default="умолчание-скрытого")
    page = manager.client.get(f"/lists/{lst.pk}/add/")
    doc = _doc(page)
    assert _all(doc, None, {"data-lists-field": visible.code})
    assert not _all(doc, None, {"data-lists-field": hidden.code}) and "умолчание-скрытого" not in _body(page)


def _v2():
    manager, lst = _world()
    visible, hidden = _pair(lst, "string")
    record = lists.make_record(lst, "Запись", {visible.code: "а", hidden.code: "б"})
    doc = _doc(manager.client.get(f"/lists/{lst.pk}/{record.number}/edit/"))
    assert _all(doc, None, {"name": visible.code}) and not _all(doc, None, {"name": hidden.code})


def _v3():
    from apps.lists.services import records

    manager, lst = _world()
    visible = lists.add_field(lst, "string", "Видимое")
    hidden = lists.add_field(lst, "string", "Скрытое", hidden=True, required=True)
    created = records.create(lst, {"title": "Запись", visible.code: "а"}, manager.user)
    assert support.model("lists.Record").objects.filter(pk=created.pk).exists()


def _v4():
    manager, lst = _world()
    visible, hidden = _pair(lst, "string")
    record = lists.make_record(lst, "Запись", {visible.code: "а", hidden.code: "б"})
    doc = _doc(lists.as_reader(lst).client.get(f"/lists/{lst.pk}/{record.number}/"))
    assert _all(doc, None, {"data-lists-value": visible.code})
    assert not _all(doc, None, {"data-lists-value": hidden.code})


def _v5():
    from apps.lists.services import grid

    manager, lst = _world()
    visible = lists.add_field(lst, "string", "Видимое", show_in_table=True)
    hidden = lists.add_field(lst, "string", "Скрытое", show_in_table=True, hidden=True)
    viewer = lists.as_reader(lst).user
    columns = grid.columns(lst, viewer)
    assert visible.code in columns and hidden.code not in columns
    support.model("lists.UserView").objects.create(user=viewer, list=lst, columns=["title", visible.code, hidden.code])
    columns = grid.columns(lst, viewer)
    assert visible.code in columns and hidden.code not in columns


def _v6():
    manager, lst = _world()
    visible, hidden = _pair(lst, "string")
    lists.make_record(lst, "Запись", {visible.code: "а", hidden.code: "б"})
    doc = _doc(lists.as_reader(lst).client.get(f"/lists/{lst.pk}/"))
    assert _all(doc, None, {"data-lists-column-choice": visible.code})
    assert _all(doc, None, {"data-lists-filter": visible.code})
    assert not _all(doc, None, {"data-lists-column-choice": hidden.code})
    assert not _all(doc, None, {"data-lists-filter": hidden.code})


def _v7():
    from apps.lists.services import grid

    manager, lst = _world()
    visible, hidden = _pair(lst, "string")
    viewer = lists.as_reader(lst).user
    for index in range(3):
        lists.make_record(lst, f"З{index}", {visible.code: "а", hidden.code: f"аа{3 - index}"})
    params = QueryDict(f"f.{hidden.code}=нет-такого&sort={hidden.code}&dir=desc")
    result = list(grid.records(lst, viewer, grid.read(params, lst, viewer)).values_list("number", flat=True))
    assert result == [1, 2, 3], result


def _v8():
    manager, lst = _world()
    visible, hidden = _pair(lst, "choice", choices=["А"])
    reader = lists.as_reader(lst)
    assert reader.client.get(f"/lists/{lst.pk}/filter-values/{visible.code}/").status_code == 200
    assert reader.client.get(f"/lists/{lst.pk}/filter-values/{hidden.code}/").status_code == 404


def _rows(client, url) -> set:
    return {int(el.attrs["data-record-number"]) for el in _doc(client.get(url)).elements
            if el.tag == "tr" and "data-record-number" in el.attrs}


def _v9():
    from apps.core.widgets import moscow_today

    manager, lst = _world()
    reader = lists.as_reader(lst)
    today = moscow_today()
    near, far = (today + timedelta(days=3)).isoformat(), (today + timedelta(days=200)).isoformat()
    hidden = lists.add_field(lst, "date", "Скрытый срок", deadline=True, thresholds=[30], hidden=True)
    first = lists.make_record(lst, "Первая", {hidden.code: near})
    second = lists.make_record(lst, "Вторая", {hidden.code: far})
    base = f"/lists/{lst.pk}/"
    assert not _all(_doc(reader.client.get(base)), "a", {"data-preset": "expiring"})
    assert _rows(reader.client, base + "?preset=expiring") == {first.number, second.number}
    visible = lists.add_field(lst, "date", "Видимый срок", deadline=True, thresholds=[30])
    support.model("lists.Record").objects.filter(pk=first.pk).update(values={hidden.code: near, visible.code: far})
    support.model("lists.Record").objects.filter(pk=second.pk).update(values={hidden.code: far, visible.code: near})
    assert _all(_doc(reader.client.get(base)), "a", {"data-preset": "expiring"})
    assert _rows(reader.client, base + "?preset=expiring") == {second.number}


def _v10():
    manager, lst = _world()
    reader = lists.as_reader(lst)
    other = support.make_employee()
    hidden = lists.add_field(lst, "person", "Скрытый сотрудник", hidden=True)
    first = lists.make_record(lst, "Первая", {hidden.code: [reader.user.pk]})
    second = lists.make_record(lst, "Вторая", {hidden.code: [other.pk]})
    base = f"/lists/{lst.pk}/"
    assert not _all(_doc(reader.client.get(base)), "a", {"data-preset": "mine"})
    assert _rows(reader.client, base + "?preset=mine") == {first.number, second.number}
    visible = lists.add_field(lst, "person", "Видимый сотрудник")
    support.model("lists.Record").objects.filter(pk=first.pk).update(values={hidden.code: [reader.user.pk],
                                                                            visible.code: [other.pk]})
    support.model("lists.Record").objects.filter(pk=second.pk).update(values={hidden.code: [other.pk],
                                                                             visible.code: [reader.user.pk]})
    assert _all(_doc(reader.client.get(base)), "a", {"data-preset": "mine"})
    assert _rows(reader.client, base + "?preset=mine") == {second.number}


def _v11():
    from apps.lists.services import export

    manager, lst = _world()
    visible = lists.add_field(lst, "string", "Видимое")
    hidden = lists.add_field(lst, "string", "Скрытое", hidden=True)
    viewer = lists.as_reader(lst).user
    lists.make_record(lst, "Запись", {visible.code: "а", hidden.code: "б"})
    support.model("lists.UserView").objects.create(user=viewer, list=lst, columns=["title", visible.code, hidden.code])
    header = lists.read_xlsx(export.build(lst, viewer, QueryDict("")).content)[0]
    assert "Видимое" in header and "Скрытое" not in header


def _v12():
    from openpyxl import load_workbook

    from apps.lists.services import imports

    manager, lst = _world()
    lists.add_field(lst, "string", "Видимое")
    lists.add_field(lst, "string", "Скрытое", hidden=True)
    book = load_workbook(io.BytesIO(imports.sample(lst)))
    cells = [cell for sheet in book.worksheets for row in sheet.iter_rows(values_only=True) for cell in row]
    assert "Видимое" in cells and "Скрытое" not in cells


def _v13():
    from apps.lists.services import imports

    manager, lst = _world()
    visible, hidden = _pair(lst, "string")
    mapping = imports.auto_mapping(lst, ["Название", "Видимое", "Скрытое"])
    assert mapping.get("0") == "title" and mapping.get("1") == visible.code and "2" not in mapping


def _v14():
    from apps.lists.services import vectors

    manager, lst = _world()
    visible, hidden = _pair(lst, "text")
    record = lists.make_record(lst, "Запись", {visible.code: "видимыйтекст", hidden.code: "скрытыйтекст"})
    # Помощник `make_record` уже посчитал вектор; стираем его, чтобы проверить именно пересчёт `vectors.refresh`.
    support.model("lists.Record").objects.filter(pk=record.pk).update(search=None)
    assert not _vector_matches(record, "видимыйтекст")
    vectors.refresh([record.pk])
    assert _vector_matches(record, "видимыйтекст") and not _vector_matches(record, "скрытыйтекст")


def _v15():
    from apps.lists.services import search as lists_search

    manager, lst = _world()
    visible, hidden = _pair(lst, "string", unique=True)
    reader = lists.as_reader(lst)
    lists.make_record(lst, "Запись видимая", {visible.code: "qvis77", hidden.code: "qhid77"})
    found = lists_search.search(reader.user, lists_search.parse_local("qvis77"), 10, 0)
    assert [row.title for row in found] == ["Запись видимая"]
    assert lists_search.search(reader.user, lists_search.parse_local("qhid77"), 10, 0) == []


def _v16():
    manager, lst = _world()
    hidden = lists.add_field(lst, "string", "Скрытый номер", unique=True, hidden=True)
    visible = lists.add_field(lst, "string", "Видимый номер", unique=True)
    record = lists.make_record(lst, "Запись", {hidden.code: "ХИДН-1", visible.code: "ВИДН-2"})
    page = lists.as_reader(lst).client.get(f"/lists/{lst.pk}/{record.number}/label/")
    doc = _doc(page)
    value = _all(doc, None, {"data-field": "value"})
    assert value and "ВИДН-2" in value[0].text and "ХИДН-1" not in _body(page)


def _v17():
    from apps.lists.services import links

    manager = lists.as_manager()
    person = support.make_employee()
    first = lists.make_list(created_by=manager.user)
    second = lists.make_list(created_by=manager.user)
    visible = lists.add_field(first, "person", "Видимый сотрудник")
    hidden = lists.add_field(second, "person", "Скрытый сотрудник", hidden=True)
    lists.make_record(first, "Первая", {visible.code: [person.pk]})
    lists.make_record(second, "Вторая", {hidden.code: [person.pk]})
    pinned = {row.field.pk for row in links.pinned(person, person)}
    assert visible.pk in pinned and hidden.pk not in pinned


def _reminder_world():
    from apps.core.widgets import moscow_today

    manager, lst = _world()
    today = moscow_today()
    return manager, lst, today, (today + timedelta(days=3)).isoformat()


def _v18():
    from apps.lists.services import reminders

    manager, lst, today, near = _reminder_world()
    visible = lists.add_field(lst, "date", "Видимый срок", deadline=True, thresholds=[7])
    hidden = lists.add_field(lst, "date", "Скрытый срок", deadline=True, thresholds=[7], hidden=True)
    record = lists.make_record(lst, "Запись", {visible.code: near, hidden.code: near})
    found = {(c.record_id, c.field_id) for c in reminders.select_candidates(today)}
    assert (record.pk, visible.pk) in found and (record.pk, hidden.pk) not in found


def _v19():
    from apps.lists.services import reminders

    manager, lst, today, near = _reminder_world()
    deadline = lists.add_field(lst, "date", "Срок", deadline=True, thresholds=[7])
    visible = lists.add_field(lst, "person", "Видимый сотрудник")
    hidden = lists.add_field(lst, "person", "Скрытый сотрудник", hidden=True)
    seen, unseen = support.make_employee(), support.make_employee()
    lists.make_record(lst, "Запись", {deadline.code: near, visible.code: [seen.pk], hidden.code: [unseen.pk]})
    reminders.run(today)
    Notification = support.model("notify.Notification")
    assert Notification.objects.filter(user=seen, kind="lists.deadline").exists()
    assert not Notification.objects.filter(user=unseen, kind="lists.deadline").exists()


def _v20():
    from apps.lists.services import records
    from apps.lists.services.errors import ValuesInvalid

    manager, lst = _world()
    visible, hidden = _pair(lst, "string")
    record = lists.make_record(lst, "Запись", {visible.code: "а"})
    with pytest.raises(ValuesInvalid) as caught:
        records.save(lst, record.number, {"title": "Запись", visible.code: "а", hidden.code: "новое"},
                     base_version=record.version, actor=manager.user)
    assert caught.value.errors[hidden.code].code == "field_hidden_meanwhile"


_V_PLACES = {"V1": _v1, "V2": _v2, "V3": _v3, "V4": _v4, "V5": _v5, "V6": _v6, "V7": _v7, "V8": _v8, "V9": _v9,
             "V10": _v10, "V11": _v11, "V12": _v12, "V13": _v13, "V14": _v14, "V15": _v15, "V16": _v16, "V17": _v17,
             "V18": _v18, "V19": _v19, "V20": _v20}


@pytest.mark.parametrize("place", list(_V_PLACES))
def test_visible_fields_places_from_contract_skip_hidden_field(place):
    """FR-509: места V1–V20 раздела 6.1 контракта не используют скрытое поле — форма, проверка обязательности,
    карточка, столбцы, фильтры, отборы, выгрузка, образец, сопоставление, вектор, точное совпадение, этикетка,
    «Закреплённое», напоминания, сохранение. (Часть «привязка к заявке» места V15 — вопрос к Opus в отчёте.)"""
    _V_PLACES[place]()


def _a1():
    from apps.lists.services import unique

    manager, lst = _world()
    hidden = lists.add_field(lst, "string", "Скрытое", unique=True, hidden=True)
    first = lists.make_record(lst, "Первая", {hidden.code: "АБВ"})
    UniqueValue = support.model("lists.UniqueValue")
    UniqueValue.objects.filter(record=first).delete()
    unique.sync(first, [hidden])
    assert list(UniqueValue.objects.filter(field=hidden).values_list("norm_value", flat=True)) == ["абв"]
    second = lists.make_record(lst, "Вторая", {hidden.code: "другое"})
    second.values = {hidden.code: "абв"}
    _raises(lambda: unique.sync(second, [hidden]), "unique_taken", hidden.code)


def _a2():
    schema = _schema()
    manager, lst = _world()
    unique_field = lists.add_field(lst, "string", "Уникальное", hidden=True)
    lists.make_record(lst, "Живая", {unique_field.code: "тот же"})
    lists.make_record(lst, "В архиве", {unique_field.code: "тот же"}, archived=True)
    _raises(lambda: schema.update_field(unique_field, actor=manager.user, unique=True), "unique_repeats", "unique")
    required_field = lists.add_field(lst, "string", "Обязательное", hidden=True)
    _raises(lambda: schema.update_field(required_field, actor=manager.user, required=True), "required_missing",
            "required")
    people = lists.add_field(lst, "person", "Люди", multiple=True, hidden=True)
    first, second = support.make_employee(), support.make_employee()
    lists.make_record(lst, "С двумя", {people.code: [first.pk, second.pk]}, archived=True)
    _raises(lambda: schema.update_field(people, actor=manager.user, multiple=False), "multiple_in_use", "multiple")


def _a3():
    schema = _schema()
    manager, lst = _world()
    hidden = lists.add_field(lst, "string", "Скрытое", hidden=True)
    lists.make_record(lst, "Запись", {hidden.code: "данные"})
    assert schema.has_data(hidden) is True


def _a4():
    schema = _schema()
    manager, lst = _world()
    for index in range(49):
        lists.add_field(lst, "string", f"Видимое {index}")
    lists.add_field(lst, "string", "Скрытое", hidden=True)
    _raises(lambda: schema.add_field(lst, name="Пятьдесят первое", type="string", actor=manager.user), "fields_limit")


def _a5():
    from apps.lists.services import archive, links

    manager = lists.as_manager()
    target_list = lists.make_list(created_by=manager.user)
    holder_list = lists.make_list(created_by=manager.user)
    hidden = lists.add_field(holder_list, "record", "Скрытая ссылка", target=target_list, hidden=True)
    target = lists.make_record(target_list, "Цель", {}, archived=True)
    holder = lists.make_record(holder_list, "Держатель", {hidden.code: [target.pk]})
    assert [row.pk for row in links.referrers(target, None).rows] == [holder.pk]
    _raises(lambda: archive.purge(target_list, target.number, actor=manager.user), "record_referenced")
    assert support.model("lists.Record").objects.filter(pk=target.pk).exists()


def _a6():
    from apps.lists.services import links

    manager, lst = _world()
    hidden = lists.add_field(lst, "unit", "Скрытое подразделение", hidden=True)
    lists.make_record(lst, "Запись", {hidden.code: [987654]})
    text = links.unit_usage(987654)
    assert text and "в 1 записях" in text


def _a7():
    manager, lst = _world()
    hidden = lists.add_field(lst, "string", "Скрытое", hidden=True)
    record = lists.make_record(lst, "Запись", {hidden.code: "новое"})
    _change(record, "changed", hidden.code, hidden.name, "старое", "новое", manager.user)
    doc = _doc(lists.as_reader(lst).client.get(f"/lists/{lst.pk}/{record.number}/history/"))
    assert _all(doc, None, {"data-change-field": hidden.code})


def _a8():
    manager, lst = _world()
    hidden = lists.add_field(lst, "string", "Скрытое", hidden=True)
    doc = _doc(manager.client.get(_fields_url(lst)))
    row = _row(doc, hidden.code)
    assert row is not None and "field-show" in _buttons(doc, row)


def _a9():
    schema = _schema()
    manager = lists.as_manager()
    target_list = lists.make_list(created_by=manager.user)
    holder_list = lists.make_list(created_by=manager.user)
    lists.add_field(holder_list, "record", "Скрытая ссылка", target=target_list, hidden=True)
    _raises(lambda: schema.delete_list(target_list, actor=manager.user), "list_referenced")
    assert support.model("lists.List").objects.filter(pk=target_list.pk).exists()


_A_PLACES = {"A1": _a1, "A2": _a2, "A3": _a3, "A4": _a4, "A5": _a5, "A6": _a6, "A7": _a7, "A8": _a8, "A9": _a9}


@pytest.mark.parametrize("place", list(_A_PLACES))
def test_all_fields_places_from_contract_include_hidden_field(place):
    """FR-509: места A1–A9 раздела 6.2 контракта учитывают скрытое поле — уникальность, проверки включения
    свойств, «поле с данными», предел 50 полей, «Ссылаются на эту запись» и запрет удаления навсегда, `unit_usage`,
    история, вкладка «Поля», удаление списка."""
    _A_PLACES[place]()


def test_hidden_required_field_is_not_checked_on_record_save():
    """FR-509, FR-515: скрытое обязательное поле при создании и сохранении записи на обязательность не
    проверяется (сценарий 14 п. 7); видимое обязательное — проверяется."""
    from apps.lists.services import records
    from apps.lists.services.errors import ValuesInvalid

    manager, lst = _world()
    visible = lists.add_field(lst, "string", "Видимое", required=True)
    hidden = lists.add_field(lst, "string", "Скрытое", required=True, hidden=True)
    created = records.create(lst, {"title": "Запись", visible.code: "а"}, manager.user)
    result = records.save(lst, created.number, {"title": "Запись 2", visible.code: "б"},
                          base_version=created.version, actor=manager.user)
    assert result.changed is True
    with pytest.raises(ValuesInvalid) as caught:
        records.create(lst, {"title": "Запись 3", hidden.code: ""}, manager.user)
    assert caught.value.errors[visible.code].code == "required" and hidden.code not in caught.value.errors


# --- значения списка значений (FR-510) --------------------------------------------------------------------

def test_choices_200_accepted_201st_title_101_and_case_duplicate_rejected_and_order_saved():
    """FR-510: до 200 значений проходят, 201-е — `choices_limit`; название длиной 101 — `too_long` у поля
    `choices` (предел «100»); дубль без учёта регистра — `choice_name_taken`; порядок сохраняется; форма на дубль
    отвечает 200 с `[data-error="choices"]`, на новый порядок — 302."""
    schema = _schema()
    manager, lst = _world()
    field = lists.add_field(lst, "choice", "Выбор")
    Choice = support.model("lists.Choice")
    items = [(None, f"Значение {index}") for index in range(200)]
    schema.save_choices(field, items, actor=manager.user)
    stored = list(Choice.objects.filter(field=field).order_by("order"))
    assert [c.title for c in stored] == [title for _, title in items]
    orders = [c.order for c in stored]
    assert orders == sorted(orders) and len(set(orders)) == 200
    existing = [(c.pk, c.title) for c in stored]
    _raises(lambda: schema.save_choices(field, existing + [(None, "Лишнее")], actor=manager.user), "choices_limit")
    error = _raises(lambda: schema.save_choices(field, existing[:5] + [(None, "я" * 101)], actor=manager.user),
                    "too_long", "choices")
    assert error.params["limit"] == "100"
    _raises(lambda: schema.save_choices(field, existing[:5] + [(None, "ЗНАЧЕНИЕ 0")], actor=manager.user),
            "choice_name_taken", "choices")
    # Латинская пара — независимо от локали кластера (кириллица выше проверяет требование для русских названий).
    latin = lists.add_field(lst, "choice", "Латиница", choices=["Alpha"])
    _raises(lambda: schema.save_choices(latin, [(lists.choice_id(latin, "Alpha"), "Alpha"), (None, "ALPHA")],
                                        actor=manager.user), "choice_name_taken", "choices")
    assert Choice.objects.filter(field=field).count() == 200
    schema.save_choices(field, list(reversed(existing)), actor=manager.user)
    assert [c.title for c in Choice.objects.filter(field=field).order_by("order")] == [t for _, t in reversed(existing)]

    small = lists.add_field(lst, "choice", "Малое", choices=["Да", "Нет"])
    yes, no = lists.choice_id(small, "Да"), lists.choice_id(small, "Нет")
    url = _field_url(lst, small) + "choices/"
    response = manager.client.post(url, {"id": [str(yes), str(no), ""], "title": ["Да", "Нет", "да"]})
    assert response.status_code == 200
    assert _all(_doc(response), None, {"data-error": "choices", "data-code": "choice_name_taken"})
    response = manager.client.post(url, {"id": [str(no), str(yes)], "title": ["Нет", "Да"]})
    assert response.status_code == 302 and response["Location"].endswith(_field_url(lst, small))
    assert [c.title for c in Choice.objects.filter(field=small).order_by("order")] == ["Нет", "Да"]


def test_used_choice_including_by_archived_record_goes_to_value_archive_and_unused_is_deleted():
    """FR-510: значение, которое есть у записи (в том числе у записи в архиве), уходит в архив значений
    (`archived=True`), неиспользованное удаляется; архивный id в новом списке возвращает значение из архива;
    журнал `lists.choices_changed` показывает названия не в архиве."""
    schema = _schema()
    manager, lst = _world()
    field = lists.add_field(lst, "choice", "Выбор", choices=["А", "Б", "В", "Г"])
    ids = {title: lists.choice_id(field, title) for title in "АБВГ"}
    lists.make_record(lst, "Живая", {field.code: [ids["А"]]})
    lists.make_record(lst, "В архиве", {field.code: [ids["Б"]]}, archived=True)
    Choice = support.model("lists.Choice")
    schema.save_choices(field, [(ids["Г"], "Г")], actor=manager.user)
    assert Choice.objects.get(pk=ids["А"]).archived is True and Choice.objects.get(pk=ids["Б"]).archived is True
    assert not Choice.objects.filter(pk=ids["В"]).exists()
    assert Choice.objects.get(pk=ids["Г"]).archived is False
    schema.save_choices(field, [(ids["Г"], "Г"), (ids["А"], "А")], actor=manager.user)
    assert Choice.objects.get(pk=ids["А"]).archived is False and Choice.objects.get(pk=ids["Б"]).archived is True
    events = _events("lists.choices_changed", field)
    assert events and events[-1].after["choices"] == "Г, А" and events[-1].before["choices"] == "Г"


def test_archived_default_choice_is_reset_and_not_put_into_new_record_form():
    """FR-510: значение по умолчанию, ушедшее в архив значений (или удалённое), сбрасывается у поля; в форму новой
    записи значение из архива не подставляется и не предлагается."""
    schema = _schema()
    manager, lst = _world()
    field = lists.add_field(lst, "choice", "Выбор", choices=["А", "Б"])
    ids = {title: lists.choice_id(field, title) for title in "АБ"}
    field.default = [ids["Б"]]
    field.save(update_fields=["default"])
    lists.make_record(lst, "Запись", {field.code: [ids["Б"]]})
    schema.save_choices(field, [(ids["А"], "А")], actor=manager.user)
    field.refresh_from_db()
    assert support.model("lists.Choice").objects.get(pk=ids["Б"]).archived is True and field.default is None
    doc = _doc(manager.client.get(f"/lists/{lst.pk}/add/"))
    container = _all(doc, None, {"data-lists-field": field.code})[0]
    assert not [el for el in doc.elements if _inside(el, container) and el.attrs.get("value") == str(ids["Б"])]
    # Пустой вариант `<option value="" selected>` допустим; выбранного непустого значения быть не должно.
    assert not [el for el in doc.elements
                if _inside(el, container) and "selected" in el.attrs and el.attrs.get("value")]
    other = lists.add_field(lst, "choice", "Второе", choices=["Х", "Ы"])
    unused = lists.choice_id(other, "Х")
    other.default = [unused]
    other.save(update_fields=["default"])
    schema.save_choices(other, [(lists.choice_id(other, "Ы"), "Ы")], actor=manager.user)
    other.refresh_from_db()
    assert other.default is None and not support.model("lists.Choice").objects.filter(pk=unused).exists()


def test_renamed_choice_shown_in_records_archived_title_rejected_and_choices_changed_journal():
    """FR-510, FR-546: переименованное значение показывается у записей по-новому (id тот же); название значения,
    уже стоящего в архиве значений, — `choice_archived_exists` у поля `choices`; журнал `lists.choices_changed`
    хранит «было / стало» названиями не в архиве через «, »."""
    schema = _schema()
    manager, lst = _world()
    field = lists.add_field(lst, "choice", "Выбор", choices=["Старое", "Другое"])
    old, other = lists.choice_id(field, "Старое"), lists.choice_id(field, "Другое")
    record = lists.make_record(lst, "Запись", {field.code: [old]})
    lists.make_record(lst, "Вторая", {field.code: [other]})
    schema.save_choices(field, [(old, "Новое"), (other, "Другое")], actor=manager.user)
    assert support.model("lists.Choice").objects.get(pk=old).title == "Новое"
    shown = _all(_doc(lists.as_reader(lst).client.get(f"/lists/{lst.pk}/{record.number}/")), None,
                 {"data-lists-value": field.code})
    assert shown and "Новое" in shown[0].text and "Старое" not in shown[0].text
    first_event = _events("lists.choices_changed", field)[0]
    assert first_event.before["choices"] == "Старое, Другое" and first_event.after["choices"] == "Новое, Другое"
    schema.save_choices(field, [(old, "Новое")], actor=manager.user)
    assert support.model("lists.Choice").objects.get(pk=other).archived is True
    _raises(lambda: schema.save_choices(field, [(old, "Новое"), (None, "другое")], actor=manager.user),
            "choice_archived_exists", "choices")
    last = _events("lists.choices_changed", field)[-1]
    assert last.before["choices"] == "Новое, Другое" and last.after["choices"] == "Новое"
